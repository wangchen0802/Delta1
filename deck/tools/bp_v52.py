"""Build SimReal (衍真) BP v52 (15 slides, Chinese) from v17's slides.

v39 (the version investors liked) with their one note acted on: show, at a glance, what the product is. A new page
draws the training-environment loop and walks it through trading; the product page shows a screen for each live
environment; the data page shows what a trajectory looks like; results are a chart. Traction comes forward after
the team (ARR, first three weeks, speed) with backers and the expert network on the next page. No round terms,
no month-by-month targets. Copy tightened.

Usage: python3 visuals_v52.py; python3 bp_v52.py SimReal-BP-v17.pptx out.pptx
       then embed_cjk.py adds the Chinese font subsets to the shipping PPTX.
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, e, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['我们是谁', '当前进展', '问题与时机', '产品', '客户与市场', '商业模式']
CONF = '机密  ·  仅供受邀投资机构内部评估使用  ·  2026年10月'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.36, 6.98, 0.5, 0.26, '衍真', 9, C['ink'], SANS, True, anchor='ctr')
    sh.t(1.85, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def contact_runs(sz=11):
    return [(R('business@simreal.co', sz, C['ink'], MONO), R('   ·   ', sz, C['grey']), R('simreal.com.cn', sz, C['ink'], MONO))]


def arrow_r(sh, x, y, w, color=None):
    """A horizontal arrow: a thin bar with a triangle head."""
    sh.rect(x, y - 0.012, w - 0.14, 0.024, color or C['grey'])
    sh.tri(x + w - 0.07, y, 0.14, 0.16, 90, color or C['grey'])


class S52(S):
    def tri(self, cx, cy, w, h, rot, fill):
        i = self._id()
        self.xml.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="SimReal {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm rot="{int((rot % 360) * 60000)}"><a:off x="{e(cx - w / 2)}" y="{e(cy - h / 2)}"/>'
            f'<a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm><a:prstGeom prst="triangle"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr/></a:p></p:txBody></p:sp>')


def p_cover(sh):
    sh.t(1.62, 0.55, 1.0, 0.25, '衍真', 13, C['ink'], SANS, True, anchor='ctr')
    sh.t(0.6, 1.35, 8.6, 1.9, ['让AI在真实世界里', '自我进化'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, 'AI的训练环境：做真实任务，按真实结果打分', 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, '面向AI实验室与企业', 16, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    for (k, v), x in zip([('年化收入（ARR）', '700万美元'), ('商业计划书', '2026年10月'), ('联系', 'business@simreal.co')], [0.6, 3.6, 6.2]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['accent'] if x == 0.6 else C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, '概述', 0, [('给AI实验室做训练环境，', F), ('模型在里面做真实任务，按真实结果打分', T)])
    cells = [
        ('做什么', '训练环境与AI数据', ['交易、AI研究、事件预测已上线', '卖环境授权、Agent轨迹、专家与评测数据'], F),
        ('收入', 'ARR 700万美元', ['创立三周交付数百万需求', '零外部融资'], T),
        ('结果', '+12%', ['Qwen3.8-27B在Xitadel训练后', '未见过的交易日上最高提升12%，多次独立复现'], F),
        ('团队', '05后量化团队', ['剑桥、LSE、杜克数学', 'Jane Street、Citadel、Optiver、Millennium'], F),
        ('速度', '14天7款产品', ['5个公开仓库，512星', '2家前沿实验室在谈'], F),
        ('市场', '85亿美元', ['训练数据与RL环境供应商年收入', 'Mercor年化毛营收16个月涨27倍'], F),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.72 + (i // 3) * 2.0
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, v, 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.8, d, 11, C['body'], line=1.15, gap=0)
    kicker(sh, 5.85, [('谁有最好的训练世界，谁就有最好的AI。', F), ('从交易出发，走向整个世界。', T)], 18)
    footer(sh)


def p_team(sh):
    header(sh, '团队', 0, [('数学、机器学习、量化与数据科学背景，', F), ('放弃量化机构转正offer，全职做衍真', T)])
    b, sz = C['body'], 10.5
    people = [
        ('Charles', 'CEO', ['United Stables首位员工：帮助U稳定币一年内从0做到14亿美元，一个月上线Binance',
                            '负责机构关系，参与SIG、DRW等合作',
                            (R('汇丰港元稳定币发行项目唯一实习生，', sz, b), BR(sz), R('全程协助推进香港金管局（HKMA）合规项目', sz, b)),
                            'X（Twitter）博主，内容数百万浏览', '伦敦Citadel对冲基金实习'],
         ['2005年生 · LSE数学本科 · 深国交', '美国数学奥林匹克（USAMO）入围']),
        ('Henry', 'CTO', [(R('剑桥机器学习暑研，师从统计学教授', sz, b), BR(sz), R('Po-Ling Loh（国际数理统计学会会士）', sz, b)),
                          '剑桥研究中心AI最年轻本科研究员', 'Jane Street、Citadel、Optiver量化经历', '设计五大Benchmark与强化学习环境'],
         ['2005年生 · 剑桥数学一等荣誉（奖学金） · 深国交', '剑桥数学竞赛全球前30 · 英国物理竞赛超级金奖']),
        ('Amaris', 'COO', ['Millennium香港数据科学家', 'Millennium另类数据团队首位应届招聘',
                           (R('参与落地Plug and Play香港首场活动', sz, b), BR(sz), R('（联合香港科技园；200+人规模）', sz, b)),
                           '强生MedTech科技峰会主持人（200+人规模）', '17岁成为出版作家；全网原创内容获10万+互动'],
         ['2005年生 · 杜克数学与统计本科 · 上海包玉刚']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.55
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.76, cw - 0.6, 2.0, lines, sz, b, line=1.1, gap=4)
        sh.rule(x + 0.3, y0 + 2.86, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 2.94, cw - 0.6, 0.55, edu, 9.5, C['grey'], line=1.12, gap=0)
    kicker(sh, 6.1, [('交易台纪律、机器学习研究、机构关系，', F), ('做专业训练环境要的三样能力', T)], 17)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_progress(sh):
    header(sh, '当前进展', 1, [('零外部融资，创立三周交付数百万需求，', F), ('ARR 700万美元', T)])
    stats = [('收入', '700万美元', 'ARR', '零外部融资'), ('交付', '3周', '创立三周', '交付数百万需求'),
             ('产品', '14天', '上线7款产品', '5个公开仓库，512星'), ('客户', '2家', '前沿实验室在谈', '下一批大客户')]
    cw, xs = cols(4, 0.25)
    top, ch = 1.72, 2.55
    for i, ((tag, v, k, m), x) in enumerate(zip(stats, xs)):
        dark = i == 0
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(x + 0.28, top + 0.22, cw - 0.5, 0.22, tag, 9.5, C['accentLt'] if dark else C['accent'], MONO, True)
        sh.t(x + 0.28, top + 0.52, cw - 0.5, 0.8, v, 36 if len(v) > 4 else 44, C['accentLt'] if dark else C['ink'], SERIF)
        sh.t(x + 0.28, top + 1.48, cw - 0.5, 0.3, k, 13, C['onDarkHi'] if dark else C['ink'], SANS, True)
        means(sh, x + 0.28, top + 1.86, cw - 0.5, m, 11.5, dark)
    label(sh, 0.6, 4.55, 8, '3个环境已上线')
    worlds = [('交易', 'Xitadel-QuantBench', '回放真实订单簿，对标专业交易员'), ('AI研究', 'SimReal-MLBench', '60个真实竞赛任务'),
              ('事件预测', 'FuturePredict Bench', '截止前锁定概率，按真实结局打分')]
    wcw, wxs = cols(3, 0.3)
    for (k, name, d), x in zip(worlds, wxs):
        sh.rule(x, 4.84, wcw, C['ink'])
        sh.text(x, 4.94, wcw, 0.44, [para([R(k, 18, C['ink'], SERIF), R('   ' + name, 9.5, C['grey'], MONO)])], 'ctr')
        sh.t(x, 5.42, wcw, 0.3, d, 11, C['body'])
    note(sh, '截至2026年10月；ARR为年化经常性收入；星标截至2026年10月2日。', 6.3)
    footer(sh)


def p_backers(sh, logos, polymarket, schools):
    header(sh, '支持者与专家网络', 1, [('交易机构的从业者支持研发，', F), ('20万+专家可触达', T)])
    label(sh, 0.6, 1.72, 8, '支持研发的从业者来自')
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 2.3, 1.45 if pic.shape_id == 403 else 1.2)
    if polymarket is not None:
        place(polymarket, slots[4], 2.3, 1.2)
    sh.rule(0.6, 2.86, W, C['ink'])
    net = [('20万+', '可触达、可验证的专家网络', T), ('7,000+', '专家候补名单', F), ('21所', '高校学生与校友网络', F)]
    cw, xs = cols(3)
    for (v, k, acc), x in zip(net, xs):
        sh.t(x, 3.0, cw, 0.74, v, 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 3.76, cw, 0.3, k, 12.5, C['ink'], SANS, True)
    label(sh, 0.6, 4.26, 8, '网络覆盖的部分高校')
    cell = W / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, 0.6 + (i % 6 + 0.5) * cell, 4.86 + (i // 6) * 0.78, min(1.2 / w, 0.56 / h, 1.25))
    note(sh, '标识仅表示支持者任职机构或网络覆盖的部分高校，不代表机构或学校背书；网络覆盖不等于已注册或参与交付。', 6.4, 8)
    footer(sh)


def p_problem(sh):
    header(sh, '问题', 2, [('AI会推理，', F), ('做不好真实世界里的事', T)])
    sh.t(0.6, 1.64, W, 0.3, '真实工作没有标准答案，只有世界的反应', 12, C['grey'])
    rows = [('交易', '历史回测一路上涨', '没见过的行情里亏钱'), ('软件工程', '测试全部通过', '上线后出错'),
            ('财务', '账看起来做完了', '月结对不平'), ('事件预测', '分析头头是道', '结果出来就错')]
    y0, rh = 2.38, 0.62
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, '表面检查：通过', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, '真实世界：失手', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 17, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 17, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    sh.t(0.6, 4.95, 1.8, 0.55, '6/32', 26, C['accent'], SERIF, anchor='ctr')
    sh.t(2.45, 4.95, 9.5, 0.55, '前沿模型用真钱交易32轮，只有6轮赚钱（Alpha Arena）', 12, C['body'], anchor='ctr')
    kicker(sh, 5.75, [('AI要学会做事，需要一个', F), ('按真实结果打分的世界', T), ('。', F)])
    footer(sh)


def p_why_now(sh):
    header(sh, '为什么是现在', 2, [('训练的瓶颈从数据转向环境，', F), ('专业领域的环境还没人做好', T)])
    items = [('01', '训练方式变了', ['推理模型靠可自动验证的结果做强化学习', '公开人类文本预计2026–2032年用尽'],
              '能自动打分的环境成了稀缺品'),
             ('02', '预算已经到位', ['每个RL任务200–2,000美元', '环境合同每季度六到七位数美元', 'Mercor 2026年7月收购环境公司Deeptune'],
              '环境成了实验室的采购项'),
             ('03', '专业任务还没解决', ['前沿模型真钱交易，32轮仅6轮盈利', (R('Meta Muse、OpenAI Dots', 11.5, C['body']), BR(11.5),
                                                              R('2026年9月开始替人做事', 11.5, C['body']))],
              '交易对错易验证、差距大，先做交易')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for (n, k, ev, m), x in zip(items, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 21, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.08, cw - 0.6, 1.2, ev, 11.5, C['body'], line=1.15, gap=5)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, m, 12)
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, '窗口期', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R('RL环境约20家早期公司，预计收敛到3–5家（Wing VC）。', 12.5, C['onDark']),
                                       R('先做出可验证的增益，先拿实验室长期合同', 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr', line=1.1)
    note(sh, '来源：Epoch AI（公开文本存量预测；《An FAQ on RL environments》，2026年1月）；SiliconANGLE、TechCrunch（2026年7月）；'
             'Nof1（2026年5月）；Meta、OpenAI发布（2026年9月）；Wing VC（2026年）。', 6.3)
    footer(sh)


def p_env(sh):
    """What a training environment is, in one picture: the loop, then the same loop for trading."""
    header(sh, '训练环境是什么', 3, [('AI的模拟舱：', F), ('做真实任务，按真实结果打分，再用分数训练', T)])
    bw, gap, by, bh = 3.43, 0.92, 1.78, 1.32
    bx = [0.6, 0.6 + bw + gap, 0.6 + 2 * (bw + gap)]
    boxes = [(bx[0], '模型 / Agent', '实验室的模型，或企业的Agent', F),
             (bx[1], '训练环境', '复刻真实任务，例如交易', T),
             (bx[2], '打分', '按真实结果算分，不靠人工判断', F)]
    for x, k, d, dark in boxes:
        sh.rect(x, by, bw, bh, C['ink'] if dark else C['tint'])
        sh.t(x + 0.3, by + 0.2, bw - 0.6, 0.5, k, 22, C['accentLt'] if dark else C['ink'], SERIF)
        sh.t(x + 0.3, by + 0.78, bw - 0.6, 0.4, d, 11.5, C['onDark'] if dark else C['body'])
    for x, lab in [(bx[0] + bw, '动作'), (bx[1] + bw, '结果')]:
        arrow_r(sh, x + 0.1, by + bh / 2 + 0.12, gap - 0.2, C['ink'])
        sh.t(x, by + bh / 2 - 0.32, gap, 0.3, lab, 12, C['ink'], SANS, True, algn='ctr')
    ry = by + bh + 0.32
    sh.rect(bx[0] + bw / 2, by + bh, 0.024, 0.32, C['accent'])
    sh.rect(bx[2] + bw / 2, by + bh, 0.024, 0.32, C['accent'])
    sh.rect(bx[0] + bw / 2, ry, bx[2] - bx[0] + 0.024, 0.024, C['accent'])
    sh.tri(bx[0] + bw / 2 + 0.012, by + bh + 0.08, 0.16, 0.14, 0, C['accent'])
    sh.t(3.0, ry + 0.06, 7.4, 0.3, '分数变成奖励，强化学习更新模型，再进环境', 11.5, C['accent'], SANS, True, algn='ctr')
    label(sh, 0.6, 4.02, 6, '以交易为例')
    steps = [('动作', '下单、撤单、改价'), ('环境', '回放真实订单簿，逐笔撮合'), ('打分', '盈亏扣回撤，再乘夏普'),
             ('结果', '训练后，未见过的交易日最高+12%')]
    cw, xs = cols(4, 0.3)
    for i, ((k, d), x) in enumerate(zip(steps, xs)):
        sh.rule(x, 4.3, cw, C['accent'] if i == 3 else C['ink'])
        sh.t(x, 4.4, cw, 0.36, k, 16, C['accent'] if i == 3 else C['ink'], SERIF)
        sh.t(x, 4.8, cw, 0.5, d, 11.5, C['body'], line=1.1)
    sh.rect(0.6, 5.5, W, 0.62, C['tint'])
    sh.t(0.85, 5.5, 1.8, 0.62, '我们交付', 14, C['accent'], SERIF, anchor='ctr')
    sh.t(2.75, 5.5, W - 2.4, 0.62, '训练环境  ·  评分工具  ·  Agent轨迹与专家数据  ·  评测报告', 14, C['ink'], SERIF, anchor='ctr')
    note(sh, '验收任务不参与训练；+12%为最佳一次，多次独立复现。', 6.36)
    footer(sh)


PRODUCT_IMAGES = ['xitadel.png', 'mlbench.png', 'forecast.png']


def p_products(sh):
    header(sh, '产品', 3, [('3个环境已上线，', F), ('5个基准已开源', T)])
    items = [('Xitadel', '交易', ['回放真实订单簿，和同资产上的专业交易员比', '下单 → 撮合成交 → 按盈亏、回撤、夏普打分']),
             ('MLBench', 'AI研究', ['60个真实机器学习竞赛任务', '读数据、写代码、训练、提交，对照真实榜单打分']),
             ('FuturePredict', '事件预测', ['截止前锁定概率和证据', '事件揭晓后，按真实结局打分'])]
    cw, xs = cols(3, 0.3)
    for (name, dom, lines), x in zip(items, xs):
        sh.text(x, 4.6, cw, 0.42, [para([R(name, 19, C['ink'], SERIF), R('   ' + dom, 10, C['accent'], MONO, True)])], 'ctr')
        sh.t(x, 5.04, cw, 0.62, lines, 11, C['body'], line=1.12, gap=2)
    sh.rule(0.6, 5.76, W)
    sh.t(0.6, 5.82, W, 0.32, [(R('另外4款  ', 10, C['accent'], MONO, True),
                               R('MathmoBench 数学  ·  Puzzle Benchmark 推理  ·  Month-End Close 财务结账  ·  SWE-Forward 软件工程', 12, C['ink']))],
         12, anchor='ctr')
    note(sh, '界面为示意图。公开仓库截至2026年10月2日，5个共512星。', 6.4)
    footer(sh)


def p_data(sh):
    header(sh, '数据', 3, [('每跑一个回合，', F), ('就留下一条能卖的数据', T)])
    label(sh, 0.6, 1.66, 8, '一条Agent轨迹（交易，示意）')
    heads = ['步', '环境给出', 'Agent动作', '反馈']
    xs_t, ws_t = [0.6, 1.1, 3.0, 4.7], [0.5, 1.9, 1.7, 2.75]
    rows = [('1', '订单簿快照', '挂单买入', '部分成交'), ('2', '新的行情', '撤单、改价', '全部成交'),
            ('3', '价格回落', '卖出平仓', '记录盈亏'), ('终', '交易日结束', '—', '得分：盈亏扣回撤 × 夏普')]
    y0, rh = 1.96, 0.56
    sh.rule(0.6, y0, 6.85, C['ink'])
    for x, w, h in zip(xs_t, ws_t, heads):
        sh.t(x, y0 + 0.04, w, 0.3, h, 9.5, C['grey'], MONO, anchor='ctr')
    for i, r in enumerate(rows):
        y = y0 + 0.36 + i * rh
        last = i == len(rows) - 1
        if last:
            sh.rect(0.6, y, 6.85, rh, C['tint'])
        for j, (x, w, v) in enumerate(zip(xs_t, ws_t, r)):
            sh.t(x + (0.08 if j == 0 else 0), y, w, rh, v, 12.5 if j else 11, C['accent'] if (last and j == 3) else (C['grey'] if j == 0 else C['ink']),
                 MONO if j == 0 else SANS, last and j == 3, anchor='ctr')
        sh.rule(0.6, y + rh, 6.85)
    kinds = [('Agent轨迹', '每一步的动作、反馈和最终结果', '用于监督微调与强化学习'),
             ('专家数据', '真实工作的示范、判断和评分标准', '用于对齐、奖励建模与任务设计'),
             ('评测数据', '私有任务与隔离测试集', '用于能力验收与持续评测')]
    for i, (k, d, use) in enumerate(kinds):
        y = 1.66 + i * 1.0
        sh.rule(7.85, y, 4.88, C['accent'] if i == 0 else C['ink'])
        sh.t(7.85, y + 0.08, 4.88, 0.36, k, 17, C['ink'], SERIF)
        sh.t(7.85, y + 0.46, 4.88, 0.26, d, 11, C['body'])
        sh.t(7.85, y + 0.7, 4.88, 0.26, use, 10.5, C['accent'], SANS, True)
    kicker(sh, 5.15, [('同一份记录，', F), ('既能训练，也能复测', T)], 20)
    sh.t(0.6, 5.7, W, 0.5, '每轮模型升级都要新任务、新轨迹和新的独立测试样本：数据需求随模型迭代持续发生。', 12, C['grey'])
    footer(sh)


def p_results(sh):
    header(sh, '结果', 3, [('分数拉得开，', F), ('训练有提升', T)])
    label(sh, 0.6, 1.66, 7, 'Xitadel公开预览版得分（满分100）')
    bars = [('GPT 6', 77.28), ('GLM 5.3', 30.12), ('Kimi K3', 27.68), ('DeepSeek V4 Pro', 24.58)]
    x0, scale, y0, bh, gap = 2.6, 4.2 / 100, 2.1, 0.42, 0.26
    for i, (k, v) in enumerate(bars):
        y = y0 + i * (bh + gap)
        sh.t(0.6, y, 1.95, bh, k, 13, C['ink'], SANS, anchor='ctr')
        sh.rect(x0, y, v * scale, bh, C['ink'])
        if v > 60:
            sh.t(x0 + v * scale - 1.0, y, 0.9, bh, f'{v:.2f}', 12, C['onDarkHi'], MONO, algn='r', anchor='ctr')
        else:
            sh.t(x0 + v * scale + 0.1, y, 0.9, bh, f'{v:.2f}', 12, C['ink'], MONO, anchor='ctr')
    hx = x0 + 80 * scale
    sh.rect(hx, y0 - 0.2, 0.02, 4 * (bh + gap) + 0.1, C['accent'])
    sh.t(hx - 0.8, y0 - 0.46, 1.8, 0.24, '人类最佳 80', 10, C['accent'], MONO, True, algn='ctr')
    sh.t(0.6, 4.95, 6.6, 0.5, ['逐笔撮合；盈亏扣回撤再乘夏普', '最后一个交易日留出，结果哈希提前公开'], 11, C['body'], line=1.1, gap=2)
    rx = 7.85
    sh.rect(rx, 1.66, 4.88, 3.78, C['ink'])
    sh.t(rx + 0.35, 1.86, 4.2, 0.22, '训练', 9.5, C['accentLt'], MONO, True)
    sh.t(rx + 0.35, 2.12, 4.2, 1.0, '+12%', 60, C['accentLt'], SERIF)
    sh.t(rx + 0.35, 3.2, 4.2, 0.9, ['Qwen3.8-27B在Xitadel训练后', '未见过的交易日上最高提升12%', '多次独立复现'], 12.5, C['onDarkHi'], line=1.15, gap=2)
    sh.rect(rx + 0.35, 4.3, 4.18, 0.01, DARK_RULE)
    sh.t(rx + 0.35, 4.4, 4.2, 0.9, ['单轮耗时降到1/4', '同等资源已评分尝试+64%'], 11.5, C['onDark'], line=1.1, gap=2)
    note(sh, '来源：Xitadel-QuantBench公开报告与计分规范（IMC Prosperity 3、4订单簿）；+12%为最佳一次，非均值；训练基础设施指标为公司内部测试。', 6.3)
    footer(sh)


def p_market(sh):
    header(sh, '市场', 4, [('今天85亿美元，', F), ('2030年公司情景测算7,000亿美元', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, '今天：训练数据与RL环境')
    sh.t(0.9, 2.25, 5, 0.72, '85亿美元 / 年', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, '50余家供应商收入估算（Deedy Das，2026年7月）', 11, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030年：公司情景测算', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '7,000亿美元 / 年', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, '约为今天的82倍；假设AI经济7万亿美元，训练投入占10%', 11, C['body'])
    growth = [('27倍', 'Mercor年化毛营收：16个月，7,500万→20亿美元', '专家数据需求在涨'),
              ('10倍', 'Mercor估值：17个月，20亿→200亿美元（洽谈中）', '资本在加注'),
              ('18倍', 'Snorkel AI新数据服务近一年增长，年化3.75亿美元', '需求转向专家与环境')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 11.5, C['body'], line=1.1)
        means(sh, x, 5.34, cw, m, 12)
    note(sh, '来源：Mercor（TechCrunch、Sacra、Dealroom）；Snorkel AI公告（2026年9月）；Deedy Das行业图谱（2026年7月）。2030年为公司情景测算，非独立研究预测。', 6.3)
    footer(sh)


def p_customers(sh):
    header(sh, '客户', 4, [('先卖给AI实验室，', F), ('再到企业和个人Agent团队', T)])
    cards = [('AI实验室', '现在', '2家前沿实验室在谈', ['采购专业数据与训练环境，用于模型训练和能力验证'],
              '环境授权、Agent轨迹、专家数据与评测'),
             ('企业Agent团队', '拓展', '', ['把业务流程转成可执行任务，在部署前训练、测试与验收'],
              '定制任务、评分工具与验收评测'),
             ('个人Agent开发团队', '下一步', '', [(R('Meta Muse：', 11, C['ink'], SANS, True), R('独立虚拟机与浏览器，跨应用执行任务，持续推进用户的长期目标', 11, C['body'])),
                                               (R('OpenAI Dots：', 11, C['ink'], SANS, True), R('常驻云端，调用工具、分派任务，用户离线后继续工作', 11, C['body']))],
              '长任务、跨应用与失败恢复的练习环境')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.75
    for (k, tag, extra, need, deliver), x in zip(cards, xs):
        dark = tag == '现在'
        main, sub, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(x + 0.3, top + 0.22, cw - 0.6, 0.22, tag, 9.5, acc, MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 20, main, SERIF)
        if extra:
            sh.t(x + 0.3, top + 0.98, cw - 0.6, 0.26, extra, 11, acc, SANS, True)
        sh.t(x + 0.3, top + 1.3, cw - 0.6, 0.2, '他们要什么', 9, sub, MONO)
        sh.t(x + 0.3, top + 1.54, cw - 0.6, 1.2, need, 11, sub if not dark else C['onDarkHi'], line=1.15, gap=4)
        sh.rect(x + 0.3, top + 2.72, cw - 0.6, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(x + 0.3, top + 2.84, cw - 0.6, 0.2, '我们交付', 9, acc, MONO)
        sh.t(x + 0.3, top + 3.08, cw - 0.6, 0.6, deliver, 12, main, SANS, True, line=1.1)
    note(sh, '个人Agent产品示例：Meta（2026年9月8日）、OpenAI（2026年9月29日）官方发布；二者为行业示例，不是衍真客户。', 5.8)
    footer(sh)


def p_competition(sh):
    header(sh, '竞争', 4, [('对手强在规模，', F), ('我们强在专业领域的计分和速度', T)])
    xs, ws = [0.6, 3.75, 6.55, 9.45], [3.0, 2.65, 2.75, 3.28]
    heads = ['客户已有的选择', '他们的强项', '他们的短板', '衍真的优势']
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 0.98, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    rows = [('数据与环境供应商', 'Scale AI、AfterQuery、Mercor', ['专家网络大，实验室关系深'],
             ['以通用领域为主；', '交易的撮合与风险计分', '需要交易台经验'],
             ['交易台出身设计撮合与风险计分，', '同一策略永远同一分']),
            ('公开基准与评测', 'Alpha Arena等榜单', ['方便比较模型能力'], ['只做评测，不产出训练数据；', '公开后很快饱和'],
             ['同一环境产出训练数据；', '储备任务持续换题']),
            ('客户内部自建', '', ['贴合自身工具和业务'], ['缺专业规则与专家，占研究人力'],
             ['现成的交易、事件预测环境', '与专家网络，按需交付'])]
    y0, rh = 2.04, 0.98
    for i, (k, names, pro, con, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.85, rh, [para([R(k, 15, C['ink'], SERIF)])] + ([para([R(names, 8.5, C['grey'], MONO)])] if names else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.15, rh, pro, 11.5, C['body'], anchor='ctr', line=1.1, gap=0)
        sh.t(xs[2], y, ws[2] - 0.15, rh, con, 11.5, C['body'], anchor='ctr', line=1.1, gap=0)
        sh.t(xs[3], y, ws[3] - 0.1, rh, us, 12, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=0)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.t(0.6, 5.15, 1.6, 0.4, '可复用资产', 13, C['accent'], SERIF, anchor='ctr')
    sh.t(2.2, 5.15, W - 1.6, 0.4, '任务库 · 评分器 · 隔离测试集 · 专家数据 · 失败案例：随客户和模型迭代持续积累', 12, C['ink'], anchor='ctr')
    kicker(sh, 5.75, [('可比交易：', F), ('据彭博，阿里拟领投UniPat 3亿美元，估值约25亿美元', T)], 17)
    footer(sh)


def p_business(sh):
    header(sh, '商业模式', 5, [('卖环境授权、数据和联合训练，', F), ('模型每升级一次就要换新', T)])
    lines = [('训练环境', '可执行任务与结果评分工具', '按期限、范围与使用权授权'),
             ('AI数据', '行动轨迹、专家示范与私有评测数据', '按交付量或批次收费'),
             ('联合训练', '训练实验、迁移验证与托管运行', '项目或持续服务合同')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 2.25
    for i, ((k, what, fee), x) in enumerate(zip(lines, xs)):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 22, C['ink'], SERIF)
        label(sh, x + 0.3, top + 1.08, cw - 0.6, '交付内容')
        sh.t(x + 0.3, top + 1.3, cw - 0.6, 0.3, what, 12, C['body'])
        label(sh, x + 0.3, top + 1.66, cw - 0.6, '收费方式', C['accent'])
        sh.t(x + 0.3, top + 1.88, cw - 0.6, 0.3, fee, 12, C['ink'], SANS, True)
    label(sh, 0.6, 4.22, 6, '合作路径')
    path = ['需求与验收', '付费试点', '授权交付', '持续更新']
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(path, pxs)):
        last = i == 3
        sh.rect(x, 4.5, pw, 0.62, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 4.5, pw - 0.5, 0.62, k, 16, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 4.5, 0.4, 0.62, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.55, [('每轮模型升级都要新任务和新测试样本，', F), ('授权之后持续更新，带来复购', T)], 18)
    sh.t(0.6, 6.25, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_progress, 17, {607}, 4),
    (p_backers, 11, {401, 402, 403, 404, 407}, 5),
    (p_problem, 4, {160}, 6),
    (p_why_now, 3, {126}, 7),
    (p_env, 5, {169}, 8),
    (p_products, 8, {300}, 9),
    (p_data, 9, {336}, 10),
    (p_results, 15, {562}, 11),
    (p_customers, 20, {666}, 12),
    (p_market, 13, {470}, 13),
    (p_competition, 14, {511}, 14),
    (p_business, 10, {376}, 15),
]
SCHOOLS = list(range(429, 441))
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')
V52 = os.path.join(ASSETS, 'v52')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S52(5000)
        if build is p_backers:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            build(sh, logos, poly, copy_pics(slides[12], s, SCHOOLS))
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_products:
            cw = (W - 2 * 0.3) / 3
            for i, img in enumerate(PRODUCT_IMAGES):
                s.shapes.add_picture(os.path.join(V52, img), Inches(0.6 + i * (cw + 0.3)), Inches(1.72), width=Inches(cw))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
