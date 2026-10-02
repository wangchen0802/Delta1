"""Build the English SimReal seed deck (13 slides) from v17's slides, in v41's design system.

Content follows the YC / Sequoia / a16z order (problem, why now, solution, product, traction, market,
competition, business model, why us, team, next six months) and carries no round terms. The words live in
bp_v41_en_content.json (drafted from v41's facts, fact-checked against bp_v41.py); this file is layout only.

Usage: python3 bp_v41_en.py SimReal-BP-v17.pptx SimReal-BP-v41-EN.pptx
"""
import copy
import io
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

import bp_v17
import bp_v34
from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, DARK_RULE, NS, S, W, recolor

_run = bp_v17.run


def en_run(text, sz, color, font=SANS, b=False):
    return _run(text, sz, color, font, b).replace('lang="zh-CN"', 'lang="en-US"')


bp_v17.run = bp_v34.run = en_run                   # every run on this deck is US English
R = bp_v34.R

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
X = json.load(open(os.path.join(HERE, 'bp_v41_en_content.json'), encoding='utf-8'))
SECTIONS = ['Problem', 'Solution', 'Traction', 'Market', 'Business', 'Team', 'Next']
CONF = 'Confidential  ·  For invited investors only  ·  October 2026'
CONTACT = [('business@simreal.co', MONO), ('   ·   ', SANS), ('simreal.com.cn', MONO)]
CUR = {'n': None}


# ------------------------------------------------------------------ frame ---
def runs(parts, sz, font=SERIF):
    return [R(t, sz, C['accent'] if acc else C['ink'], font) for t, acc in parts]


def claim(c):
    return [(c['lead'] + (' ' if c['lead'] and not c['lead'].endswith(' ') else ''), False), (c['accent'], True)]


def header(sh, page, section):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, page['label'], 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(runs(claim(page['claim']), 16))])


def footer(sh):
    sh.t(1.4, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, c, sz=19):
    sh.text(0.6, y, W, 0.46, [para(runs(claim(c), sz))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False, h=0.5):
    sh.t(x, y, w, h, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                       R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


def band(sh, y, h, head, text, accent, hw=2.0):
    sh.rect(0.6, y, W, h, C['ink'])
    sh.t(0.85, y, hw, h, head, 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(0.85 + hw + 0.1, y, W - hw - 0.6, h, [(R(text + ' ', 12.5, C['onDark']), R(accent, 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr', line=1.1)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


# ------------------------------------------------------------------ pages ---
def p_cover(sh):
    c = X['cover']
    sh.t(0.6, 1.35, 8.8, 1.9, [c['line1'], c['line2']], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, c['tagline'], 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, c['subline'], 16, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    for (k, v), x in zip([('Seed deck', 'October 2026'), ('Contact', 'business@simreal.co')], [0.6, 4.6]):
        sh.t(x, 5.12, 3.4, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.6, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_summary(sh):
    c = X['summary']
    header(sh, c, -1)
    cw, xs = cols(3)
    for i, (cell, x) in enumerate(zip(c['cells'], xs * 2)):
        acc = i in (2, 5)
        y = 1.72 + (i // 3) * 2.0
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, cell['k'], C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, cell['v'], 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.06, cw, 0.85, cell['lines'], 11, C['body'], line=1.12, gap=2)
    kicker(sh, 5.85, c['kicker'], 18)
    footer(sh)


def p_problem(sh):
    c = X['problem']
    header(sh, c, 0)
    sh.t(0.6, 1.64, W, 0.3, c['sub'], 12, C['grey'])
    y0, rh = 2.38, 0.62
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * 4, C['tint'])
    sh.t(2.4, y0 - 0.32, 4.4, 0.26, c['col_left'], 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 5, 0.26, c['col_right'], 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.t(0.6, y, 1.75, rh, r['domain'], 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, r['looks'], 16, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.05, rh, r['reality'], 16, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    cw, xs = cols(2, 0.5)
    for ev, x in zip(c['evidence'], xs):
        sh.t(x, 4.95, 1.7, 0.5, ev['v'], 22, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.75, 4.95, cw - 1.75, 0.5, ev['text'], 11.5, C['body'], anchor='ctr', line=1.1)
    kicker(sh, 5.75, c['kicker'], 20)
    footer(sh)


def p_why_now(sh):
    c = X['why_now']
    header(sh, c, 0)
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for i, (card, x) in enumerate(zip(c['cards'], xs)):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, card['title'], 20, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.06, cw - 0.6, 1.3, card['evidence'], 11.5, C['body'], line=1.12, gap=6)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, card['means'], 12, h=0.7)
    w = c['window']
    band(sh, 5.28, 0.72, w['head'], w['text'], w['accent'], 1.6)
    note(sh, c['sources'], 6.3)
    footer(sh)


def p_solution(sh):
    c = X['solution']
    header(sh, c, 1)
    cw, xs = cols(4, 0.4)
    top, bh = 1.72, 1.35
    for i, (st, x) in enumerate(zip(c['steps'], xs)):
        acc = i == 2
        sh.rect(x, top, cw, bh, C['accent'] if acc else C['tint'])
        sh.t(x + 0.25, top + 0.14, 1, 0.2, f'0{i + 1}', 9, C['onDarkHi'] if acc else C['accent'], MONO)
        sh.t(x + 0.25, top + 0.36, cw - 0.5, 0.42, st['name'], 18, C['onDarkHi'] if acc else C['ink'], SERIF)
        sh.t(x + 0.25, top + 0.84, cw - 0.5, 0.48, st['lines'], 10.5, C['onDarkHi'] if acc else C['body'], line=1.1, gap=0)
        if i < 3:
            sh.t(x + cw, top, 0.4, bh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    cx = xs[2] + cw / 2
    sh.rect(cx, top + bh, 0.012, 0.3, C['mid'])
    sh.t(cx + 0.08, top + bh + 0.02, 3.2, 0.26, c['connector'], 9, C['grey'], MONO, anchor='ctr')
    lt = 3.42
    sh.rect(0.6, lt, W, 1.25, C['tint'])
    sh.t(0.85, lt + 0.16, 6, 0.22, c['loop_title'], 9.5, C['accent'], MONO, True)
    lw, lxs = cols(4, 0.4, 0.85, W - 0.5)
    for i, (st, x) in enumerate(zip(c['loop'], lxs)):
        sh.t(x, lt + 0.46, lw, 0.34, st['name'], 14.5, C['accent'] if i == 3 else C['ink'], SERIF)
        sh.t(x, lt + 0.82, lw, 0.4, st['desc'], 10, C['body'], line=1.05)
        if i < 3:
            sh.t(x + lw, lt + 0.46, 0.4, 0.34, '→', 12, C['grey'], algn='ctr', anchor='ctr')
    sh.t(0.6, 4.78, W, 0.26, c['loop_note'], 10, C['grey'])
    sh.rect(0.6, 5.2, W, 0.66, C['ink'])
    sh.t(0.85, 5.2, 2.4, 0.66, c['deliver_head'], 14, C['accentLt'], SERIF, anchor='ctr')
    sh.t(3.2, 5.2, W - 2.8, 0.66, '  ·  '.join(c['deliver_items']), 15, C['onDarkHi'], SERIF, anchor='ctr')
    note(sh, 'References: InstructGPT (arXiv 2203.02155); DeepSeek-R1 (arXiv 2501.12948); Absolute Zero (arXiv 2505.03335).', 6.2)
    footer(sh)


def p_product(sh):
    c = X['product']
    header(sh, c, 1)
    x0, y0, rh = 6.55, 1.72, 0.54
    sh.rule(x0, y0, W + 0.6 - x0, C['ink'])
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.t(x0, y + 0.04, 6.2, 0.26, r['name'], 13.5, C['ink'], SERIF)
        sh.t(x0, y + 0.29, 6.2, 0.24, r['desc'], 10.5, C['body'])
        sh.rule(x0, y + rh, W + 0.6 - x0)
    label(sh, 0.6, 4.78, 8, c['infra_label'], C['accent'])
    cw, xs = cols(3)
    for m, x in zip(c['infra'], xs):
        sh.rect(x, 5.08, cw, 0.02, C['ink'])
        sh.t(x, 5.16, 1.4, 0.7, m['v'], 30, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.45, 5.16, cw - 1.45, 0.7, [m['k'], (R(m['how'], 9.5, C['grey']),)], 11, C['ink'], SANS, True,
             anchor='ctr', line=1.1, gap=1)
    note(sh, c['note'], 6.3)
    footer(sh)


def p_traction(sh, logos, polymarket, schools):
    c = X['traction']
    header(sh, c, 2)
    cw, xs = cols(4, 0.3)
    for i, (st, x) in enumerate(zip(c['stats'], xs)):
        acc = i == 0
        sh.rect(x, 1.72, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, 1.8, cw, 0.74, st['v'], 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.56, cw, 0.3, st['k'], 12, C['ink'], SANS, True)
        means(sh, x, 2.88, cw, st['means'], 11)
    sh.t(0.6, 3.28, W, 0.22, c['date_note'], 8.5, C['grey'])
    label(sh, 0.6, 3.66, 8, c['backers_label'])
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 4.16, 1.25 if pic.shape_id == 403 else 1.0)
    if polymarket is not None:
        place(polymarket, slots[4], 4.16)
    sh.rule(0.6, 4.62, W, C['ink'])
    for i, n in enumerate(c['network']):
        y = 4.74 + i * 0.52
        sh.t(0.6, y, 1.6, 0.52, n['v'], 24, C['accent'] if i == 0 else C['ink'], SERIF, anchor='ctr')
        sh.t(2.25, y, 2.5, 0.52, n['k'], 11, C['body'], anchor='ctr', line=1.05)
    gx = 4.95
    cell = (W + 0.6 - gx) / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, gx + (i % 6 + 0.5) * cell, 5.14 + (i // 6) * 0.68, min(1.0 / w, 0.46 / h, 1.0))
    note(sh, c['disclaimer'], 6.45, 8)
    footer(sh)


def p_market(sh):
    c = X['market']
    header(sh, c, 3)
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    for side, x, acc in [(c['now'], 0.9, False), (c['future'], 6.9, True)]:
        label(sh, x, 1.97, 5.6, side['label'], C['accent'] if acc else None)
        sh.t(x, 2.25, 5.6, 0.72, side['value'], 38, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 3.05, 5.6, 0.5, side['sub'], 11, C['body'], line=1.1)
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    cw, xs = cols(3)
    for m, x in zip(c['comps'], xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, m['v'], 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, m['d'], 11.5, C['body'], line=1.1)
        means(sh, x, 5.38, cw, m['means'], 12)
    note(sh, c['note'], 6.3)
    footer(sh)


def p_competition(sh):
    c = X['competition']
    header(sh, c, 3)
    xs, ws = [0.6, 3.75, 6.55, 9.45], [3.0, 2.65, 2.75, 3.28]
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 1.0, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, c['heads'])):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    y0, rh = 2.04, 1.0
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.85, rh, [para([R(r['alt'], 15, C['ink'], SERIF)])] + ([para([R(r['names'], 8.5, C['grey'], MONO)])] if r['names'] else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.2, rh, r['pro'], 11, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[2], y, ws[2] - 0.2, rh, r['con'], 11, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[3], y, ws[3] - 0.1, rh, r['ours'], 11.5, C['ink'], SANS, True, anchor='ctr', line=1.1)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.t(0.6, 5.2, 2.0, 0.4, c['moat']['head'], 13, C['accent'], SERIF, anchor='ctr')
    sh.t(2.6, 5.2, W - 2.0, 0.4, c['moat']['text'], 11.5, C['ink'], anchor='ctr')
    kicker(sh, 5.78, c['kicker'], 17)
    footer(sh)


def p_business(sh):
    c = X['business']
    header(sh, c, 4)
    label(sh, 0.6, 1.66, 4, 'Who pays')
    cw, xs = cols(3, 0.3)
    for i, (sg, x) in enumerate(zip(c['segments'], xs)):
        now = i == 0
        sh.rect(x, 1.92, cw, 1.02, C['ink'] if now else C['tint'])
        sh.t(x + 0.25, 2.02, cw - 0.5, 0.2, sg['stage'], 9, C['accentLt'] if now else C['accent'], MONO, True)
        sh.t(x + 0.25, 2.24, cw - 0.5, 0.34, sg['name'], 15, C['onDarkHi'] if now else C['ink'], SERIF)
        sh.t(x + 0.25, 2.58, cw - 0.5, 0.34, sg['line'], 10.5, C['onDark'] if now else C['body'], line=1.05)
    label(sh, 0.6, 3.1, 4, 'What they buy')
    for i, (ln, x) in enumerate(zip(c['lines'], xs)):
        sh.rect(x, 3.36, cw, 1.6, C['tint'])
        sh.t(x + 0.25, 3.46, 1, 0.2, f'0{i + 1}', 9, C['accent'], MONO, True)
        sh.t(x + 0.25, 3.66, cw - 0.5, 0.4, ln['name'], 18, C['ink'], SERIF)
        sh.t(x + 0.25, 4.08, cw - 0.5, 0.4, ln['what'], 10.5, C['body'], line=1.05)
        sh.t(x + 0.25, 4.56, cw - 0.5, 0.34, [(R('Pricing  ', 9, C['accent'], MONO), R(ln['pricing'], 10.5, C['ink'], SANS, True))], 10.5, line=1.05)
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(c['path'], pxs)):
        last = i == 3
        sh.rect(x, 5.14, pw, 0.5, C['ink'] if last else C['tint'])
        sh.t(x + 0.22, 5.14, pw - 0.44, 0.5, k, 14, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 5.14, 0.4, 0.5, '→', 13, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.8, c['kicker'], 17)
    note(sh, c['note'], 6.42)
    footer(sh)


def p_why_us(sh):
    c = X['why_us']
    header(sh, c, 5)
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.35
    for card, x in zip(c['cards'], xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, cw - 0.6, 0.46, card['title'], 20, C['ink'], SERIF)
        sh.t(x + 0.3, top + 0.82, cw - 0.6, 0.2, 'How', 9, C['grey'], MONO)
        sh.t(x + 0.3, top + 1.05, cw - 0.6, 1.0, card['how'], 11, C['body'], line=1.1, gap=2)
        sh.rule(x + 0.3, top + 2.08, cw - 0.6, C['mid'])
        sh.t(x + 0.3, top + 2.2, cw - 0.6, 0.2, 'Result', 9, C['accent'], MONO)
        sh.t(x + 0.3, top + 2.44, cw - 0.6, 0.85, card['result'], 12, C['ink'], SANS, True, line=1.1, gap=2)
    m = c['moat']
    band(sh, 5.32, 0.72, m['head'], m['text'], m['accent'], 2.2)
    note(sh, c['note'], 6.3)
    footer(sh)


def p_team(sh):
    c = X['team']
    header(sh, c, 5)
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.62
    for p, x in zip(c['people'], xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(p['name'], 24, C['ink'], SERIF), R('   ' + p['role'], 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.74, cw - 0.6, 2.2, p['bullets'], 10, C['body'], line=1.08, gap=3)
        sh.rule(x + 0.3, y0 + 2.92, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.0, cw - 0.6, 0.58, p['edu'], 9, C['grey'], line=1.1, gap=1)
    kicker(sh, 6.16, c['kicker'], 17)
    footer(sh)


def p_next(sh):
    c = X['next']
    header(sh, c, 6)
    cw, xs = cols(2, 0.3)
    top, ch = 1.72, 3.2
    for i, (t, x) in enumerate(zip(c['tracks'], xs)):
        dark = i == 1
        main, sub, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.32, cw - 0.64
        sh.t(ix, top + 0.2, iw, 0.22, t['tag'], 9.5, acc, MONO, True)
        sh.t(ix, top + 0.44, iw, 0.46, t['title'], 24, acc, SERIF)
        sh.t(ix, top + 0.96, iw, 0.28, t['line'], 12, main)
        sh.t(ix, top + 1.32, iw, 0.2, 'Focus', 9, sub, MONO)
        sh.t(ix, top + 1.54, iw, 0.32, t['focus'], 11, sub)
        sh.rect(ix, top + 1.95, iw, 0.01, DARK_RULE if dark else C['rule'])
        mw = iw / 2
        for j, m in enumerate(t['milestones']):
            mx = ix + j * mw
            sh.t(mx, top + 2.07, mw - 0.15, 0.2, m['when'], 9, acc, MONO)
            sh.t(mx, top + 2.3, mw - 0.15, 0.4, m['what'], 16, main, SERIF)
            sh.t(mx, top + 2.72, mw - 0.15, 0.45, m['detail'], 10.5, sub, line=1.1)
    if c.get('extra'):
        sh.t(xs[1] + 0.32, top + ch - 0.4, cw - 0.64, 0.3, c['extra'], 11, C['onDarkHi'], SANS, True)
    v = c['vision']
    sh.t(0.6, 5.12, W, 0.9, [v['line1'], v['line2']], 22, C['ink'], SERIF, line=1.05, gap=0)
    sh.t(0.6, 6.25, W, 0.3, [tuple(R(t, 11, C['grey'] if f is SANS else C['ink'], f) for t, f in CONTACT)], 11)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_summary, 1, {64}, 2),
    (p_problem, 4, {160}, 3),
    (p_why_now, 3, {126}, 4),
    (p_solution, 6, {237}, 5),
    (p_product, 8, {300}, 6),
    (p_traction, 11, {401, 402, 403, 404, 407}, 7),
    (p_market, 13, {470}, 8),
    (p_competition, 14, {511}, 9),
    (p_business, 10, {376}, 10),
    (p_why_us, 15, {562}, 11),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 12),
    (p_next, 16, {607}, 13),
]
SCHOOLS = list(range(429, 441))
SCREENSHOT = os.path.join(ASSETS, 'repos-2026-10-02.png')
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        sh = S(5000)
        if build is p_traction:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3))
            build(sh, logos, poly, copy_pics(slides[12], s, SCHOOLS))
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                place(p, 0.6 + (i + 0.5) * W / len(pics), 1.98)
            build(sh)
        elif build is p_product:
            s.shapes.add_picture(SCREENSHOT, Inches(0.6), Inches(1.72), width=Inches(5.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            s.shapes._spTree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
