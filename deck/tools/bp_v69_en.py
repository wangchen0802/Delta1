"""Build SimReal BP v69-EN (14 slides, English) from v17's slides.

v68-EN with the product-detail storyboards from visuals_v69_en.py: the third (scoring) panel of each strip is cleaner
(one font for every number; Xitadel's caption dropped, since the best-human row explains the scale; MLBench shows
"84% of teams beaten -> 70.1 score"; FuturePredict shows both errors and "Reward 0.20 = baseline error - AI error").
Everything else is as in v68-EN (the RSI picture is still deck/assets/v68en/rsi.png).

Usage: python3 visuals_v68_en_rsi.py; python3 visuals_v69_en.py; python3 bp_v69_en.py SimReal-BP-v17.pptx out.pptx
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['Company', 'Problem', 'Product', 'Traction', 'Market', 'Edge', 'Round']
CONF = 'Confidential  ·  For invited investors only  ·  October 2026'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.45, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def contact(sh, y):
    sh.t(0.6, y, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_cover(sh):
    sh.t(0.6, 1.35, 8.6, 1.9, ["Every industry's best AI", 'will come from our worlds'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.45, 8.2, 0.9, 'An AI-native neolab building RL environments, verifiers and training data for agents that improve themselves',
         18, C['accent'], SERIF, line=1.15)
    sh.rule(0.6, 4.95, 9.0)
    for (k, v), x in zip([('Founded', 'September 10, 2026'), ('Raising', '$6M seed round'), ('Investor presentation', 'October 2026')], [0.6, 3.6, 6.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, 'Overview', 0, [('An AI-native neolab for recursive self-improvement (RSI): ', F),
                               ('the environments, verifiers and data that agents learn from', T)])
    cells = [
        ('Revenue', '$7M ARR', ['Reached 24 days after founding', 'No outside capital'], T),
        ('Product', '7 environments', ['Trading, ML research, math, reasoning,', 'forecasting, software and accounting'], T),
        ('Proof', '+12%', ['A model trained in our trading environment', 'did better on trading days it had never seen'], F),
        ('Open source', '602 GitHub stars', ['Across 5 open benchmarks,', 'all shipped in our first 14 days'], F),
        ('Team', 'Quant founders, all 21', ['Cambridge, LSE and Duke mathematics', 'Jane Street, Citadel, Optiver, Millennium'], F),
        ('Market', '$8.5B a year', ['Spent today on training data', 'and RL environments'], F),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.98 + (i // 3) * 2.05
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, v, 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.8, d, 11.5, C['body'], line=1.15, gap=1)
    footer(sh)


def p_team(sh):
    header(sh, 'Team', 0, [('Three founders, all 21. ', F), ('Together they turned down $0.9M+ in first-year pay to build SimReal', T)])
    b, sz = C['body'], 11
    people = [
        ('Charles', 'CEO', ['First employee at United Stables, where $U grew from zero to $1.4B in a year and listed on Binance within a month',
                            'Led institutional partnerships with YZi Labs, SIG, DRW and Wintermute',
                            "Only intern on HSBC's HKD stablecoin project, working on HKMA compliance",
                            'Hedge fund intern, Citadel London'],
         ['LSE, Mathematics', 'USAMO qualifier']),
        ('Henry', 'CTO', ["Cambridge Mathematics, First Class; scholar of St John's College",
                          'Youngest researcher at the Cambridge AI research center, under Prof. Po-Ling Loh (IMS Fellow)',
                          'Ex-quant at Jane Street, Citadel and Optiver', 'Designed all 5 open benchmarks'],
         ['Top 30 worldwide, Cambridge maths competition', 'British Physics Olympiad Top Gold']),
        ('Amaris', 'COO', ['Data scientist, Millennium Hong Kong', "First graduate hire on Millennium's alternative-data team",
                           "Co-organized Plug and Play's first Hong Kong event, with HKSTP (200+ attendees)"],
         ['Duke, Mathematics and Statistics']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.75
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role + ' · 21', 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.8, cw - 0.6, 2.2, lines, sz, b, line=1.12, gap=6)
        sh.rule(x + 0.3, y0 + 2.98, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.08, cw - 0.6, 0.6, edu, 9.5, C['grey'], line=1.12, gap=1)
    footer(sh)


def p_problem(sh):
    header(sh, 'Problem', 1, [('Agents can reason. ', F), ('They still fail at real work', T)])
    rows = [('Trading', 'Wins on past data', 'Loses money live'), ('Software', 'All tests pass', 'Breaks in production'),
            ('Accounting', 'Books look closed', "Month-end doesn't reconcile"), ('Forecasting', 'Analysis sounds right', 'Outcome proves it wrong')]
    y0, rh = 2.15, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, 'Looks right', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, 'What actually happens', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10.5, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 19, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 19, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    sh.t(0.6, 5.0, 1.8, 0.6, '6/32', 30, C['accent'], SERIF, anchor='ctr')
    sh.t(2.5, 5.0, 9.5, 0.6, 'In a live trading contest (Alpha Arena), frontier models made money in only 6 of 32 runs', 13, C['body'], anchor='ctr')
    kicker(sh, 5.85, [('The gap: ', F), ('agents have nowhere realistic to practice, and no source of truth to learn from', T)], 19)
    footer(sh)


RSI_IMG = os.path.join(ASSETS, 'v68en', 'rsi.png')


def p_rsi(sh):
    header(sh, 'Solution', 2, [('A loop that turns real outcomes into better agents, ', F), ('with AI inside it from day one', T)])
    rx, rw = 8.35, 4.38
    label(sh, rx, 1.86, rw, "What's different", C['accent'])
    sh.rule(rx, 2.14, rw, C['mid'])
    sh.t(rx, 2.26, rw, 0.36, 'Today', 17, C['ink'], SERIF)
    sh.t(rx, 2.64, rw, 0.42, 'Humans hand-write every task, label and rubric', 11.5, C['body'], line=1.15)
    sh.rule(rx, 3.12, rw, C['mid'])
    sh.t(rx, 3.24, rw, 0.36, 'SimReal', 17, C['accent'], SERIF)
    sh.t(rx, 3.62, rw, 0.8, 'Agents help write new tasks and data, find failure modes and improve the environments they '
                             'train in. Next, they will write the verifiers too', 11.5, C['body'], line=1.15)
    sh.rule(rx, 4.38, rw, C['mid'])
    sh.rect(rx, 4.52, rw, 1.18, C['ink'])
    sh.t(rx + 0.3, 4.62, 3.9, 0.22, 'Proof point', 9.5, C['accentLt'], MONO, True)
    sh.t(rx + 0.3, 4.86, 1.6, 0.7, '+12%', 32, C['accentLt'], SERIF, anchor='ctr')
    sh.t(rx + 1.75, 4.86, 2.5, 0.7, ['on trading days it never saw,', 'after training in our loop'], 11.5, C['onDarkHi'], line=1.1, gap=0, anchor='ctr')
    note(sh, '+12% is the best of several independent replications, measured on trading days held out from training.', 6.36)
    footer(sh)


def p_products(sh):
    header(sh, 'Product', 2, [('Seven environments in 14 days. ', F), ("Each is scored by something the agent can't argue with", T)])
    sh.t(0.6, 1.62, 3.3, 1.0, '602', 64, C['accent'], SERIF, anchor='ctr')
    sh.t(0.6, 2.66, 3.3, 0.6, ['GitHub stars across', '5 open benchmarks'], 14, C['ink'], SANS, True, line=1.15, gap=0)
    sh.t(0.6, 3.3, 3.3, 0.3, 'As of October 4, 2026', 10, C['grey'], MONO)
    sh.rule(0.6, 3.82, 3.3, C['mid'])
    sh.t(0.6, 3.96, 3.3, 1.9, ['Give the agent a world to practice in.',
                               'Make the reward come from a market, a proof, hidden labels, an executable system or a real-world outcome.'],
         12.5, C['body'], line=1.2, gap=8)
    x0 = 4.3
    ws = [1.45, 2.25, 3.85, 0.88]
    xs = [x0, x0 + ws[0], x0 + ws[0] + ws[1], x0 + ws[0] + ws[1] + ws[2]]
    heads = ['Domain', 'Environment', 'Scored by', 'Stars']
    rows = [('Trading', 'Xitadel', 'Profit on unseen market days, vs. the best human', '102'),
            ('ML research', 'SimReal-MLBench', 'The official Kaggle leaderboard, on hidden labels', '302'),
            ('Math', 'MathmoBench', 'Whether the proof is correct', '101'),
            ('Reasoning', 'Puzzle Benchmark', 'Hidden answers to 749 human-written puzzles', '70'),
            ('Forecasting', 'FuturePredict', 'The real outcome, after the deadline', '27'),
            ('Software', 'SWE-Forward', 'Hidden tests the code must pass', 'Private'),
            ('Accounting', 'Month-End Close', 'Whether the books reconcile', 'Private')]
    sh.rule(x0, 1.72, W + 0.6 - x0, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x, 1.74, w - 0.1, 0.3, h, 9.5, C['grey'], MONO, algn='r' if i == 3 else 'l', anchor='ctr')
    y0, rh = 2.06, 0.56
    for i, (dom, env, rew, st) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(x0, y, W + 0.6 - x0)
        sh.t(xs[0], y, ws[0] - 0.1, rh, dom, 10.5, C['grey'], MONO, anchor='ctr')
        sh.t(xs[1], y, ws[1] - 0.1, rh, env, 13, C['ink'], SERIF, anchor='ctr')
        sh.t(xs[2], y, ws[2] - 0.15, rh, rew, 11.5, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[3], y, ws[3], rh, st, 11 if st[0].isdigit() else 9.5, C['accent'] if st[0].isdigit() else C['grey'],
             MONO, st[0].isdigit(), algn='r', anchor='ctr')
    sh.rule(x0, y0 + len(rows) * rh, W + 0.6 - x0)
    note(sh, 'SWE-Forward and Month-End Close are not open-sourced.', 6.36)
    footer(sh)


def p_traction(sh):
    header(sh, 'Traction', 3, [('Founded September 10, 2026. ', F), ('$7M ARR in 24 days, with no outside capital', T)])
    ms = [('Day 14', '7', 'Environments shipped', 'Trading training loop working'),
          ('Week 3', '10+', 'Types of data delivered', ''),
          ('Day 24', '$7M', 'ARR', ''),
          ('Day 24', '602', 'GitHub stars', 'Across 5 open benchmarks')]
    cw, xs = cols(4, 0.3)
    ax = 2.08
    sh.rect(0.6, ax, W, 0.02, C['ink'])
    for i, ((day, v, k, d), x) in enumerate(zip(ms, xs)):
        acc = i >= 2
        sh.rect(x, ax - 0.07, 0.16, 0.16, C['accent'])
        sh.t(x + 0.24, ax - 0.4, cw - 0.3, 0.26, day, 10.5, C['accent'], MONO, True, anchor='ctr')
        sh.t(x, ax + 0.22, cw, 0.7, v, 40, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(x, ax + 0.96, cw, 0.3, k, 13, C['ink'], SANS, True)
        if d:
            sh.t(x, ax + 1.28, cw, 0.3, d, 10.5, C['body'])
    sh.t(0.6, 3.72, W, 0.24, 'Days counted from founding; as of October 4, 2026', 8.5, C['grey'])
    sh.rule(0.6, 4.12, W, C['ink'])
    label(sh, 0.6, 4.24, 6, 'Customers')
    sh.t(0.6, 4.5, 2.6, 0.7, '2', 34, C['ink'], SERIF, anchor='ctr')
    sh.t(0.6, 5.22, 2.7, 0.3, 'Frontier labs in talks', 13, C['ink'], SANS, True)
    label(sh, 3.75, 4.24, 8, 'Training infrastructure: more training per unit of compute', C['accent'])
    infra = [('1/4', 'Each practice run takes a quarter of the time'), ('+64%', 'More scored attempts on the same compute'),
             ('−1/3', 'Less time spent saving model checkpoints')]
    iw, ixs = cols(3, 0.3, 3.75, W + 0.6 - 3.75)
    for (v, k), x in zip(infra, ixs):
        sh.rect(x, 4.5, iw, 0.02, C['ink'])
        sh.t(x, 4.56, iw, 0.66, v, 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x, 5.22, iw, 0.5, k, 11.5, C['body'], line=1.1)
    note(sh, 'Infrastructure figures come from separate internal tests and should not be added together.', 6.3)
    footer(sh)


def p_why_now(sh):
    header(sh, 'Why now', 4, [('Agents are starting to improve themselves. ', F), ('Environments are now the bottleneck', T)])
    xl, wl, xr = 0.6, 6.2, 7.45
    wr = W + 0.6 - xr
    sh.rect(xr - 0.25, 1.72, wr + 0.25, 0.34 + 3 * 1.08, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    sh.t(xl + 0.2, 1.74, 4, 0.32, 'What changed in 2026', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(xr, 1.74, 4, 0.32, 'What we have built', 9.5, C['accent'], MONO, True, anchor='ctr')
    rows = [('AI is doing AI research', 'OpenAI: agents do 3.1 days of work for every researcher-day', 'OpenAI, Sep 2026',
             ['SimReal-MLBench: an exam for AI research', '302 GitHub stars']),
            ('AI is improving itself', 'An agent rewrote its own code and improved 7 times in 8 days', 'Weco AIDE², Sep 2026',
             ['RSI working in trading', 'Up to +12% on unseen trading days']),
            ('Environments are the bottleneck', 'When the environment gets noisy, top agents drop from 83.9% to 57.6%', 'Breaking the Environment Wall, Sep 2026',
             ['Seven environments in 14 days', '200K+ experts to build more'])]
    y0, rh = 2.06, 1.08
    for i, (k, ev, src, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.t(xl + 0.2, y + 0.14, wl, 0.42, k, 20, C['ink'], SERIF)
        sh.t(xl + 0.2, y + 0.56, wl, 0.26, ev, 12, C['body'])
        sh.t(xl + 0.2, y + 0.82, wl, 0.2, src, 8.5, C['grey'], MONO)
        sh.t(xr - 0.7, y, 0.4, rh, '→', 16, C['accent'], algn='ctr', anchor='ctr')
        sh.t(xr, y, wr - 0.1, rh, [(R(us[0], 15, C['ink'], SANS, True),), (R(us[1], 12, C['body']),)], 15, anchor='ctr', line=1.1, gap=3)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.rect(0.6, 5.62, W, 0.68, C['ink'])
    sh.t(0.85, 5.62, 1.7, 0.68, 'Our position', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.6, 5.62, W - 2.2, 0.68, [(R('Environments gate RSI; they are our core research. ', 12.5, C['onDark']),
                                      R('Most neolabs have no revenue; we hit $7M ARR in 24 days', 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr')
    note(sh, 'Sources: OpenAI (Sep 6, 2026); Weco AI, arXiv 2609.26457; Breaking the Environment Wall, arXiv 2609.29773; '
             'Deedy Das neolab list (May 2026). +12% is the best of several replications.', 6.45)
    footer(sh)


def p_market(sh):
    header(sh, 'Market', 4, [('$8.5B a year today. ', F), ('One vendor, Mercor, grew revenue 27x in 16 months', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, 'Today: training data and RL environments')
    sh.t(0.9, 2.25, 5, 0.72, '$8.5B / year', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, 'Combined revenue of 50+ vendors', 11.5, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030: the AI economy', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '$700B / year', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, 'If AI becomes a $7T economy and 10% goes to training', 11.5, C['body'])
    growth = [('27x', ['Mercor run-rate revenue:', '$75M to $2B in 16 months'], 'Demand for expert data is exploding'),
              ('10x', ['Mercor valuation:', '$2B to $20B in 17 months (in talks)'], 'Investors are paying up'),
              ('18x', ["Snorkel AI's new data business:", 'growth in one year'], 'Demand is shifting to experts and environments')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 12, C['body'], line=1.1, gap=0)
        means(sh, x, 5.38, cw, m, 12)
    note(sh, 'Sources: Deedy Das market map (Jul 2026); Mercor (TechCrunch, Dealroom); Snorkel AI (Sep 2026). '
             'The 2030 figure is our estimate.', 6.3)
    footer(sh)


def p_customers(sh):
    header(sh, 'Business model', 4, [('We sell environments, data and evals. ', F),
                                     ('Frontier labs pay today; enterprise and personal agents next', T)])
    cards = [('Now', 'Frontier AI labs', 'Environments, verified tasks and evals to post-train their models',
              'Per environment and per task, on quarterly contracts', '$300K–$1M+', 'Per lab, per quarter',
              'Market rates: $200–$2,000 a task; about $300K for a complex environment'),
             ('Next', 'Enterprise agents', 'Custom environments and tests that prove an agent can do the job',
              'A paid pilot, then an annual license', '$200K–$800K', 'Per customer, per year', 'Pilot: $50K–$150K over 8–12 weeks'),
             ('Later', 'Personal agents', "Short practice runs that fit an agent to one person's habits before it acts for them",
              'A monthly subscription, direct or through agent apps', '$5–$10', 'Per user, per month',
              'For reference, consumer AI plans cost about $20 a month')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.66, 4.12
    for (tag, k, buy, charge, price, unit, detail), x in zip(cards, xs):
        dark = tag == 'Now'
        main, sub, acc, rule = ((C['onDarkHi'], C['onDark'], C['accentLt'], DARK_RULE) if dark else (C['ink'], C['body'], C['accent'], C['mid']))
        ix, iw = x + 0.3, cw - 0.6
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(ix, top + 0.2, iw, 0.22, tag, 10, acc, MONO, True)
        sh.t(ix, top + 0.44, iw, 0.5, k, 22, main, SERIF)
        sh.t(ix, top + 0.98, iw, 0.2, 'They buy', 9, sub, MONO)
        sh.t(ix, top + 1.2, iw, 0.44, buy, 12, main, line=1.15)
        sh.rect(ix, top + 1.74, iw, 0.01, rule)
        sh.t(ix, top + 1.84, iw, 0.2, 'How we charge', 9, sub, MONO)
        sh.t(ix, top + 2.06, iw, 0.44, charge, 12, main, line=1.15)
        sh.rect(ix, top + 2.6, iw, 0.01, rule)
        sh.t(ix, top + 2.7, iw, 0.2, 'Estimated price', 9, acc, MONO)
        sh.t(ix, top + 2.9, iw, 0.44, price, 24, acc, SERIF)
        sh.t(ix, top + 3.34, iw, 0.22, unit, 11, main, SANS, True)
        sh.t(ix, top + 3.58, iw, 0.36, detail, 9.5, sub, line=1.1)
    kicker(sh, 5.9, [('One set of environments serves all three, ', F), ('and every new model needs new tasks', T)], 18)
    note(sh, 'Estimates, not signed prices. Lab prices are market rates from Epoch AI interviews with environment builders and AI labs '
             '(January 2026). Enterprise and personal prices are our assumptions.', 6.46)
    footer(sh)


PRODUCT_IMG = os.path.join(ASSETS, 'v69en')
DETAIL = [('Trading', 'Xitadel', 'The AI writes a trading strategy.', 'Score: risk-adjusted profit vs. the best human', 'xitadel.png'),
          ('AI research', 'SimReal-MLBench', 'The AI runs an ML project.', 'Score: rank on the real Kaggle leaderboard', 'mlbench.png'),
          ('Forecasting', 'FuturePredict', 'The AI forecasts real events.', 'Score: error once the outcome is known', 'forecast.png')]
DY0, DRH, DGAP = 1.66, 1.48, 0.06


def p_product_detail(sh):
    header(sh, 'How it works', 2, [('Three environments up close. ', F),
                                   ('Each scores the AI against something real: a human, a leaderboard, an outcome', T)])
    for i, (tag, name, does, score, _img) in enumerate(DETAIL):
        y = DY0 + i * (DRH + DGAP)
        sh.rect(0.6, y, 2.53, DRH, C['tint'])                 # left column only: the storyboard picture brings its own tint
        sh.rect(12.66, y, W + 0.6 - 12.66, DRH, C['tint'])
        sh.rect(0.6, y, 0.05, DRH, C['accent'])
        sh.t(0.85, y + 0.16, 2.2, 0.22, tag, 9.5, C['accent'], MONO, True)
        sh.t(0.85, y + 0.38, 2.2, 0.36, name, 16, C['ink'], SERIF)
        sh.t(0.85, y + 0.74, 2.15, 0.66, [(R(does, 10, C['body']),), (R(score, 10, C['ink'], SANS, True),)], 10, line=1.1, gap=3)
    note(sh, 'Xitadel: published results; the human benchmark is the best real strategy from IMC Prosperity, a global trading competition. '
             'MLBench: a baseline run, scored on the official leaderboard. FuturePredict: the numbers are illustrative.', 6.3)
    footer(sh)


def p_why_us(sh):
    header(sh, 'Why us', 5, [('Data vendors sell expert hours. ', F), ('We build assets that compound with every customer', T)])
    xs, ws = [0.6, 2.75, 6.75], [2.15, 4.0, 5.98]
    sh.rect(xs[2] - 0.15, 1.72, ws[2] + 0.15, 0.34 + 5 * 0.7, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    sh.t(xs[1], 1.74, ws[1], 0.32, 'Expert-data vendors', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(xs[2], 1.74, ws[2], 0.32, 'SimReal', 9.5, C['accent'], MONO, True, anchor='ctr')
    rows = [('What compounds', 'Expert hours. Each project starts from scratch.',
             'Proprietary environments, verifiers and maps of where models fail, plus the models we train in them. Every customer and loop adds to them.'),
            ('Cost curve', 'Rises with headcount',
             'Falls with scale: once a verifier exists, new data is generated and checked automatically'),
            ('Proof of quality', 'Spot checks and reviewer judgment',
             'Hidden tests, out-of-sample results, hash-locked published results and measured training gains (+12%)'),
            ('Velocity', 'A new domain means new recruiting',
             '7 environments in 14 days; $7M ARR in 24 days'),
            ('Network', 'Large expert pools, paid by the hour',
             'Practitioners from top trading firms and 200K+ reachable experts, whose work our verifiers check')]
    y0, rh = 2.06, 0.7
    for i, (k, them, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.t(xs[0] + 0.15, y, ws[0] - 0.2, rh, k, 15, C['ink'], SERIF, anchor='ctr')
        sh.t(xs[1], y, ws[1] - 0.3, rh, them, 11.5, C['grey'], anchor='ctr', line=1.1)
        sh.t(xs[2], y, ws[2] - 0.15, rh, us, 11.5, C['ink'], SANS, True, anchor='ctr', line=1.1)
    sh.rule(0.6, y0 + 5 * rh, W)
    sh.rect(0.6, 5.74, W, 0.6, C['ink'])
    sh.t(0.85, 5.74, 1.6, 0.6, 'In short', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.5, 5.74, W - 2.1, 0.6, [(R('They sell hours; we sell verified gains. ', 13, C['onDark']),
                                     R('Their costs rise with scale; ours fall', 13, C['onDarkHi'], SANS, True))], 13, anchor='ctr')
    footer(sh)


def p_network(sh, logos, polymarket, schools):
    header(sh, 'Network', 5, [('Experts from the most competitive entry-level jobs and from research, ', F),
                              ('the work agents are now built to do', T)])
    label(sh, 0.6, 1.66, 10, 'Practitioners in our R&D come from top trading firms and markets, including')
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 2.2, 1.3 if pic.shape_id == 403 else 1.0)
    if polymarket is not None:
        place(polymarket, slots[4], 2.2, 0.92)
    sh.rule(0.6, 2.62, W, C['ink'])
    cover = [('The most competitive entry-level jobs', 'Trading, finance, software and more', 'Data for agents that streamline everyday work'),
             ('Research', 'AI, mathematics and quantitative research', 'Data for agents that automate research')]
    cw, xs = cols(2, 0.3)
    for (k, d, m), x in zip(cover, xs):
        sh.rect(x, 2.76, cw, 1.06, C['tint'])
        sh.rect(x, 2.76, 0.05, 1.06, C['accent'])
        sh.t(x + 0.3, 2.86, cw - 0.5, 0.36, k, 17, C['ink'], SERIF)
        sh.t(x + 0.3, 3.2, cw - 0.5, 0.26, d, 11.5, C['body'])
        means(sh, x + 0.3, 3.48, cw - 0.5, m, 12)
    stats = [('200K+', 'Reachable experts', T), ('7,000+', 'Experts on the waitlist', F), ('21', 'Top universities', F)]
    cw, xs = cols(3)
    for (v, k, acc), x in zip(stats, xs):
        sh.t(x, 3.94, 1.55, 0.52, v, 28, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(x + 1.6, 3.94, cw - 1.6, 0.52, k, 12.5, C['ink'], SANS, True, anchor='ctr')
    sh.rule(0.6, 4.56, W, C['mid'])
    label(sh, 0.6, 4.64, 6, 'Students and alumni of')
    cell = W / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, 0.6 + (i % 6 + 0.5) * cell, 5.18 + (i // 6) * 0.7, min(1.2 / w, 0.5 / h, 1.5))
    note(sh, 'Logos show where practitioners work or study; no endorsement implied.', 6.45, 8)
    footer(sh)


def p_raise(sh):
    header(sh, 'The round', 6, [('Raising a $6M seed round ', F), ('to scale environments, RSI and delivery', T)])
    top, ch, lw = 1.72, 3.86, 2.9
    sh.rect(0.6, top, lw, ch, C['ink'])
    sh.t(0.9, top + 0.3, lw - 0.6, 0.22, 'Raising', 10, C['accentLt'], MONO, True)
    sh.t(0.9, top + 0.62, lw - 0.6, 1.0, '$6M', 60, C['accentLt'], SERIF)
    sh.t(0.9, top + 1.66, lw - 0.6, 0.4, 'Seed round', 18, C['onDarkHi'], SERIF)
    sh.t(0.9, top + ch - 0.62, lw - 0.6, 0.4, 'Three directions, each with a milestone  →', 11, C['onDark'], line=1.1)
    dirs = [('01', 'Environments and data', 'Take the trading playbook to more expert domains', 'Environment engineers  ·  Expert network',
             '10+ expert-domain environments'),
            ('02', 'RSI', 'Scale self-improvement to larger models and more domains', 'Compute  ·  Researchers',
             'RSI on open-weight large models'),
            ('03', 'Delivery', 'Turn early demand into lasting contracts', 'Delivery team  ·  Sales', '20 customers')]
    x0 = 0.6 + lw + 0.25
    cw, xs = cols(3, 0.25, x0, W + 0.6 - x0)
    for (n, k, d, inv, goal), x in zip(dirs, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        ix, iw = x + 0.28, cw - 0.5
        sh.t(ix, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(ix, top + 0.46, iw, 0.66, k, 20, C['ink'], SERIF, line=1.0)
        sh.t(ix, top + 1.3, iw, 0.62, d, 12, C['ink'], SANS, True, line=1.15)
        sh.rect(ix, top + 2.0, iw, 0.01, C['mid'])
        sh.t(ix, top + 2.1, iw, 0.2, 'Where the money goes', 9, C['body'], MONO)
        sh.t(ix, top + 2.32, iw, 0.42, inv, 11.5, C['ink'], line=1.1)
        sh.t(ix, top + 2.86, iw, 0.2, 'Milestone', 9, C['accent'], MONO)
        sh.t(ix, top + 3.08, iw, 0.62, goal, 15, C['accent'], SERIF, line=1.05)
    contact(sh, 5.86)
    note(sh, 'Converted from RMB at 6.67 per US dollar.', 6.3)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_problem, 4, {160}, 4),
    (p_rsi, 6, {237}, 5),
    (p_products, 8, {300}, 6),
    (p_product_detail, 9, {336}, 7),
    (p_traction, 16, {607}, 8),
    (p_why_now, 15, {562}, 9),
    (p_market, 13, {470}, 10),
    (p_customers, 20, {666}, 11),
    (p_why_us, 14, {511}, 12),
    (p_network, 11, {401, 402, 403, 404, 407}, 13),
    (p_raise, 17, {607}, 14),
]
SCHOOLS = list(range(429, 441))
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out



def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_product_detail:
            for i, (*_rest, img) in enumerate(DETAIL):
                s.shapes.add_picture(os.path.join(PRODUCT_IMG, img), Inches(3.12), Inches(DY0 + i * (DRH + DGAP)), width=Inches(9.55))
            build(sh)
        elif build is p_network:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            build(sh, logos, poly, copy_pics(slides[12], s, SCHOOLS))
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_rsi:
            s.shapes.add_picture(RSI_IMG, Inches(0.45), Inches(1.6), width=Inches(7.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
