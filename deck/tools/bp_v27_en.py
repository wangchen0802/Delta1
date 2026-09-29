"""Build the English SimReal BP v27 from v17's slides: the same 22 pages and
layout as bp_v27.py, in English, with widths and sizes adjusted where
English runs longer than Chinese.

Usage: python3 bp_v27_en.py SimReal-BP-v17.pptx SimReal-BP-v27-EN.pptx
"""
import sys

from lxml import etree
from pptx import Presentation

from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v27 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

SECTIONS = ['Team', 'Problem', 'Solution', 'Market', 'Traction', 'Raise']
CONF = 'Confidential  ·  For invited investors only  ·  September 2026'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('     ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    n = CUR['n']
    sh.t(1.5, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if n is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{n:02d}' if isinstance(n, int) else n, 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20, algn='l'):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)), algn)], 'ctr')


# ------------------------------------------------------------------ pages ---
def p_cover(sh):
    sh.t(0.6, 1.7, 8.2, 1.9, ['Real workplaces', 'for AI that actually works'], 42, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.85, 8.2, 0.36, 'Environments, AI data and self-improvement for AI labs and enterprises', 15, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    cols = [('This round', 'RMB 40M (paid in USD)'), ('Post-money', 'RMB 500M'), ('Business plan', 'September 2026')]
    for (k, v), x in zip(cols, [0.6, 3.95, 6.1]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.3, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, 'Overview', 0, [('Give AI a real place to work, and let it fail, learn and ', False), ('improve itself', True)])
    cells = [
        ('What we do', 'Environments & AI data', ['For AI labs and enterprises: agent trajectories,', 'expert and eval data; exclusive environments for RSI'], False),
        ('Done so far', '7 products in 14 days', ['Trading, event prediction, AI research, math,', 'reasoning, finance, software; 5 public repos'], False),
        ('Key result', '+12%', ['Trained in Xitadel, Qwen3.8-27B traded up to', '12% better on unseen live markets; replicated'], True),
        ('Team', 'Quant founders, born 2005', ['Math at Cambridge, LSE and Duke', 'Jane Street, Citadel, Optiver, Millennium'], False),
        ('Market', [(R('$8.5B ', 26, C['ink'], SERIF), R('→', 24, C['ink'], SANS), R(' $700B', 26, C['ink'], SERIF))],
         ['Labs buy $8.5B a year of data and RL environments', 'The 2030 training market: ~$700B, 82x today'], False),
        ('This round', 'RMB 40M', ['Paid in USD; RMB 500M post-money', 'RMB 20M each for the revenue and RSI engines'], True),
    ]
    cw, gap = (W - 2 * 0.35) / 3, 0.35
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = 0.6 + (i % 3) * (cw + gap), 1.66 + (i // 3) * 2.0
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, y + 0.14, cw, 0.22, k, 9.5, C['accent'] if acc else C['grey'], MONO)
        sh.t(x, y + 0.42, cw, 0.56, v, 24, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.72, d, 10.5, C['body'], line=1.15, gap=0)
    kicker(sh, 5.9, [('Own the best training world, own the best AI. ', False),
                     ('From trading to every industry: AI that improves itself.', True)], 16)
    footer(sh)


def p_team(sh):
    header(sh, 'Team', 0, [('From olympiads to Wall Street’s toughest desks: ', False), ('we know the gap between tests and real work', True)])
    b, sz = C['body'], 10
    people = [
        ('Charles', 'CEO', ['First employee at United Stables: helped take its U stablecoin from 0 to $1.4B in a year; listed on Binance in a month',
                            'Ran institutional relations; worked on partnerships with SIG, DRW and others',
                            'Sole intern on HSBC’s HKD stablecoin issuance; supported HKMA compliance end to end',
                            'X (Twitter) writer with millions of views', 'Hedge fund intern at Citadel, London'],
         ['Born 2005 · LSE Mathematics · SCIE', 'USAMO qualifier']),
        ('Henry', 'CTO', ['ML research at Cambridge with renowned statistician Po-Ling Loh (IMS Fellow)',
                          'Youngest undergraduate AI researcher at a Cambridge research centre',
                          'Quant roles at Jane Street, Citadel, Optiver', 'Designed five benchmarks and RL environments'],
         ['Born 2005 · Cambridge Mathematics, First · SCIE', 'Scholar · top 30 worldwide, Cambridge math contest', 'UK Physics Olympiad Super Gold']),
        ('Amaris', 'COO', ['Data scientist at Millennium Hong Kong', 'The alt-data team’s first-ever graduate hire',
                           'Helped launch Plug and Play’s first Hong Kong event (with HK Science Park); hosted J&J MedTech’s tech summit (200+ attendees each)',
                           'Published author at 17; 100K+ engagements online'],
         'Born 2005 · Duke Math & Statistics · YK Pao School'),
    ]
    cw, gap, y0, ch = (W - 2 * 0.3) / 3, 0.3, 2.42, 3.55
    for i, (name, role, lines, edu) in enumerate(people):
        x = 0.6 + i * (cw + gap)
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.76, cw - 0.6, 2.0, lines, sz, b, line=1.08, gap=4)
        sh.rule(x + 0.3, y0 + 2.84, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 2.92, cw - 0.6, 0.55, edu, 9, C['grey'], line=1.12, gap=0)
    kicker(sh, 6.12, [('Scale AI, Mercor and AfterQuery were founded by people around 20. Early investors made ', False),
                      ('17,000x, 40x, 1,800x', True)], 15)
    footer(sh)


def p_problem(sh):
    header(sh, 'Problem', 1, [('Passing checks isn’t enough. ', False), ('We make AI do professional work in the real world', True)])
    sh.t(0.6, 1.64, W, 0.3, 'Real work has no answer key', 12, C['grey'])
    rows = [('Trading', 'The backtest climbs all the way', 'Loses money on unseen live markets'),
            ('Software', 'All tests pass', 'Breaks in production'),
            ('Finance', 'The books look finished', 'Month-end won’t reconcile'),
            ('Prediction', 'The analysis sounds convincing', 'Wrong once the result is in')]
    y0, rh = 2.38, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, 'Surface check: passed', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, 'Real world: failed', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 17, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 17, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    kicker(sh, 5.62, [('To learn real work, AI needs ', False), ('a world that reacts to its every move', True), ('.', False)])
    footer(sh)


def p_insight(sh):
    header(sh, 'Core insight', 1, [('Every leap in AI came from ', False), ('a better feedback signal', True)])
    gens = [('Gen 1 · Internet data', 'Imitation', 'Models learn to copy', 'AI that can talk'),
            ('Gen 2 · Human preference', 'Alignment', 'People rate the output', 'Useful assistants'),
            ('Gen 3 · Answer keys', 'Verification', 'Checked against answers', 'Reasoning models'),
            ('Gen 4 · Real-world feedback', 'Practice', 'AI enters the real economy; the world reacts to every move', ['Autonomy and', 'self-improvement'])]
    cw, gap, y0, ch = (W - 3 * 0.25) / 4, 0.25, 1.75, 3.1
    for i, (lab, big, how, res) in enumerate(gens):
        x, dark = 0.6 + i * (cw + gap), i == 3
        sh.rect(x, y0, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(x + 0.25, y0 + 0.2, cw - 0.5, 0.22, lab, 9, C['accentLt'] if dark else C['grey'], SANS)
        sh.t(x + 0.25, y0 + 0.5, cw - 0.5, 0.6, big, 26, C['onDarkHi'] if dark else C['ink'], SERIF)
        sh.t(x + 0.25, y0 + 1.16, cw - 0.5, 0.6, how, 11, C['onDark'] if dark else C['grey'], line=1.12)
        sh.rule(x + 0.25, y0 + 1.9, cw - 0.5, DARK_RULE if dark else C['mid'])
        sh.t(x + 0.25, y0 + 2.02, cw - 0.5, 0.5, res, 14, C['accentLt'] if dark else C['ink'], SERIF, line=1.05, gap=0)
        if dark:
            sh.t(x + 0.25, y0 + 2.62, cw - 0.5, 0.26, 'Just beginning', 9.5, C['accentLt'], MONO, True)
    kicker(sh, 5.3, [('Acing tests isn’t doing the job. The next signal ', False), ('comes from the real world', True), ('.', False)], 22, 'ctr')
    footer(sh)


def p_solution(sh):
    header(sh, 'Solution', 2, [('AI does the task; ', False), ('real outcomes become the next training signal', True)])
    parts = [('01', 'Environments', 'Real work, rebuilt for practice'), ('02', 'Outcome grading', 'Scored by real results, not another AI’s opinion'),
             ('03', 'Expert feedback', 'Professional judgment and grading standards')]
    lw, step, bh, top = 6.6, 0.9, 0.65, 1.8
    for i, (n, k, d) in enumerate(parts):
        y = top + i * step
        sh.rect(0.6, y, lw, bh, C['tint'])
        sh.t(0.85, y, 0.5, bh, n, 10, C['accent'], MONO, anchor='ctr')
        sh.t(1.35, y, 2.2, bh, k, 17, C['ink'], SERIF, anchor='ctr')
        sh.t(3.6, y, lw - 3.2, bh, d, 11.5, C['grey'], anchor='ctr', line=1.1)
    y = top + 3 * step
    sh.rect(0.6, y, lw, bh, C['ink'])
    sh.t(0.85, y, 2.2, bh, 'SimReal', 10, C['accentLt'], MONO, anchor='ctr')
    sh.t(3.6, y, lw - 3.2, bh, 'One environment for evals, data and training', 14, C['onDarkHi'], SERIF, anchor='ctr', line=1.05)
    rx, rw = 7.75, 4.98
    steps = [('AI acts', C['tint'], C['ink']), ('The world responds', C['tint'], C['ink']), ('Graded by outcome; trajectory saved', C['tint'], C['ink']),
             ('Next round of training: self-improvement', C['accent'], C['onDarkHi'])]
    for i, (k, fill, col) in enumerate(steps):
        y = top + i * step
        sh.rect(rx, y, rw, bh, fill)
        sh.t(rx + 0.3, y, rw - 0.6, bh, k, 15, col, SERIF, anchor='ctr')
        if i < 3:
            sh.t(rx, y + bh, rw, step - bh, '↓', 11, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.6, [('Each round starts where the last one ended: that is ', False), ('self-improvement (RSI)', True), ('.', False)])
    footer(sh)


def p_xitadel(sh):
    header(sh, 'Flagship proof · Xitadel', 2, [('In two weeks we ran ', False), ('the world’s first self-improving market-making loop', True)])
    rows = [('Real battlefield', 'Replays real trading days and order books, with data licensed through IMC Trading'),
            ('Market as judge', 'The market settles every trade'),
            ('Learns from results', 'Reviews each day’s P&L and orders; each round beats the last')]
    lw = 6.4
    sh.rule(0.6, 1.7, lw, C['ink'])
    for i, (k, d) in enumerate(rows):
        y = 1.7 + i * 0.5
        sh.t(0.6, y, 1.6, 0.5, k, 10.5, C['ink'], SANS, True, anchor='ctr')
        sh.t(2.25, y, lw - 1.65, 0.5, d, 10.5, C['body'], anchor='ctr', line=1.05)
        sh.rule(0.6, y + 0.5, lw)
    sh.t(0.6, 3.4, 2.4, 0.95, '+12%', 54, C['accent'], SERIF, anchor='ctr')
    sh.t(3.0, 3.44, lw - 2.4, 0.9, [(R('Open-source Qwen3.8-27B, after training,', 12, C['ink']), BR(12), R('traded up to 12% better on unseen live markets', 12, C['ink'])),
                                    (R('Controlled experiment, replicated across independent runs; logs available in diligence', 9.5, C['grey']),)],
         12, anchor='ctr', line=1.1)
    sh.rule(0.6, 4.54, lw)
    sh.t(0.6, 4.66, lw, 0.34, 'Our environments let AI improve itself in real financial markets', 14, C['ink'], SERIF)
    sh.t(0.6, 5.04, lw, 0.34, 'Next: more strategies, more markets, more industries', 14, C['accent'], SERIF)
    px, pw, py, ph = 7.45, 5.28, 1.7, 3.68
    sh.rect(px, py, pw, ph, C['tint'])
    sh.t(px + 0.3, py + 0.22, pw - 0.6, 0.22, 'Xitadel public preview', 9.5, C['grey'], MONO)
    sh.t(px + 0.3, py + 0.48, pw - 0.6, 0.36, 'No model has passed the human baseline yet', 15, C['ink'], SERIF)
    models = [('GPT 6', 77.28), ('GLM 5.3', 30.12), ('Kimi K3', 27.68), ('DeepSeek V4 Pro', 24.58)]
    bx, scale, by = px + 1.65, 0.031, py + 1.45
    human_x = bx + 80 * scale
    sh.rect(human_x, by - 0.2, 0.015, len(models) * 0.46 + 0.1, C['accent'])
    sh.t(human_x - 0.8, by - 0.44, 1.6, 0.2, 'Human baseline 80', 9, C['accent'], MONO, algn='ctr')
    for i, (m, v) in enumerate(models):
        y, end = by + i * 0.46, bx + v * scale
        sh.t(px + 0.3, y, 1.35, 0.32, m, 11, C['ink'], anchor='ctr')
        sh.rect(bx, y + 0.05, v * scale, 0.22, C['ink'])
        if end + 0.66 > human_x:
            sh.t(end - 0.66, y, 0.6, 0.32, f'{v:.2f}', 10, C['onDarkHi'], MONO, algn='r', anchor='ctr')
        else:
            sh.t(end + 0.06, y, 0.7, 0.32, f'{v:.2f}', 10, C['ink'], MONO, anchor='ctr')
    footer(sh)


def p_products(sh):
    header(sh, 'Products', 2, [('7 products shipped in 14 days, ', False), ('5 of them in public repos', True)])
    rows = [('Xitadel', 'Xitadel-QuantBench', 'Trading', 'Replays real markets and settles on P&L; self-improvement proven', 'Open source', '102 stars'),
            ('SimReal-MLBench', 'Simreal-MLBench', 'AI research', '60 research tasks across 7 data types, modelled on OpenAI’s MLE-bench', 'Open source', '178 stars'),
            ('FuturePredict Bench', 'future-prediction-bench', 'Prediction', 'Predicts real events and trains on the outcome; live data', 'Partly open', ''),
            ('MathmoBench', 'MathmoBench', 'Math proofs', 'Makes AI prove its answers instead of guessing', 'Open source', '102 stars'),
            ('Puzzle Benchmark', 'hard-puzzle-benchmark', 'Logic', '749 hand-written reasoning puzzles', 'Open source', '17 stars'),
            ('Month-End Close', '', 'Finance', 'AI closes the books with zero errors', '', ''),
            ('SWE-Forward', '', 'Software', 'Tests whether AI-written code survives the next release', '', '')]
    cols = [(0.6, 'Product'), (3.55, 'Domain'), (4.85, 'What it does'), (10.4, 'Status')]
    sh.rule(0.6, 1.66, W, C['ink'])
    for x, h in cols:
        sh.t(x, 1.7, 2.5, 0.26, h, 9.5, C['grey'], MONO, anchor='ctr')
    y0, rh = 2.0, 0.53
    for i, (name, repo, dom, what, status, stars) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.6, y, 2.9, rh, [para([R(name, 14, C['ink'], SERIF)])] +
                ([para([R(repo, 8.5, C['grey'], MONO)])] if repo else []), 'ctr')
        sh.t(3.55, y, 1.25, rh, dom, 10.5, C['grey'], anchor='ctr')
        sh.t(4.85, y, 5.45, rh, what, 11, C['body'], anchor='ctr')
        if status:
            st = (R(status, 10.5, C['accent'], SANS, True),) + ((R(f'  ·  {stars}', 10.5, C['grey']),) if stars else ())
        else:
            st = (R('Live  ·  private', 10.5, C['grey']),)
        sh.t(10.4, y, 2.33, rh, [st], 10.5, anchor='ctr')
    sh.rule(0.6, y0 + len(rows) * rh, W)
    kicker(sh, 5.9, [('All built with ', False), ('zero outside funding', True), ('.', False)], 18)
    sh.t(6.6, 5.9, 6.13, 0.46, '400 stars in total  ·  September 29, 2026', 10.5, C['grey'], SANS, algn='r', anchor='ctr')
    footer(sh)


def p_data(sh):
    header(sh, 'Data business', 2, [('Environments and experts produce three kinds of AI data: ', False), ('trajectories, expert data, evals', True)])
    kinds = [('01', 'Agent trajectories', 'A full record of AI doing a multi-step task: every action, the world’s response, the result',
              'Pass or fail set by real outcomes', 'For fine-tuning and RL'),
             ('02', 'Expert data', 'From our 200,000+ expert network: demonstrations, judgments and grading standards from real work',
              'Entry-level roles to top research', 'For alignment, reward models and grading'),
             ('03', 'Eval data', 'Private eval and held-out sets, used only for testing, never for training',
              'Red-teamed before launch', 'For model acceptance and ongoing evals')]
    cw, gap, y0, ch = (W - 2 * 0.3) / 3, 0.3, 1.72, 2.6
    for i, (n, name, what, edge, use) in enumerate(kinds):
        x = 0.6 + i * (cw + gap)
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.t(x + 0.3, y0 + 0.2, 1, 0.22, n, 10, C['accent'], MONO)
        sh.t(x + 0.3, y0 + 0.44, cw - 0.6, 0.5, name, 21, C['ink'], SERIF)
        sh.t(x + 0.3, y0 + 1.0, cw - 0.6, 0.78, what, 10.5, C['body'], line=1.1)
        sh.rule(x + 0.3, y0 + 1.84, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 1.92, cw - 0.6, 0.3, edge, 11.5, C['ink'], SANS, True)
        sh.t(x + 0.3, y0 + 2.22, cw - 0.6, 0.28, use, 10, C['grey'])
    sh.t(0.6, 4.52, 6, 0.22, 'Unlike human-only data vendors', 10, C['grey'], MONO)
    diffs = [('Same source', 'Data, grading and training share one environment'),
             ('Verifiable', 'Real outcomes decide pass or fail, not human opinion'),
             ('Compounding', 'Where the model fails, the next batch of data goes')]
    dw = (W - 2 * 0.3) / 3
    for i, (k, d) in enumerate(diffs):
        x = 0.6 + i * (dw + 0.3)
        sh.rule(x, 4.8, dw, C['ink'])
        sh.t(x, 4.86, 1.35, 0.44, k, 15, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.35, 4.86, dw - 1.35, 0.44, d, 10, C['body'], anchor='ctr', line=1.05)
    kicker(sh, 5.72, [('Labs spend $8.5B a year on training data and RL environments: ', False), ('we sell both', True), ('.', False)], 18)
    footer(sh)


def p_why_now(sh):
    header(sh, 'Why now', 3, [('Mercor grew revenue 27x and valuation 10x in 18 months: ', False), ('the market is just starting', True)])
    tops = [('Revenue growth', '27x', 'Mercor gross run-rate: $75M → $2B in 16 months'),
            ('Valuation growth', '10x', 'Mercor valuation: $2B → $20B in 17 months (in talks)'),
            ('Demand growth', '14x', 'AI agent market forecast: $7.9B → ~$111B over 7 years')]
    cw, gap = (W - 2 * 0.35) / 3, 0.35
    for i, (k, v, d) in enumerate(tops):
        x = 0.6 + i * (cw + gap)
        sh.rect(x, 1.66, cw, 0.02, C['ink'])
        sh.t(x, 1.8, cw, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 2.06, cw, 0.8, v, 44, C['accent'], SERIF)
        sh.t(x, 2.9, cw, 0.5, d, 10.5, C['body'], line=1.1)
    facts = [('18x', 'Snorkel AI run-rate revenue in a year, now $375M'), ('$29B', 'Scale AI valuation; Meta bought 49% for $14.3B'),
             ('14 months', 'AfterQuery: from founding to $100M run-rate'), ('Nearly half', 'of Scale AI’s new training projects use RL environments'),
             ('5 months', 'AfterQuery valuation: $300M → $3.2B'), ('Before 2032', 'Public human text is projected to run out')]
    fw, fg, fy, fh = (W - 0.5) / 2, 0.5, 3.78, 0.56
    for i, (v, d) in enumerate(facts):
        x, y = 0.6 + (i % 2) * (fw + fg), fy + (i // 2) * fh
        sh.rule(x, y, fw)
        sh.t(x, y, 1.6, fh, v, 18, C['ink'], SERIF, anchor='ctr')
        sh.t(x + 1.65, y, fw - 1.65, fh, d, 11, C['body'], anchor='ctr')
    sh.rule(0.6, fy + 3 * fh, fw)
    sh.rule(0.6 + fw + fg, fy + 3 * fh, fw)
    sh.t(0.6, 5.76, W, 0.5, 'Sources: Mercor (TechCrunch, Sacra, Dealroom, Bloomberg); AfterQuery (Business Wire, Forbes); Snorkel AI (company); '
                            'Scale AI (Reuters, company blog); AI agents (Precedence Research); text stock (Epoch AI). See A2.', 8, C['grey'], line=1.1)
    footer(sh)


def p_market(sh):
    header(sh, 'Market size', 3, [('By 2030 the training market reaches ', False), ('$700B a year', True), (', 82x today', False)])
    gap, top, ch = 0.5, 1.66, 3.7
    cw = (W - 2 * gap) / 3
    xs = [0.6 + i * (cw + gap) for i in range(3)]
    for i in range(2):
        sh.t(xs[i] + cw, top + 0.4, gap, 0.6, '→', 14, C['grey'], algn='ctr', anchor='ctr')

    def rows(x0, w, items, y0, dark=False, pad=0.0):
        for i, (k, v) in enumerate(items):
            y, last = y0 + i * 0.4, i == len(items) - 1
            kc = (C['onDarkHi'] if last else C['onDark']) if dark else (C['ink'] if last else C['body'])
            vc = (C['accentLt'] if last else C['onDarkHi']) if dark else C['ink']
            sh.t(x0 + pad, y, w - 2 * pad - 1.1, 0.4, k, 10.5, kc, SANS, last, anchor='ctr')
            sh.t(x0 + w - pad - 1.1, y, 1.1, 0.4, v, 14, vc, SERIF, algn='r', anchor='ctr')
            if not last:
                sh.rect(x0 + pad, y + 0.4, w - 2 * pad, 0.01, DARK_RULE if dark else C['rule'])

    x0 = xs[0]
    sh.rect(x0, top, cw, 0.02, C['ink'])
    sh.t(x0, top + 0.14, cw, 0.22, 'Today · what labs buy', 10, C['grey'], MONO)
    sh.t(x0, top + 0.42, cw, 0.6, '$8.5B / year', 28, C['ink'], SERIF)
    sh.t(x0, top + 1.1, cw, 0.3, 'Training data and RL environments, 50+ vendors', 10.5, C['body'])
    sh.t(x0, top + 1.5, cw, 0.22, 'Top three by revenue (Surge: 2024; others run-rate)', 9, C['grey'])
    for i, (co, v, lab, col) in enumerate([('Mercor', 2.0, '$2.0B', C['ink']), ('Surge AI', 1.2, '$1.2B', C['mid']),
                                           ('Snorkel AI', 0.375, '$375M', C['mid'])]):
        y, w = top + 1.82 + i * 0.44, 1.5 * v / 2.0
        sh.t(x0, y, 1.15, 0.3, co, 12, C['ink'], SERIF, anchor='ctr')
        sh.rect(x0 + 1.15, y + 0.03, w, 0.24, col)
        sh.t(x0 + 1.15 + w + 0.08, y, 1.1, 0.3, lab, 11, C['ink'], anchor='ctr')
    x1 = xs[1]
    sh.rect(x1, top, cw, 0.02, C['ink'])
    sh.t(x1, top + 0.14, cw, 0.22, '2030 · AI economy', 10, C['grey'], MONO)
    sh.t(x1, top + 0.42, cw, 0.6, '~$7T / year', 28, C['ink'], SERIF)
    rows(x1, cw, [('US (McKinsey)', '$2.9T'), ('÷ US share of world GDP', '25.7%'), ('= world, same penetration', '$11.3T'),
                  ('= half penetration outside US', '$7.1T')], top + 1.12)
    x2, pad = xs[2], 0.26
    sh.rect(x2, top, cw, ch, C['ink'])
    sh.t(x2 + pad, top + 0.14, cw - 2 * pad, 0.22, '2030 · Our market', 10, C['accentLt'], MONO)
    sh.t(x2 + pad, top + 0.42, cw - 2 * pad, 0.6, '$700B / year', 28, C['accentLt'], SERIF)
    rows(x2, cw, [('AI economy', '$7T'), ('× training share', '10%'), ('= training market', '$700B')], top + 1.12, dark=True, pad=pad)
    sh.t(x2 + pad, top + 2.42, cw - 2 * pad, 0.5, '82x today', 22, C['onDarkHi'], SERIF, anchor='ctr')
    sh.t(x2 + pad, top + 3.0, cw - 2 * pad, 0.5, 'Training is the AI economy’s R&D budget. Big tech spends 10–15% of revenue on R&D; we use 10%.',
         9.5, C['onDark'], line=1.1)
    kicker(sh, 5.62, [('Today we sell to labs. By 2030, to ', False), ('the whole AI economy', True), ('.', False)], 20, 'ctr')
    sh.t(0.6, 6.32, W, 0.4, ['Sources: Menlo Ventures (Jul 2026); Mercor gross run-rate (Jun 2026), Surge revenue (2024), Snorkel run-rate (Sep 2026); '
                             'McKinsey (Nov 2025); IMF (Apr 2026); company filings.',
                             'Training share and non-US penetration are assumptions; see A2.'], 8, C['grey'], line=1.1, gap=0)
    footer(sh)


def p_competition(sh):
    header(sh, 'Competition', 3, [('Others do one piece. We run ', False), ('the whole loop that makes AI better', True)])
    rows = [('Expert data platforms', 'Mercor, Surge, AfterQuery', ['Expert demos and judgments, human-graded'], False),
            ('China peers', 'UniPat, Humanlaya', ['Mainly serve Chinese labs: expert data and evals'], False),
            ('AI eval companies', '', ['Measure today’s level; the score is the product'], False),
            ('Labs in-house', '', ['Build only for work they already know'], False),
            ('SimReal', 'Environments, data, grading, RSI', ['One environment for data, grading and training', 'Every result feeds the next round'], True)]
    lw, rh, gap, y0 = 7.0, 0.74, 0.1, 1.66
    for i, (k, sub, d, dark) in enumerate(rows):
        y = y0 + i * (rh + gap)
        sh.rect(0.6, y, lw, rh, C['ink'] if dark else C['tint'])
        sh.t(0.85, y + (0.08 if sub else 0.24), 2.75, 0.4, k, 16, C['onDarkHi'] if dark else C['ink'], SERIF)
        if sub:
            sh.t(0.85, y + 0.46, 2.75, 0.22, sub, 8.5, C['accentLt'] if dark else C['grey'], MONO)
        sh.t(3.7, y, lw - 3.3, rh, d, 11, C['onDarkHi'] if dark else C['body'], anchor='ctr', line=1.1)
    rx, rw = 8.0, 4.73
    sh.t(rx, y0, rw, 0.22, 'Our edge', 10, C['accent'], MONO)
    wins = [('Own assets', 'Environment tooling, verifiers, a professional task library, accumulated failure cases'),
            ('Stickiness', 'Every model upgrade needs new tasks, new data, new held-out sets and verifier upkeep'),
            ('Speed', '7 products in 14 days; self-improving market making in 2 weeks'),
            ('Neutral', 'Owned by no lab: US and Chinese labs can both buy with confidence')]
    for i, (k, d) in enumerate(wins):
        y = y0 + 0.32 + i * 0.98
        sh.rule(rx, y, rw)
        sh.t(rx, y + 0.1, rw, 0.38, k, 17, C['ink'], SERIF)
        sh.t(rx, y + 0.5, rw, 0.42, d, 10.5, C['body'], line=1.1)
    sh.text(0.6, 6.04, W, 0.44, [para([R('China today  ', 10, C['accent'], SANS, True),
                                      R('Per Bloomberg, Alibaba led a $300M round in UniPat at a $2.5B valuation; Humanlaya closed a CDH-led Pre-A '
                                        'of several hundred million RMB;', 10, C['body']), BR(10),
                                      R('Alibaba, ByteDance and DeepSeek already buy from both.', 10, C['body'])], line=1.15)], 'ctr')
    footer(sh)


def p_why_us(sh):
    header(sh, 'Why us', 3, [('The faster the hot spots change, ', False), ('the better for us', True)])
    xs, cw = [0.6 + i * (2.88 + 0.2) for i in range(4)], 2.88
    sh.t(0.6, 1.64, 9, 0.22, 'Every wave was won by the fastest young team', 10, C['grey'], MONO)
    for i, (bx, h) in enumerate(zip(xs, ['Data labeling', 'Expert data', 'RL environments', 'Self-improvement'])):
        dark = i == 3
        sh.rect(bx, 1.9, cw, 0.46, C['ink'] if dark else C['tint'])
        sh.t(bx + 0.26, 1.9, cw - 0.4, 0.46, h, 14, C['onDarkHi'] if dark else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(bx + cw, 1.9, 0.2, 0.46, '→', 11, C['grey'], algn='ctr', anchor='ctr')
    cards = [('Scale AI', '$29B', 'Valuation, 2025', ['Founded by Alexandr Wang at 19', 'MIT dropout, Y Combinator alum', ['Once the youngest', 'self-made billionaire']]),
             ('Mercor', '$10B', 'Valuation, 2025', [['Founded by three', 'high-school friends'], 'Youngest self-made billionaires at 22', 'Valuation up ~40x in 13 months']),
             ('AfterQuery', '$3.2B', 'Valuation, 18 months after YC', ['Founded by two students around 21', ['A founder interned at', 'Citadel Securities'], 'Fastest unicorn in YC history']),
             ('SimReal', '14 days', '7 products shipped', ['Quant founders, born 2005', 'Final year at Cambridge, LSE, Duke',
                                                         ['Jane Street, Citadel,', 'Optiver, Millennium']])]
    top, chh = 2.48, 3.4
    for i, (bx, (co, v, cap, lines)) in enumerate(zip(xs, cards)):
        dark = i == 3
        main, sub = (C['onDarkHi'], C['onDark']) if dark else (C['ink'], C['body'])
        sh.rect(bx, top, cw, chh, C['ink'] if dark else C['tint'])
        ix, iw = bx + 0.26, cw - 0.52
        sh.t(ix, top + 0.16, iw, 0.36, co, 17, main, SERIF)
        sh.t(ix, top + 0.52, iw, 0.7, v, 40, C['accentLt'] if dark else C['ink'], SERIF)
        sh.t(ix, top + 1.22, iw, 0.44, cap, 10, sub, line=1.0, gap=0)
        for j, ln in enumerate(lines):
            y = top + 1.7 + j * 0.54
            sh.rect(ix, y, iw, 0.01, DARK_RULE if dark else C['mid'])
            sh.t(ix, y + 0.06, iw, 0.46, ln, 10, sub, line=1.0, gap=0)
    kicker(sh, 6.06, [('Bet on us, and you bet on ', False), ('every paradigm shift in AI', True), ('.', False)], 20)
    footer(sh)


def p_business(sh):
    header(sh, 'Business model', 4, [('Every round brings a new delivery: ', False), ('the deeper the work, the more we earn', True)])
    sh.text(0.6, 1.64, 8.4, 0.34, [para([R('Three lines   ', 10, C['grey'], MONO),
                                         R('Environments  ·  AI data (trajectories, experts, evals)  ·  RSI service', 13, C['ink'], SERIF)])], 'ctr')
    sh.t(9.0, 1.64, 3.73, 0.34, 'Same model: AfterQuery hit $100M run-rate in 14 months', 9, C['accent'], algn='r', anchor='ctr')
    cols = [(0.8, 2.3, 'Stage'), (3.2, 3.3, 'What they buy'), (6.65, 2.6, 'How we charge'), (9.4, 3.15, 'Why they buy again')]
    ty = 2.22
    sh.rule(0.6, ty - 0.02, W, C['ink'])
    for cx, cw, h in cols:
        sh.t(cx, ty + 0.04, cw, 0.26, h, 9.5, C['grey'], MONO, anchor='ctr')
    stages = [('01', 'Paid pilot', ['Environment, data and eval pack', 'for one capability'], 'Fixed scope and acceptance criteria', 'Proves integration and training value'),
              ('02', 'License', ['Environments, task sets, verifiers,', 'agreed usage rights'], ['Term, scope, exclusivity;', 'data by volume'], 'Wider domains and task coverage'),
              ('03', 'Updates', ['New scenarios, tasks, trajectories,', 'private held-out sets'], 'Annual contract or batch orders', 'Old tasks saturate; new skills need new data'),
              ('04', 'Joint RSI training', ['Training runs, ablations,', 'transfer tests, hosted runs'], 'Project or ongoing service contract', 'Closes the customer’s next capability gap')]
    rh, y0 = 0.62, ty + 0.36
    sh.rule(0.6, y0, W)
    for i, (n, name, buy, fee, why) in enumerate(stages):
        y, dark = y0 + i * rh, i == 3
        if dark:
            sh.rect(0.6, y, W, rh, C['ink'])
        main, sub, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.t(0.8, y, 0.45, rh, n, 10, acc, MONO, anchor='ctr')
        sh.t(1.25, y, 1.9, rh, name, 14, main, SERIF, anchor='ctr')
        sh.t(3.2, y, 3.3, rh, buy, 10.5, sub, anchor='ctr', line=1.05, gap=0)
        sh.t(6.65, y, 2.6, rh, fee, 10.5, sub, anchor='ctr', line=1.05, gap=0)
        sh.t(9.4, y, 3.15, rh, why, 10.5, acc if dark else main, SANS, True, anchor='ctr', line=1.05)
        if not dark:
            sh.rule(0.6, y + rh, W)
    notes = [('Open source vs. paid', ['Open benchmarks: public, to build trust', 'Paid: private tasks, data, held-out sets, verifiers']),
             ('Non-exclusive vs. exclusive', ['Non-exclusive: one environment, many labs', 'Exclusive: by domain and term, premium (Epoch AI)']),
             ('Delivery is just the start', ['Verifier upkeep, difficulty upgrades', 'New tasks and data, hosted runs'])]
    by, bh, gap = 5.34, 1.12, 0.25
    bw = (W - 2 * gap) / 3
    for i, (h, lines) in enumerate(notes):
        bx = 0.6 + i * (bw + gap)
        sh.rect(bx, by, bw, bh, C['tint'])
        sh.t(bx + 0.26, by + 0.16, bw - 0.52, 0.34, h, 14, C['ink'], SERIF, anchor='ctr')
        sh.t(bx + 0.26, by + 0.54, bw - 0.52, 0.5, lines, 10, C['body'], gap=2)
    footer(sh)


def p_progress(sh, logos):
    header(sh, 'Traction', 4, [('Products are public; ', False), ('lab purchasing is under way', True)])
    stats = [('2', 'Frontier labs in talks', '', True), ('7,000+', 'Experts on the waitlist', '', False),
             ('400', 'GitHub stars in total', 'September 29, 2026', False), ('5', 'Angels reached out to us', '', False)]
    cw, gap = (W - 3 * 0.3) / 4, 0.3
    for i, (v, k, note, acc) in enumerate(stats):
        x = 0.6 + i * (cw + gap)
        sh.rect(x, 1.66, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, 1.82, cw, 0.8, v, 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.66, cw, 0.3, k, 12.5, C['ink'], SANS, True)
        if note:
            sh.t(x, 2.98, cw, 0.26, note, 10, C['grey'])
    sh.t(0.6, 3.6, 6, 0.22, 'Frontier lab purchasing', 10, C['grey'], MONO)
    stages = ['Contact', 'Talks', 'Paid pilot', 'Purchase']
    x0, x1, ly = 0.9, 12.43, 4.2
    step = (x1 - x0) / 3
    sh.rect(x0, ly - 0.005, x1 - x0, 0.01, C['mid'])
    sh.rect(x0, ly - 0.015, step, 0.03, C['accent'])
    for i, name in enumerate(stages):
        cx, done = x0 + i * step, i <= 1
        sh.dot(cx - 0.09, ly - 0.09, 0.18, C['accent'] if done else C['mid'])
        sh.t(cx - 0.8, ly + 0.2, 1.6, 0.3, name, 13, C['accent'] if i == 1 else (C['ink'] if done else C['grey']), SANS, i == 1, algn='ctr')
        if i == 1:
            sh.t(cx - 0.8, ly - 0.46, 1.6, 0.22, 'Now', 9.5, C['accent'], MONO, algn='ctr')
    sh.t(0.6, 5.0, 6, 0.22, 'Traders who support our R&D', 10, C['grey'], MONO)
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for (pic, _w, _h), cx in zip(logos, slots[:4]):
        if pic.shape_id == 403:
            pic.width, pic.height = int(pic.width * 1.45), int(pic.height * 1.45)
        pic.left, pic.top = int((cx - pic.width / 914400 / 2) * 914400), int((5.6 - pic.height / 914400 / 2) * 914400)
    sh.t(slots[4] - 1.0, 5.4, 2.0, 0.4, 'Polymarket', 16, C['ink'], SANS, True, algn='ctr', anchor='ctr')
    sh.t(0.6, 6.1, W, 0.26, 'Logos show where supporters work; they do not imply endorsement by these firms.', 8.5, C['grey'])
    footer(sh)


def p_network(sh):
    header(sh, 'Expert network', 4, [('200,000+ reachable, verified experts from 21 top universities: ', False), ('the source of our expert data', True)])
    for x, v, k in [(0.6, '200K+', 'Reachable, verified experts'), (2.85, '21', 'Top universities')]:
        sh.t(x, 1.72, 2.2, 0.8, v, 40, C['ink'], SERIF)
        sh.t(x, 2.52, 2.2, 0.26, k, 10.5, C['grey'])
    sh.rule(0.6, 3.0, 4.3)
    rows = [('Entry roles', 'The work AI takes over first, in every industry', C['ink']),
            ('Top research', 'The hardest problems and judgment calls', C['ink']),
            ('Moving up', 'Members get promoted; the network climbs to senior levels', C['accent'])]
    for i, (k, d, col) in enumerate(rows):
        y = 3.12 + i * 0.44
        sh.t(0.6, y, 1.2, 0.44, k, 13, col, SERIF, anchor='ctr')
        sh.t(1.85, y, 3.05, 0.44, d, 10, C['grey'], anchor='ctr', line=1.05)
    sh.t(5.53, 2.25, 5, 0.22, 'Selected universities in the network', 10, C['grey'], MONO)
    sh.rule(0.6, 5.92, W)
    sh.text(0.6, 6.0, W, 0.36, [para([R('Each domain leaves assets behind: ', 12, C['ink'], SERIF),
                                      R('experts → data and standards → environments and verifiers → failure library → ', 12, C['body'], SERIF),
                                      R('reuse in the next domain', 12, C['accent'], SERIF)])], 'ctr')
    sh.t(0.6, 6.5, W, 0.24, 'Network coverage does not mean registration or delivery; university logos do not imply endorsement.', 8, C['grey'])
    footer(sh)


def p_raise(sh):
    header(sh, 'The round', 5, [('Raising RMB 40M ', False), ('at RMB 500M post-money', True)])
    tiles = [('Raising', 'RMB 40M', 'Paid in USD', True), ('Dilution', '8%', '', False),
             ('Pre-money', 'RMB 460M', '', False), ('Post-money', 'RMB 500M', '', False)]
    tw, tg = (W - 3 * 0.25) / 4, 0.25
    for i, (k, v, n, dark) in enumerate(tiles):
        x = 0.6 + i * (tw + tg)
        sh.rect(x, 1.66, tw, 1.12, C['ink'] if dark else C['tint'])
        sh.t(x + 0.26, 1.8, tw - 0.5, 0.22, k, 9.5, C['accentLt'] if dark else C['grey'], MONO)
        sh.t(x + 0.26, 2.02, tw - 0.5, 0.56, v, 28, C['accentLt'] if dark else C['ink'], SERIF)
        if n:
            sh.t(x + 0.26, 2.5, tw - 0.5, 0.22, n, 9.5, C['onDark'] if dark else C['grey'])
    sh.text(0.6, 2.94, W, 0.36, [para([R('This round   ', 10, C['accent'], MONO, True),
                                       R('Environments: pilot → license → renewals and updates; AI data delivered at scale; RSI live in trading and event prediction',
                                         12, C['ink'])])], 'ctr')
    sh.t(0.6, 3.54, 9, 0.22, 'Valuation references: top investors have already priced this space', 10, C['grey'], MONO)
    comps = [('China · UniPat', '$2.5B', 'Reported valuation (Sep 2026)', 'Alibaba led; Tencent, HongShan joined'),
             ('China · Humanlaya', 'Several 100M RMB', 'Pre-A round size (Sep 2026)', 'CDH led; HongShan and others joined'),
             ('Global · Applied Compute', '$100M → ~$3B', 'Seed post → 14 months later (in talks)', '$20M seed round'),
             ('Global · AfterQuery', '$300M → $3.2B', 'Series A → 5 months later', 'Series A at $100M run-rate')]
    for i, (tag, v, when, who) in enumerate(comps):
        x = 0.6 + i * (tw + tg)
        sh.rect(x, 3.82, tw, 1.5, C['tint'])
        sh.t(x + 0.26, 3.96, tw - 0.5, 0.22, tag, 9, C['accent'], MONO)
        sh.t(x + 0.26, 4.22, tw - 0.5, 0.42, v, 17, C['ink'], SERIF)
        sh.t(x + 0.26, 4.66, tw - 0.5, 0.24, when, 9.5, C['grey'])
        sh.t(x + 0.26, 4.92, tw - 0.5, 0.36, who, 9.5, C['body'], line=1.1)
    kicker(sh, 5.54, [('HongShan, CDH, Capital Today, BAI, Alibaba and Tencent are already in: ', False), ('the category is validated', True), ('.', False)], 16)
    sh.t(0.6, 6.12, W, 0.24, 'UniPat: reported valuation; Applied Compute ~$3B: valuation in talks; Humanlaya: round size; others post-money. Sources in A2.',
         8.5, C['grey'])
    footer(sh)


def p_funds(sh):
    header(sh, 'Use of funds', 5, [('Two engines: revenue earns today’s money; ', False), ('RSI wins every industry', True)])
    top, colh, gap, pad = 1.66, 4.5, 0.3, 0.32
    cw = (W - gap) / 2
    iw = cw - 2 * pad
    eq = [('Environment revenue ', 12.5, 'main', SERIF, False), ('=', 12.5, 'main', SANS, False), (' environments ', 12.5, 'main', SERIF, False),
          ('×', 12.5, 'main', SANS, False), (' licenses ', 12.5, 'main', SERIF, False), ('×', 12.5, 'main', SANS, False),
          (' price', 12.5, 'main', SERIF, False)]
    engines = [
        dict(dark=False, label='01  Revenue engine  ·  RMB 20M', big='Environments & data', sub='Build once, license to many labs; data billed by volume',
             uses=[('RMB 7M', 'Environments', 'Batch-build environments by industry'), ('RMB 4M', 'Data', 'Agent trajectories, expert and eval data'),
                   ('RMB 6M', 'Delivery', 'Integration, acceptance, updates'), ('RMB 3M', 'Sales & ops', 'Frontier labs and enterprises')],
             block=[eq, [('Environment price $20K–$300K (Epoch AI)', 10, 'sub', SANS, False)]],
             ms=[('3 months', 'First paid pilot'), ('6 months', 'Licenses, renewals'), ('12 months', 'Update contracts')]),
        dict(dark=True, label='02  RSI engine  ·  RMB 20M', big='Self-improvement', sub='The whole AI economy: in every industry, AI that improves itself',
             uses=[('RMB 9M', 'Compute', 'Models train themselves, each round stronger'), ('RMB 7M', 'Research', 'From trading to 10 industries'),
                   ('RMB 4M', 'Live trading', 'Real money and settlement; compliance')],
             block=[[('Trading RSI  ·  Prediction RSI  ·  10 industries', 12.5, 'main', SERIF, False)],
                    [('A self-improving environment for every industry, in parallel', 10, 'sub', SANS, False)]],
             ms=[('3 months', 'Trading RSI live'), ('6 months', 'Prediction RSI live'), ('12 months', 'Across 10 industries')]),
    ]
    for k, eng in enumerate(engines):
        x0, dark = 0.6 + k * (cw + gap), eng['dark']
        col = dict(main=C['onDarkHi'] if dark else C['ink'], sub=C['onDark'] if dark else C['body'],
                   acc=C['accentLt'] if dark else C['accent'], rule=DARK_RULE if dark else C['rule'])
        sh.rect(x0, top, cw, colh, C['ink'] if dark else C['tint'])
        ix = x0 + pad
        sh.t(ix, top + 0.24, iw, 0.22, eng['label'], 10, col['acc'], MONO, True)
        sh.t(ix, top + 0.5, iw, 0.56, eng['big'], 28, col['acc'], SERIF)
        sh.t(ix, top + 1.12, iw, 0.3, eng['sub'], 11.5, col['main'], anchor='ctr')
        y0, rh = top + 1.56, 0.34
        sh.rect(ix, y0, iw, 0.01, col['rule'])
        for i, (amt, item, det) in enumerate(eng['uses']):
            y = y0 + i * rh
            sh.t(ix, y, 0.95, rh, amt, 13, col['acc'], SERIF, anchor='ctr')
            sh.t(ix + 1.0, y, 1.45, rh, item, 12.5, col['main'], SERIF, anchor='ctr')
            sh.t(ix + 2.5, y, iw - 2.5, rh, det, 9.5, col['sub'], anchor='ctr')
            sh.rect(ix, y + rh, iw, 0.01, col['rule'])
        sh.text(ix, top + 3.0, iw, 0.7, [para([R(t, sz, col[c], f, bold) for t, sz, c, f, bold in line], before=0 if j == 0 else 4)
                                         for j, line in enumerate(eng['block'])])
        my = top + 3.82
        sh.rect(ix, my - 0.06, iw, 0.01, col['rule'])
        mw = iw / 3
        for i, (t, d) in enumerate(eng['ms']):
            sh.t(ix + i * mw, my, mw - 0.1, 0.2, t, 9.5, col['acc'], MONO)
            sh.t(ix + i * mw, my + 0.22, mw - 0.05, 0.34, d, 11.5, col['main'], SERIF)
    sh.t(0.6, 6.3, W, 0.22, 'Milestones are this round’s targets.', 8.5, C['grey'])
    footer(sh)


def p_closing(sh):
    sh.t(0.6, 2.4, 8.4, 1.7, ['Every industry’s best AI', 'will come from our worlds.'], 40, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 4.4, 7.8, 0.34, 'Three worlds live: trading, AI research, event prediction', 15, C['grey'])
    sh.t(0.6, 4.84, 8.4, 0.4, 'SimReal  ·  The data and self-improvement engine of the AI economy', 17, C['accent'], SERIF)
    sh.rule(0.6, 5.55, 5.6)
    sh.t(0.6, 5.7, 7.8, 0.3, 'business@simreal.co  ·  simreal.co  ·  x.com/simrealhq', 11, C['ink'], MONO)
    footer(sh)


def appendix_header(sh, title):
    sh.text(0.6, 0.42, 6, 0.24, [para([R(CUR['n'], 10, C['accent'], MONO, True), R('    Appendix', 11, C['grey'], SANS, True)])], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, title, 30, C['ink'], SERIF)


def p_a1(sh):
    appendix_header(sh, 'Xitadel: what we tested and how')
    rows = [('Environment', 'Replays real trading days and order books with data licensed through IMC Trading; the market settles every trade'),
            ('Model', 'Open-source Qwen3.8-27B, before vs. after training'),
            ('Test data', 'Real trading days the model had never seen'),
            ('Result', 'Up to 12% better trading than the base model, replicated across independent runs (controlled experiment)'),
            ('Public benchmark', 'Xitadel public preview: human reference 80; best frontier model 77.28 (GPT 6); no model above the human line yet'),
            ('Diligence', 'Run logs, metric definitions and scripts')]
    for i, (k, d) in enumerate(rows):
        y = 1.7 + i * 0.62
        sh.rule(0.6, y, W)
        sh.t(0.6, y, 1.6, 0.62, k, 11, C['grey'], anchor='ctr')
        sh.t(2.3, y, W - 1.7, 0.62, d, 12.5, C['ink'], anchor='ctr')
    sh.rule(0.6, 1.7 + 6 * 0.62, W)
    footer(sh)


SOURCES_L = [
    ('Mercor run-rate: ', 'TechCrunch (Feb 2025, $75M); CEO (Sep 2025, $500M); Sacra (Dec 2025, $760M); Mercor (early 2026, $1B); '
                          'Dealroom (Jun 2026, $2B). All gross; experts take 60–70% (Bloomberg)'),
    ('Mercor valuation: ', 'Series A $250M (Sep 2024), Series B $2B (Feb 2025), Series C $10B (Oct 2025); a $20B round in early talks (Bloomberg, Jul 9, 2026)'),
    ('Surge AI: ', '2024 revenue $1.2B (TechCrunch, Forbes)'),
    ('Snorkel AI: ', 'Company and TechCrunch, Sep 22, 2026 ($350M raised at a $3.5B valuation; $375M run-rate, up 18x in a year)'),
    ('AfterQuery: ', 'Founded Feb 2025; Forbes, Sep 1, 2026 (18 months after YC, $3.2B valuation); Business Wire (Apr 2026, Series A at $300M, $100M run-rate); '
                     'YC profile (co-founder interned at Citadel Securities)'),
    ('Scale AI: ', 'Founded in 2016 by Alexandr Wang, then 19 (Forbes); Meta bought 49% for $14.3B, ~$29B valuation (Reuters, Jun 2025); '
                   'nearly half of new training projects involve RL environments (Scale AI blog, Feb 2026)'),
    ('UniPat: ', 'Bloomberg, Sep 10, 2026 (Alibaba-led $300M round, reported $2.5B valuation, with Tencent and HongShan; terms may change; '
                 'Alibaba, ByteDance, DeepSeek and others buy from UniPat and Humanlaya)'),
    ('Humanlaya: ', 'Founded Oct 2025; Pre-A of several hundred million RMB on Sep 9, 2026, led by CDH with HongShan, Capital Today and BAI (Jiemian, Eastmoney)'),
]
SOURCES_R = [
    ('Training data spend today: ', 'Menlo Ventures (Jul 2026): 50+ vendors, ~$8.5B combined'),
    ('2030 AI economy: ', 'McKinsey, Agents, robots, and us (Nov 2025): US ~$2.9T; IMF WEO (Apr 2026): US GDP $32.4T, world $126.3T; '
                          'half the US penetration assumed elsewhere'),
    ('Training market: ', '2030 AI economy × 10%; 10% is the low end of big tech’s 10–15% R&D-to-revenue ratio; an assumption'),
    ('AI agent market: ', 'Precedence Research ($7.92B in 2025; 45.82% CAGR); 2028 and 2032 extrapolated'),
    ('RL environment pricing: ', 'Epoch AI, An FAQ on RL environments (Jan 2026)'),
    ('Public text stock: ', 'Epoch AI, Will we run out of data? (2024): exhausted 2026–2032 (80% confidence interval)'),
    ('Rounds and valuations: ', 'Applied Compute: Upstarts (Jun 2025, $20M seed at $100M post); The Information (Aug 2026, ~$3B in talks)'),
]


def p_a2(sh):
    appendix_header(sh, 'Sources')
    for x, head, items in [(0.6, 'Companies', SOURCES_L), (6.95, 'Market and research', SOURCES_R)]:
        sh.t(x, 1.62, 5.78, 0.24, head, 10, C['accent'])
        sh.rule(x, 1.92, 5.78, C['ink'])
        sh.text(x, 2.04, 5.78, 4.7, [para([R(k, 8.5, C['ink'], SANS, True), R(v, 8.5, C['body'])], before=0 if i == 0 else 5, line=1.1)
                                      for i, (k, v) in enumerate(items)])
    footer(sh)


def p_a3(sh):
    appendix_header(sh, 'Glossary')
    terms = [('Training environment', 'A system where AI does real work again and again and learns from the results; an RL environment'),
             ('Agent trajectory', 'The full record of AI doing a task: every action, the world’s response, the result'),
             ('Real feedback', 'The world’s reaction to AI’s actions: P&L, the books, whether code runs, whether events happen'),
             ('Self-improvement (RSI)', 'AI trains on its own results in real environments, stronger each round'),
             ('Held-out set', 'Test-only tasks, never used in training, so models can’t memorize answers'),
             ('Red-teaming', 'Simulated cheating and attacks before launch, to find and fix grading loopholes'),
             ('Controlled experiment', 'Change one thing (training in Xitadel or not) and compare before and after'),
             ('Gross run-rate', 'Current revenue annualized, including payouts to experts')]
    cw, gap = (W - 3 * 0.3) / 4, 0.3
    for i, (k, d) in enumerate(terms):
        x, y = 0.6 + (i % 4) * (cw + gap), 1.8 + (i // 4) * 2.1
        sh.rule(x, y, cw, C['ink'])
        sh.t(x, y + 0.14, cw, 0.22, f'{i + 1:02d}', 9.5, C['accent'], MONO)
        sh.t(x, y + 0.42, cw, 0.44, k, 16, C['ink'], SERIF)
        sh.t(x, y + 0.94, cw, 0.9, d, 10.5, C['body'], line=1.12)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_problem, 4, {160}, 4),
    (p_insight, 5, {169}, 5),
    (p_solution, 6, {237}, 6),
    (p_xitadel, 7, {266}, 7),
    (p_products, 8, {300}, 8),
    (p_data, 9, {336}, 9),
    (p_why_now, 3, {126}, 10),
    (p_market, 13, {470}, 11),
    (p_competition, 14, {511}, 12),
    (p_why_us, 15, {562}, 13),
    (p_business, 10, {376}, 14),
    (p_progress, 11, {401, 402, 403, 404, 407}, 15),
    (p_network, 12, set(range(429, 441)) | {442}, 16),
    (p_raise, 16, {607}, 17),
    (p_funds, 17, {607}, 18),
    (p_closing, 18, {615, 616, 617, 623}, None),
    (p_a1, 19, {654}, 'A1'),
    (p_a2, 21, {682}, 'A2'),
    (p_a3, 22, {716}, 'A3'),
]


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s = slides[idx]
        tree = s.shapes._spTree
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                tree.remove(shp._element)
        sh = S(5000)
        if build is p_progress:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            build(sh, [(p, p.width, p.height) for p in logos])
        elif build is p_closing:
            for p in s.shapes:
                if p.shape_id == 615:
                    p.top = int(-0.15 * 914400)
            build(sh)
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(idx)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for idx in order:
        lst.append(ids[idx])
    for idx, el in enumerate(ids):
        if idx not in order:
            prs.part.drop_rel(el.rId)
    for s in prs.slides:            # English deck: mark every run as en-US
        for el in s._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'):
            el.set('lang', 'en-US')
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
