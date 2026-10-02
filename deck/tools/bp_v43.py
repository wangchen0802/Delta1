"""Build SimReal (衍真) BP v43 (14 slides, Chinese, no round terms) from v17's slides.

v41 without the raise: no amount, valuation or use of funds anywhere; the last page is the team's edge
instead (delivery speed, and moving into new tracks: personal-agent environments in development). The
why-us page keeps scoring and training; delivery moves to the last page. Wording tightened throughout.

Usage: python3 bp_v43.py SimReal-BP-v17.pptx out.pptx
       then embed_cjk.py adds the Chinese font subsets to the shipping PPTX.
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['我们是谁', '问题与机会', '我们的答案', '优势与竞争', '商业与进展', '团队优势']
CONF = '机密  ·  仅供受邀投资机构内部评估使用  ·  2026年10月'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.36, 6.98, 0.5, 0.26, '衍真', 9, C['ink'], SANS, True, anchor='ctr')
    sh.t(1.85, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def p_cover(sh):
    sh.t(1.62, 0.55, 1.0, 0.25, '衍真', 13, C['ink'], SANS, True, anchor='ctr')
    sh.t(0.6, 1.35, 8.6, 1.9, ['让AI在真实世界里', '自我进化'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, '用真实数据重建真实世界', 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, 'AI实验室、企业与个人Agent的训练场', 16, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    for (k, v), x in zip([('商业计划书', '2026年10月'), ('联系', 'business@simreal.co')], [0.6, 4.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, '项目概述', 0, [('给AI一个真实世界，让它反复犯错、学习、', F), ('自我进化', T)])
    cells = [
        ('做什么', '训练环境与AI数据', ['面向AI实验室与企业', 'Agent轨迹、专家与评测数据；独家环境（RSI）'], F),
        ('已做到', '14天7款产品', ['3个世界已上线：交易、AI研究、事件预测', '5个公开仓库；零外部融资'], F),
        ('核心成果', '+12%', ['Qwen3.8-27B经Xitadel训练', '未见过的交易日上最高提升12%，多次独立复现'], T),
        ('团队', '05后量化创始团队', ['剑桥、LSE、杜克数学本科', 'Jane Street、Citadel、Optiver、Millennium'], F),
        ('市场', '85亿 → 7,000亿美元', ['今天：训练数据与RL环境供应商年收入', '2030年情景：今天的82倍'], F),
        ('下一个赛道', '个人Agent', ['练习环境开发中', '跨应用长任务、失败恢复'], T),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.72 + (i // 3) * 2.0
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, v, 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.8, d, 11, C['body'], line=1.15, gap=0)
    kicker(sh, 5.85, [('谁有最好的训练世界，谁就有最好的AI：', F), ('从交易出发，走向整个世界。', T)], 18)
    footer(sh)


def p_team(sh):
    header(sh, '团队', 0, [('数学、机器学习、量化与数据科学背景；', F), ('放弃顶级量化机构转正offer来做这件事', T)])
    b, sz = C['body'], 10.5
    people = [
        ('Charles', 'CEO', [(R('United Stables首位员工：帮助U稳定币', sz, b), BR(sz), R('一年内从0做到14亿美元，一个月上线Binance', sz, b)),
                            '负责机构关系，参与SIG、DRW等合作',
                            (R('汇丰港元稳定币项目唯一实习生，', sz, b), BR(sz), R('协助推进香港金管局（HKMA）合规', sz, b)),
                            'X（Twitter）博主，内容数百万浏览', '伦敦Citadel对冲基金实习'],
         ['2005年生 · LSE数学本科 · 深国交', '美国数学奥林匹克（USAMO）入围']),
        ('Henry', 'CTO', [(R('剑桥机器学习暑研，师从统计学教授', sz, b), BR(sz), R('Po-Ling Loh（国际数理统计学会会士）', sz, b)),
                          '剑桥研究中心AI最年轻本科研究员', 'Jane Street、Citadel、Optiver量化经历', '设计五大Benchmark与强化学习环境'],
         ['2005年生 · 剑桥数学一等荣誉（奖学金） · 深国交', '剑桥数学竞赛全球前30 · 英国物理竞赛超级金奖']),
        ('Amaris', 'COO', ['Millennium香港数据科学家', 'Millennium另类数据团队史上首位应届招聘',
                           (R('参与落地Plug and Play香港首场活动', sz, b), BR(sz), R('（联合香港科技园，200+人）', sz, b)),
                           '强生MedTech科技峰会主持人（200+人）', '17岁出版作家；原创内容获10万+互动'],
         ['2005年生 · 杜克数学与统计本科 · 上海包玉刚']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.55
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.76, cw - 0.6, 2.0, lines, sz, b, line=1.1, gap=4)
        sh.rule(x + 0.3, y0 + 2.86, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 2.94, cw - 0.6, 0.55, edu, 9.5, C['grey'], line=1.12, gap=0)
    kicker(sh, 6.1, [('交易台纪律、机器学习研究、机构关系：', F), ('做专业训练环境要的三样能力', T)], 17)
    footer(sh)


def p_problem(sh):
    header(sh, '问题', 1, [('AI已经学会推理，', F), ('却还做不好真实世界里的事', T)])
    sh.t(0.6, 1.64, W, 0.3, '真实工作没有标准答案，只有世界的反应', 12, C['grey'])
    rows = [('交易', '历史回测一路上涨', '进入未见过的真实行情就亏钱'), ('软件工程', '测试全部通过', '上线后出错'),
            ('财务', '账看起来已经做完', '月结对不平'), ('事件预测', '分析头头是道', '结果一出就落空')]
    y0, rh = 2.38, 0.62
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, '表面检查：通过', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, '真实世界：失手', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 17, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 17, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    ev = [('6/32', '前沿模型用真钱交易，32轮只有6轮赚钱（Alpha Arena）'), ('77 < 80', 'Xitadel上最强的GPT 6总分77，仍低于人类80')]
    cw, xs = cols(2, 0.5)
    for (v, d), x in zip(ev, xs):
        sh.t(x, 4.95, 1.7, 0.5, v, 22, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.75, 4.95, cw - 1.75, 0.5, d, 11.5, C['body'], anchor='ctr', line=1.1)
    kicker(sh, 5.75, [('AI要学会做事，需要一个', F), ('用真实结果结算每个动作的世界', T), ('。', F)])
    footer(sh)


def p_why_now(sh):
    header(sh, '为什么是现在', 1, [('训练的瓶颈正从数据转向环境：', F), ('预算已到位，专业环境还没人做好', T)])
    items = [('01', '训练方式变了', ['推理模型靠可自动验证的结果做强化学习', '公开人类文本预计2026–2032年用尽'],
              '稀缺的是能自动打分的环境'),
             ('02', '预算已经到位', [(R('每个RL任务200–2,000美元；', 11.5, C['body']), BR(11.5), R('环境合同每季度六到七位数美元', 11.5, C['body'])),
                                 'Mercor在2026年7月收购环境公司Deeptune'],
              '环境已是实验室的新采购品类'),
             ('03', '专业任务还没解决', ['前沿模型真钱交易，32轮仅6轮盈利',
                                    (R('Meta Muse、OpenAI Dots', 11.5, C['body']), BR(11.5), R('在2026年9月开始替人做事', 11.5, C['body']))],
              '交易对错易验证、差距大：最适合先做')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for (n, k, ev, m), x in zip(items, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 21, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.08, cw - 0.6, 1.2, ev, 11.5, C['body'], line=1.15, gap=6)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, m, 12)
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, '窗口期', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R('RL环境约20家早期公司，预计收敛到3–5家（Wing VC）：', 12.5, C['onDark']),
                                       R('先做出可验证的增益，先拿实验室长期合同', 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr', line=1.1)
    note(sh, '来源：Epoch AI（公开文本存量预测；《An FAQ on RL environments》，2026年1月）；SiliconANGLE、TechCrunch（2026年7月）；'
             'Nof1（2026年5月）；Meta、OpenAI发布（2026年9月）；Wing VC（2026年）。', 6.3)
    footer(sh)


def p_market(sh):
    header(sh, '市场空间', 1, [('训练数据与环境的支出正在加速：', F), ('2030年情景约为今天的82倍', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, '当前：训练数据与RL环境')
    sh.t(0.9, 2.25, 5, 0.72, '85亿美元 / 年', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, '50余家供应商收入估算（Deedy Das，2026年7月）', 11, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030年：公司情景测算', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '7,000亿美元 / 年', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, '约为当前的82倍；假设AI经济规模7万亿美元、训练投入占10%', 11, C['body'])
    growth = [('27倍', 'Mercor年化毛营收：16个月，7,500万→20亿美元', '专家数据需求在爆发'),
              ('10倍', 'Mercor估值：17个月，20亿→200亿美元（洽谈中）', '资本已经为这个赛道定价'),
              ('18倍', 'Snorkel AI近一年增长，年化收入3.75亿美元', '数据服务正转向专家与环境')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 11.5, C['body'], line=1.1)
        means(sh, x, 5.34, cw, m, 12)
    note(sh, '来源：Mercor（TechCrunch、Sacra、Dealroom）；Snorkel AI公告（2026年9月）；Deedy Das行业图谱（2026年7月）。2030年为公司情景测算，非独立研究预测。', 6.3)
    footer(sh)


def p_customers(sh):
    header(sh, '客户与应用', 2, [('从专业任务训练，', F), ('拓展到个人Agent', T)])
    cards = [('AI实验室', '现在', '2家前沿实验室在谈', [(R('采购专业数据与训练环境，', 11, C['onDarkHi']), BR(11), R('用于模型训练和能力验证', 11, C['onDarkHi']))],
              '环境授权、Agent轨迹、专家数据与评测'),
             ('企业Agent团队', '拓展', '', [(R('把业务流程转成可执行任务，', 11, C['body']), BR(11), R('在部署前训练、测试与验收', 11, C['body']))],
              '定制任务、评分工具与验收评测'),
             ('个人Agent开发团队', '下一步', '', [(R('Meta Muse：', 11, C['ink'], SANS, True), R('独立虚拟机与浏览器，', 11, C['body']), BR(11),
                                                R('跨应用执行任务，持续推进用户的长期目标', 11, C['body'])),
                                               (R('OpenAI Dots：', 11, C['ink'], SANS, True), R('常驻云端，', 11, C['body']), BR(11),
                                                R('调用工具、分派任务，用户离线后继续工作', 11, C['body']))],
              '长任务、跨应用与失败恢复的练习环境')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.75
    for (k, tag, extra, need, deliver), x in zip(cards, xs):
        dark = tag == '现在'
        main, sub, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(x + 0.3, top + 0.22, cw - 0.6, 0.22, tag, 9.5, acc, MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 20, main, SERIF)
        if extra:
            sh.t(x + 0.3, top + 0.98, cw - 0.6, 0.26, extra, 11, acc, SANS, True)
        sh.t(x + 0.3, top + 1.3, cw - 0.6, 0.2, '他们要什么', 9, sub, MONO)
        sh.t(x + 0.3, top + 1.54, cw - 0.6, 1.2, need, 11, sub if not dark else C['onDarkHi'], line=1.15, gap=4)
        sh.rect(x + 0.3, top + 2.72, cw - 0.6, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(x + 0.3, top + 2.84, cw - 0.6, 0.2, '我们交付', 9, acc, MONO)
        sh.t(x + 0.3, top + 3.08, cw - 0.6, 0.6, deliver, 12, main, SANS, True, line=1.1)
    note(sh, '个人Agent产品示例：Meta（2026年9月8日）、OpenAI（2026年9月29日）官方发布；二者为行业示例，不是衍真客户。', 5.72)
    kicker(sh, 6.05, [('同一批环境，三类客户：', F), ('先卖实验室，再到企业与个人Agent', T)], 16)
    footer(sh)


def p_solution(sh):
    header(sh, '解决方案', 2, [('用任务结果训练模型，', F), ('用独立评测检验改进', T)])
    steps = [('基础Agent', ['接入客户已有模型', '与工具调用能力'], F), ('按需适配', ['监督微调，学习示范', '人类反馈，对齐偏好'], F),
             ('环境训练', ['按结果奖励，强化学习', '更新模型或策略'], T), ('独立评测', ['用未见任务验证效果', '达到约定标准后交付'], F)]
    cw, xs = cols(4, 0.4)
    top, bh = 1.72, 1.35
    for i, ((k, lines, acc), x) in enumerate(zip(steps, xs)):
        sh.rect(x, top, cw, bh, C['accent'] if acc else C['tint'])
        sh.t(x + 0.25, top + 0.14, 1, 0.2, f'0{i + 1}', 9, C['onDarkHi'] if acc else C['accent'], MONO)
        sh.t(x + 0.25, top + 0.36, cw - 0.5, 0.42, k, 19, C['onDarkHi'] if acc else C['ink'], SERIF)
        sh.t(x + 0.25, top + 0.84, cw - 0.5, 0.48, lines, 10.5, C['onDarkHi'] if acc else C['body'], line=1.1, gap=0)
        if i < 3:
            sh.t(x + cw, top, 0.4, bh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    cx = xs[2] + cw / 2
    sh.rect(cx, top + bh, 0.012, 0.3, C['mid'])
    sh.t(cx + 0.08, top + bh + 0.02, 2.4, 0.26, '训练中的薄弱任务', 9, C['grey'], MONO, anchor='ctr')
    loop_top = 3.42
    sh.rect(0.6, loop_top, W, 1.25, C['tint'])
    sh.t(0.85, loop_top + 0.16, 4, 0.22, '自我出题循环 · 研发方向', 9.5, C['accent'], MONO, True)
    loop = [('分析训练失败', '定位训练中的薄弱任务'), ('智能体出题', '生成或挑选新任务'), ('任务与评分校验', '规则、执行器与专家验证'), ('回到环境训练', '校验通过后进入下一轮')]
    lw, lxs = cols(4, 0.4, 0.85, W - 0.5)
    for i, ((k, d), x) in enumerate(zip(loop, lxs)):
        sh.t(x, loop_top + 0.46, lw, 0.34, k, 15, C['accent'] if i == 3 else C['ink'], SERIF)
        sh.t(x, loop_top + 0.82, lw, 0.28, d, 10.5, C['body'])
        if i < 3:
            sh.t(x + lw, loop_top + 0.46, 0.4, 0.34, '→', 12, C['grey'], algn='ctr', anchor='ctr')
    sh.t(0.6, 4.78, W, 0.26, '自主出题为研发方向；独立评测任务不参与训练或出题', 10, C['grey'])
    sh.rect(0.6, 5.2, W, 0.66, C['ink'])
    sh.t(0.85, 5.2, 2.2, 0.66, '衍真交付', 14, C['accentLt'], SERIF, anchor='ctr')
    sh.t(3.0, 5.2, W - 2.6, 0.66, '专业数据  ·  训练环境  ·  评分工具  ·  评测报告', 15, C['onDarkHi'], SERIF, anchor='ctr')
    note(sh, '参考：InstructGPT（arXiv 2203.02155）；DeepSeek-R1（arXiv 2501.12948）；Absolute Zero（arXiv 2505.03335）。', 6.2)
    footer(sh)


def p_products(sh):
    header(sh, '产品', 2, [('五款评测产品已公开；', F), ('训练基础设施压低每轮成本', T)])
    rows = [('SimReal-MLBench', '60个真实竞赛任务，评测机器学习研究能力'),
            ('Xitadel-QuantBench', '回放订单簿，对标同资产上的专业交易员；保留轨迹'),
            ('MathmoBench', '高阶数学题，检验推理与解题能力'),
            ('Puzzle Benchmark', '749道人工推理题，覆盖15类主题'),
            ('FuturePredict Bench', '截止前锁定概率与证据，按真实结局评分')]
    x0, y0, rh = 6.55, 1.72, 0.54
    sh.rule(x0, y0, W + 0.6 - x0, C['ink'])
    for i, (k, d) in enumerate(rows):
        y = y0 + i * rh
        sh.t(x0, y + 0.04, 6.2, 0.26, k, 13.5, C['ink'], SERIF)
        sh.t(x0, y + 0.28, 6.2, 0.24, d, 10.5, C['body'])
        sh.rule(x0, y + rh, W + 0.6 - x0)
    label(sh, 0.6, 4.78, 8, '训练基础设施：同样的算力，跑更多训练', C['accent'])
    infra = [('1/4', '单轮耗时', 'microVM深度隔离'), ('+64%', '同等资源完成的已评分尝试', '任务资源分配（Rollout allocation）'),
             ('−1/3', '存档耗时，存档次数减半', '智能存档：自动判断何时保存检查点')]
    cw, xs = cols(3)
    for (v, k, how), x in zip(infra, xs):
        sh.rect(x, 5.08, cw, 0.02, C['ink'])
        sh.t(x, 5.16, 1.55, 0.62, v, 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.6, 5.16, cw - 1.6, 0.62, [k, (R(how, 9.5, C['grey']),)], 11, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=0)
    note(sh, '公开仓库截至2026年10月2日。三项基础设施指标各有比较口径，不能相乘或相加。', 6.3)
    footer(sh)


def p_why_us(sh):
    header(sh, '为什么是我们', 3, [('计分与训练，', F), ('两件最难的事都有结果', T)])
    cards = [('计分做对', ['交易台经验写进规则：逐笔撮合；盈亏扣除回撤再乘夏普', '最后一个交易日留出，结果哈希预先公开'],
              ['分数拉得开：GPT 6得77分，其余前沿模型24–30分', '人类最佳80分，尚无模型越过']),
             ('训练有增益', ['同一环境既评测又训练', 'microVM隔离、任务资源分配、智能存档压低每轮成本'],
              ['Qwen3.8-27B最高提升12%', '同等资源多完成64%的已评分尝试'])]
    cw, xs = cols(2, 0.3)
    top, ch = 1.72, 3.35
    for (k, how, res), x in zip(cards, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.35, top + 0.24, cw - 0.7, 0.5, k, 24, C['ink'], SERIF)
        sh.t(x + 0.35, top + 0.9, cw - 0.7, 0.2, '做法', 9, C['grey'], MONO)
        sh.t(x + 0.35, top + 1.14, cw - 0.7, 0.8, how, 13, C['body'], line=1.15, gap=4)
        sh.rule(x + 0.35, top + 2.0, cw - 0.7, C['mid'])
        sh.t(x + 0.35, top + 2.14, cw - 0.7, 0.2, '结果', 9, C['accent'], MONO)
        sh.t(x + 0.35, top + 2.38, cw - 0.7, 0.85, res, 15.5, C['ink'], SANS, True, line=1.15, gap=4)
    sh.rect(0.6, 5.32, W, 0.72, C['ink'])
    sh.t(0.85, 5.32, 2.0, 0.72, '越做越难复制', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.95, 5.32, W - 2.6, 0.72, [(R('每个客户都会留下任务库、评分器、失败案例与验证过的专家，', 12.5, C['onDark']),
                                       R('复用到下一个客户和下一个领域', 12.5, C['onDarkHi'], SANS, True))], 12.5, anchor='ctr', line=1.1)
    note(sh, '来源：Xitadel-QuantBench公开报告与计分规范；训练基础设施指标为公司内部测试。', 6.3)
    footer(sh)


def p_competition(sh):
    header(sh, '竞争与优势', 3, [('从交易与事件预测切入：', F), ('别人强在规模，我们强在专业计分与速度', T)])
    xs, ws = [0.6, 3.75, 6.55, 9.45], [3.0, 2.65, 2.75, 3.28]
    heads = ['客户已有的选择', '他们的强项', '他们的短板', '衍真的优势']
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 0.98, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    rows = [('数据与环境供应商', 'Scale AI、AfterQuery、Mercor', ['专家网络大，实验室关系深'],
             ['以通用领域为主；', '交易的撮合与风险计分', '需要交易台经验'],
             ['交易台出身设计撮合与风险计分，', '同一策略永远同一分']),
            ('公开基准与评测', 'Alpha Arena等榜单', ['方便比较模型能力'], ['只做评测，不产出训练数据；', '公开后很快饱和'],
             ['同一环境产出训练数据；', '储备任务持续换题']),
            ('客户内部自建', '', ['贴合自身工具和业务'], ['缺专业规则与专家，占研究人力'],
             ['现成的交易、事件预测环境', '与专家网络，按需交付'])]
    y0, rh = 2.04, 0.98
    for i, (k, names, pro, con, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.85, rh, [para([R(k, 15, C['ink'], SERIF)])] + ([para([R(names, 8.5, C['grey'], MONO)])] if names else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.15, rh, pro, 11.5, C['body'], anchor='ctr', line=1.1, gap=0)
        sh.t(xs[2], y, ws[2] - 0.15, rh, con, 11.5, C['body'], anchor='ctr', line=1.1, gap=0)
        sh.t(xs[3], y, ws[3] - 0.1, rh, us, 12, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=0)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.t(0.6, 5.15, 1.6, 0.4, '可复用资产', 13, C['accent'], SERIF, anchor='ctr')
    sh.t(2.2, 5.15, W - 1.6, 0.4, '任务库 · 评分器 · 隔离测试集 · 专家数据 · 失败案例：随客户和模型迭代持续积累', 12, C['ink'], anchor='ctr')
    kicker(sh, 5.75, [('赛道已被定价：', F), ('据彭博，阿里拟领投UniPat 3亿美元，估值约25亿美元', T), ('。', F)], 17)
    footer(sh)


def p_business(sh):
    """v39's business model with the data page folded in: its three data types are what the AI-data line delivers."""
    header(sh, '商业模式', 4, [('面向AI实验室与企业，', F), ('按交付范围与使用权收费', T)])
    b = C['body']
    data = [(R('Agent轨迹', 11.5, C['ink'], SANS, True), R('  用于监督微调与强化学习', 11, b)),
            (R('专家数据', 11.5, C['ink'], SANS, True), R('  用于对齐、奖励建模与任务设计', 11, b)),
            (R('评测数据', 11.5, C['ink'], SANS, True), R('  用于能力验收与持续评测', 11, b))]
    lines = [('训练环境', ['可执行任务与结果评分工具'], '按期限、范围与使用权授权'),
             ('AI数据', data, '按交付量或批次收费'),
             ('联合训练', ['训练实验、迁移验证与托管运行'], '项目或持续服务合同')]
    xs, ws = [0.6, 4.22, 9.11], [3.32, 4.59, 3.62]
    top, ch = 1.72, 2.8
    for i, ((k, what, fee), x, cw) in enumerate(zip(lines, xs, ws)):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 22, C['ink'], SERIF)
        label(sh, x + 0.3, top + 1.06, cw - 0.6, '交付内容')
        sh.t(x + 0.3, top + 1.3, cw - 0.6, 0.7, what, 12, b, line=1.1, gap=1.5)
        label(sh, x + 0.3, top + 2.22, cw - 0.6, '收费方式', C['accent'])
        sh.t(x + 0.3, top + 2.44, cw - 0.6, 0.3, fee, 12, C['ink'], SANS, True)
    sh.t(0.6, 4.64, W, 0.26, [(R('以交易为例：', 10, C['accent'], SANS, True),
                               R('一次回合留下订单、成交与计分记录，同一份记录既能训练，也能复测', 10, C['grey']))], 10)
    label(sh, 0.6, 4.98, 6, '合作路径')
    path = ['需求与验收', '付费试点', '授权交付', '持续更新']
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(path, pxs)):
        last = i == 3
        sh.rect(x, 5.22, pw, 0.52, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 5.22, pw - 0.5, 0.52, k, 16, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 5.22, 0.4, 0.52, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.9, [('每轮模型升级都要新任务与新测试样本：', F), ('持续更新带来复购', T)], 17)
    note(sh, '收费方式为拟议商业模式，不代表已签署合同。', 6.46)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_progress(sh, logos, polymarket, schools):
    """v39's progress and expert-network pages on one page: traction, who backs the research, who we can reach."""
    header(sh, '当前进展', 4, [('零外部融资，14天做出7款产品；', F), ('2家前沿实验室在谈，20万+专家可触达', T)])
    stats = [('2家', '前沿实验室在谈', '首个付费试点的来源'), ('14天', '完成首批7款产品开发', '交付速度'),
             ('512', '5个公开仓库的GitHub星标', '开发者与研究者的关注'), ('0', '外部融资', '每一分钱都花在产品上')]
    cw, xs = cols(4, 0.3)
    for i, ((v, k, m), x) in enumerate(zip(stats, xs)):
        acc = i == 0
        sh.rect(x, 1.72, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, 1.8, cw, 0.74, v, 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.56, cw, 0.3, k, 12.5, C['ink'], SANS, True)
        means(sh, x, 2.88, cw, m, 11)
    sh.t(0.6, 3.28, W, 0.22, '截至2026年9月29日；星标截至2026年10月2日', 8.5, C['grey'])
    label(sh, 0.6, 3.66, 8, '支持研发的从业者背景')
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 4.16, 1.25 if pic.shape_id == 403 else 1.0)
    if polymarket is not None:
        place(polymarket, slots[4], 4.16)
    sh.rule(0.6, 4.62, W, C['ink'])
    net = [('20万+', '可触达、可验证的专家网络', T), ('7,000+', '专家候补名单', F), ('21所', '高校学生与校友网络', F)]
    for i, (v, k, acc) in enumerate(net):
        y = 4.74 + i * 0.52
        sh.t(0.6, y, 1.6, 0.52, v, 26, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(2.25, y, 2.4, 0.52, k, 11.5, C['body'], anchor='ctr')
    gx, gw = 4.95, W + 0.6 - 4.95
    cell = gw / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, gx + (i % 6 + 0.5) * cell, 5.14 + (i // 6) * 0.68, min(1.0 / w, 0.46 / h, 1.0))
    note(sh, '标识仅表示支持者任职机构或网络覆盖的部分高校，不代表机构或学校背书；网络覆盖不等于已注册或参与交付。', 6.45, 8)
    footer(sh)


def p_edge(sh):
    """The team's edge: ships fast, and is already building for the next track (personal agents)."""
    header(sh, '团队优势', 5, [('交付快，转向快：', F), ('14天做出7个领域的产品，下一站个人Agent', T)])
    cards = [
        ('交付能力', '7个领域，14天', '交易、AI研究、事件预测、数学、推理、财务结账、软件工程',
         [('自有环境构建工具 + 训练基础设施', '单轮耗时降至1/4；同等资源多完成64%的已评分尝试'),
          ('20万+可触达专家网络', '7,000+已报名候补：示范、评分标准与验收都有人做'),
          ('独立评测后交付', '用未见任务验证，达到约定标准才交付')],
         '新环境从立项到交付，以周计', F),
        ('紧跟新赛道 · 开发中', '下一站：个人Agent', 'Meta Muse、OpenAI Dots已于2026年9月上线，Agent开始替人做事',
         [('个人Agent练习场', '在虚拟机与浏览器中复刻邮件、日历、文档等真实应用'),
          ('跨应用长任务 + 失败恢复', '多步骤任务按最终结果计分；注入弹窗、断网，测能否自救'),
          ('复用现有底座', '环境构建、计分器与专家网络直接沿用')],
         '下一条产品线：面向个人Agent开发团队', T),
    ]
    cw, xs = cols(2, 0.3)
    top, ch = 1.72, 3.78
    for (tag, big, sub, items, m, dark), x in zip(cards, xs):
        main, body, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.35, cw - 0.7
        sh.t(ix, top + 0.22, iw, 0.22, tag, 9.5, acc, MONO, True)
        sh.t(ix, top + 0.46, iw, 0.5, big, 26, acc if dark else C['ink'], SERIF)
        sh.t(ix, top + 1.02, iw, 0.26, sub, 10.5, body)
        sh.rect(ix, top + 1.38, iw, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(ix, top + 1.5, iw, 1.65, [(R(h, 12, main, SANS, True), BR(12), R(d, 11, body)) for h, d in items], 12, line=1.12, gap=7)
        means(sh, ix, top + 3.3, iw, m, 12, dark)
    kicker(sh, 5.72, [('每个新赛道都需要新的训练世界：', F), ('我们交付得快，也转向得快', T)], 19)
    sh.t(0.6, 6.32, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_problem, 4, {160}, 4),
    (p_why_now, 3, {126}, 5),
    (p_market, 13, {470}, 6),
    (p_customers, 20, {666}, 7),
    (p_solution, 6, {237}, 8),
    (p_products, 8, {300}, 9),
    (p_why_us, 15, {562}, 10),
    (p_competition, 14, {511}, 11),
    (p_business, 10, {376}, 12),
    (p_progress, 11, {401, 402, 403, 404, 407}, 13),
    (p_edge, 16, {607}, 14),
]
SCHOOLS = list(range(429, 441))                   # university logos on v17's network slide, copied onto progress
SCREENSHOT = os.path.join(ASSETS, 'repos-2026-10-02.png')
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_progress:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            schools = copy_pics(slides[12], s, SCHOOLS)
            build(sh, logos, poly, schools)
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_products:
            s.shapes.add_picture(SCREENSHOT, Inches(0.6), Inches(1.72), width=Inches(5.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
