"""Build the English SimReal deck v48-EN (14 slides) from v17's slides, in v48's layout.

English edition of v48: same order and design, copy in bp_v48_en_content.json, team bios as v42-EN.
No round amounts or valuation.

Usage: python3 bp_v48_en.py SimReal-BP-v17.pptx out.pptx
"""
import copy
import io
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

import bp_v17
import bp_v34
from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

_run = bp_v17.run
bp_v17.run = bp_v34.run = lambda *a, **k: _run(*a, **k).replace('lang="zh-CN"', 'lang="en-US"')

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['Who we are', 'Problem', 'Solution', 'Market', 'Strategy']
X = json.load(open(sys.argv[3] if len(sys.argv) > 3 else os.path.join(HERE, 'bp_v48_en_content.json'), encoding='utf-8'))
D = X['deck']
CONF = 'Confidential  ·  For invited investors only  ·  October 2026'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.4, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def cl(c):
    lead = c['lead'] if c['lead'].endswith(' ') else c['lead'] + ' '
    return [(lead, F), (c['accent'], T)]


def src_note(key, default=''):
    for n in X.get('notes', {}).get('sources', []):
        if key in n['page']:
            return n['text']
    return default


def p_cover(sh):
    c = D['cover']
    sh.t(0.6, 1.35, 8.6, 1.9, ['AI that improves itself', 'in the real world'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, c['tagline'], 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, c['subline'], 16, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    for (k, v), x in zip([('Company overview', 'October 2026'), ('Contact', 'business@simreal.co')], [0.6, 4.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    c = D['overview']
    header(sh, c['label'], 0, cl(c['claim']))
    cw, xs = cols(3)
    for i, cell in enumerate(c['cells']):
        acc = i in (1, 5)
        x, y = xs[i % 3], 1.86 + (i // 3) * 2.2
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, cell['k'], C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.6, cell['v'], 30, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.14, cw, 0.8, cell['lines'], 11.5, C['body'], line=1.15, gap=2)
    footer(sh)


def p_team(sh):
    header(sh, 'Team', 0, cl(D['team']['claim']))
    people = json.load(open(os.path.join(HERE, 'bp_v42_en_content.json'), encoding='utf-8'))['team']['people']
    for p in people:
        p['bullets'] = [b.replace('Helped launch', 'Helped organize') for b in p['bullets']]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 4.2
    for p, x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.18, cw - 0.6, 0.5, [para([R(p['name'], 26, C['ink'], SERIF), R('   ' + p['role'], 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.82, cw - 0.6, 2.4, p['bullets'], 10.5, C['body'], line=1.1, gap=4)
        sh.rule(x + 0.3, y0 + 3.3, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.4, cw - 0.6, 0.7, p['edu'], 9, C['grey'], line=1.15, gap=1.5)
    footer(sh)


def p_speed(sh):
    c = D['speed']
    header(sh, c['label'], 0, cl(c['claim']))
    cw, xs = cols(3)
    for i, (m, x) in enumerate(zip(c['metrics'], xs)):
        acc = i == 0
        sh.rect(x, 1.72, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, 1.84, cw, m['tag'], C['accent'] if acc else None)
        sh.t(x, 2.08, cw, 0.76, m['v'], 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.88, cw, 0.3, m['k'], 13, C['ink'], SANS, True)
        means(sh, x, 3.22, cw, m['means'], 11.5)
    sh.rect(0.6, 3.82, W, 0.92, C['tint'])
    pw, pxs = cols(3, 0.35, 0.85, W - 0.5)
    for p, x in zip(c['proof'], pxs):
        sh.t(x, 3.82, pw, 0.92, [(R(p['head'], 13, C['ink'], SANS, True),), (R(p['text'], 11, C['body']),)], 13, anchor='ctr', line=1.1, gap=3)
    sh.rect(0.6, 4.98, W, 1.18, C['ink'])
    sh.t(0.85, 4.98, 2.0, 1.18, c['next_head'], 16, C['accentLt'], SERIF, anchor='ctr')
    nw, nxs = cols(3, 0.3, 3.0, W - 2.65)
    for n, x in zip(c['next'], nxs):
        sh.t(x, 5.16, nw, 0.22, n['when'], 9.5, C['accentLt'], MONO, True)
        sh.t(x, 5.42, nw, 0.66, n['what'], 13.5, C['onDarkHi'], SERIF, line=1.1)
    note(sh, src_note('04'), 6.38)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_progress(sh, logos, polymarket, schools):
    c = D['traction']
    header(sh, c['label'], 0, cl(c['claim']))
    cw, xs = cols(4, 0.3)
    for i, (st, x) in enumerate(zip(c['stats'], xs)):
        acc = i == 0
        sh.rect(x, 1.72, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, 1.8, cw, 0.74, st['v'], 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.56, cw, 0.3, st['k'], 12.5, C['ink'], SANS, True)
        if st['means']:
            means(sh, x, 2.88, cw, st['means'], 10.5)
    sh.t(0.6, 3.28, W, 0.22, 'As of September 29, 2026; GitHub stars as of October 2, 2026', 8.5, C['grey'])
    label(sh, 0.6, 3.66, 8, c['backers_label'])
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 4.16, 1.25 if pic.shape_id == 403 else 1.0)
    if polymarket is not None:
        place(polymarket, slots[4], 4.16)
    sh.rule(0.6, 4.62, W, C['ink'])
    for i, n in enumerate(c['network']):
        y = 4.74 + i * 0.52
        sh.t(0.6, y, 1.6, 0.52, n['v'], 26, C['accent'] if i == 0 else C['ink'], SERIF, anchor='ctr')
        sh.t(2.25, y, 2.4, 0.52, n['k'], 11.5, C['body'], anchor='ctr')
    gx, gw = 4.95, W + 0.6 - 4.95
    cell = gw / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, gx + (i % 6 + 0.5) * cell, 5.14 + (i // 6) * 0.68, min(1.0 / w, 0.46 / h, 1.0))
    note(sh, 'Logos show where supporters work and some universities in our network; they do not imply endorsement. Network reach is not registration or delivery.', 6.45, 8)
    footer(sh)


def p_problem(sh):
    c = D['problem']
    header(sh, c['label'], 1, cl(c['claim']))
    y0, rh = 2.2, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * 4, C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, 'Looks like', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, 'Actually', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, r['domain'], 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, r['looks'], 18, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, r['reality'], 18, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    ev = c['evidence']
    sh.t(0.6, 5.0, 1.8, 0.6, ev['v'], 26, C['accent'], SERIF, anchor='ctr')
    sh.t(2.45, 5.0, 9, 0.6, ev['text'], 12, C['body'], anchor='ctr')
    kicker(sh, 5.9, cl(c['kicker']), 20)
    footer(sh)


def p_why_now(sh):
    c = D['why_now']
    header(sh, c['label'], 1, cl(c['claim']))
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for i, (card, x) in enumerate(zip(c['cards'], xs)):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, card['title'], 22, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.1, cw - 0.6, 1.2, card['evidence'], 12, C['body'], line=1.15, gap=5)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, card['means'], 12.5)
    w = c['window']
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, w['head'], 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R(w['text'], 12.5, C['onDark']), R(w['accent'], 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr', line=1.1)
    note(sh, 'Sources: Epoch AI (public-text projection; An FAQ on RL environments, Jan 2026); SiliconANGLE, TechCrunch (Jul 2026); '
             'Meta, OpenAI launches (Sep 2026); Wing VC (2026). Muse and Dots are industry examples, not SimReal customers.', 6.3)
    footer(sh)


def p_solution(sh):
    c = D['solution']
    header(sh, c['label'], 2, cl(c['claim']))
    cw, xs = cols(4, 0.4)
    top, bh = 1.72, 1.35
    for i, (st, x) in enumerate(zip(c['steps'], xs)):
        acc = i == 2
        sh.rect(x, top, cw, bh, C['accent'] if acc else C['tint'])
        sh.t(x + 0.25, top + 0.14, 1, 0.2, f'0{i + 1}', 9, C['onDarkHi'] if acc else C['accent'], MONO)
        sh.t(x + 0.25, top + 0.36, cw - 0.5, 0.42, st['name'], 21, C['onDarkHi'] if acc else C['ink'], SERIF)
        sh.t(x + 0.25, top + 0.86, cw - 0.5, 0.48, st['lines'], 11, C['onDarkHi'] if acc else C['body'], line=1.1, gap=0)
        if i < 3:
            sh.t(x + cw, top, 0.4, bh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    cx = xs[2] + cw / 2
    sh.rect(cx, top + bh, 0.012, 0.3, C['mid'])
    sh.t(cx + 0.08, top + bh + 0.02, 2.6, 0.26, c['connector'], 9, C['grey'], MONO, anchor='ctr')
    lt = 3.42
    sh.rect(0.6, lt, W, 1.25, C['tint'])
    sh.t(0.85, lt + 0.16, 4, 0.22, c['loop_title'], 9.5, C['accent'], MONO, True)
    lw, lxs = cols(4, 0.4, 0.85, W - 0.5)
    for i, (st, x) in enumerate(zip(c['loop'], lxs)):
        sh.t(x, lt + 0.46, lw, 0.34, st['name'], 16, C['accent'] if i == 3 else C['ink'], SERIF)
        sh.t(x, lt + 0.82, lw, 0.28, st['desc'], 11, C['body'])
        if i < 3:
            sh.t(x + lw, lt + 0.46, 0.4, 0.34, '→', 12, C['grey'], algn='ctr', anchor='ctr')
    sh.t(0.6, 4.78, W, 0.26, c['note'], 10, C['grey'])
    sh.rect(0.6, 5.2, W, 0.66, C['ink'])
    sh.t(0.85, 5.2, 2.2, 0.66, 'We deliver', 14, C['accentLt'], SERIF, anchor='ctr')
    sh.t(3.0, 5.2, W - 2.6, 0.66, '  ·  '.join(c['deliver']), 15, C['onDarkHi'], SERIF, anchor='ctr')
    note(sh, 'References: InstructGPT (arXiv 2203.02155); DeepSeek-R1 (arXiv 2501.12948); Absolute Zero (arXiv 2505.03335).', 6.2)
    footer(sh)


def p_products(sh):
    c = D['product']
    header(sh, c['label'], 2, cl(c['claim']))
    x0, y0, rh = 6.55, 1.72, 0.54
    sh.rule(x0, y0, W + 0.6 - x0, C['ink'])
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.t(x0, y + 0.04, 6.2, 0.26, r['name'], 13.5, C['ink'], SERIF)
        sh.t(x0, y + 0.28, 6.2, 0.24, r['desc'], 10.5, C['body'])
        sh.rule(x0, y + rh, W + 0.6 - x0)
    label(sh, 0.6, 4.78, 8, 'Training infrastructure', C['accent'])
    cw, xs = cols(3)
    for m, x in zip(c['infra'], xs):
        sh.rect(x, 5.08, cw, 0.02, C['ink'])
        sh.t(x, 5.16, 1.55, 0.62, m['v'], 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.6, 5.16, cw - 1.6, 0.62, [m['k'], (R(m['how'], 9.5, C['grey']),)], 11.5, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=0)
    note(sh, 'Public repos as of Oct 2, 2026. Infra metrics are internal tests with different baselines; they can’t be multiplied or added.', 6.3)
    footer(sh)


def p_agent(sh):
    header(sh, 'Personal agents', 2, [('Personal agents have launched; ', F), ('our practice environment is in development', T)])
    sh.t(0.6, 1.62, W, 0.3, 'Meta Muse and OpenAI Dots launched in Sep 2026 to do work for people. Agents need practice before they act.', 12, C['grey'])
    cards = [('Real app replicas', ['Email, calendar, and docs', 'in VMs and browsers']),
             ('Long cross-app tasks', ['Multi-step tasks', 'scored on the final result']),
             ('Failure recovery', ['Pop-ups, dropped connections,', 'expired logins: can it recover?'])]
    cw, xs = cols(3, 0.3)
    top, ch = 2.2, 2.55
    for i, ((k, lines), x) in enumerate(zip(cards, xs)):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.32, top + 0.24, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.32, top + 0.52, cw - 0.64, 0.5, k, 24, C['ink'], SERIF)
        sh.t(x + 0.32, top + 1.28, cw - 0.64, 0.9, lines, 14, C['body'], line=1.15, gap=2)
    sh.rect(0.6, 5.0, W, 0.86, C['ink'])
    sh.t(0.85, 5.0, 1.6, 0.86, 'Who pays', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.5, 5.0, W - 2.15, 0.86, [(R('Personal-agent teams pay for practice; ', 13, C['onDark']), R('labs buy the trajectories it produces', 13, C['onDarkHi'], SANS, True))],
         13, anchor='ctr', line=1.1)
    note(sh, 'In development; features are planned. Muse and Dots are industry examples, not SimReal customers.', 6.2)
    footer(sh)


def p_market(sh):
    c = D['market']
    header(sh, c['label'], 3, cl(c['claim']))
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    for side, x, acc in [(c['now'], 0.9, False), (c['future'], 6.9, True)]:
        label(sh, x, 1.97, 5.5, side['label'], C['accent'] if acc else None)
        sh.t(x, 2.25, 5.6, 0.72, side['value'], 38, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 3.05, 5.6, 0.3, side['sub'], 11, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    cw, xs = cols(3)
    for m, x in zip(c['comps'], xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, m['v'], 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, m['d'], 11.5, C['body'], line=1.1)
        means(sh, x, 5.34, cw, m['means'], 12)
    sh.t(0.6, 5.86, W, 0.3, [(R('Comparable  ', 10, C['accent'], MONO, True), R(c['comparable'], 12, C['ink']))], 12, anchor='ctr')
    note(sh, 'Sources: Mercor (TechCrunch, Sacra, Dealroom); Snorkel AI announcement (Sep 2026); Deedy Das industry map (Jul 2026); UniPat (Bloomberg, Sep 2026). 2030 is a company scenario, not an independent forecast.', 6.36)
    footer(sh)


def p_business(sh):
    c = D['business']
    header(sh, c['label'], 3, cl(c['claim']))
    cw, xs = cols(3, 0.3)
    label(sh, 0.6, 1.66, 4, 'Who pays')
    for i, (sg, x) in enumerate(zip(c['segments'], xs)):
        now = i == 0
        sh.rect(x, 1.92, cw, 1.05, C['ink'] if now else C['tint'])
        sh.t(x + 0.28, 2.02, cw - 0.5, 0.2, sg['stage'], 9, C['accentLt'] if now else C['accent'], MONO, True)
        sh.t(x + 0.28, 2.24, cw - 0.5, 0.36, sg['name'], 17, C['onDarkHi'] if now else C['ink'], SERIF)
        sh.t(x + 0.28, 2.62, cw - 0.5, 0.3, sg['line'], 11, C['onDark'] if now else C['body'])
    label(sh, 0.6, 3.12, 4, 'What they buy')
    for ln, x in zip(c['lines'], xs):
        sh.rect(x, 3.38, cw, 1.42, C['tint'])
        sh.t(x + 0.28, 3.5, cw - 0.5, 0.4, ln['name'], 19, C['ink'], SERIF)
        sh.t(x + 0.28, 3.94, cw - 0.5, 0.3, ln['what'], 11, C['body'])
        sh.t(x + 0.28, 4.32, cw - 0.5, 0.36, [(R('Pricing  ', 9, C['accent'], MONO), R(ln['fee'], 11.5, C['ink'], SANS, True))], 11.5)
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(c['path'], pxs)):
        last = i == 3
        sh.rect(x, 5.0, pw, 0.5, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 5.0, pw - 0.5, 0.5, k, 15, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 5.0, 0.4, 0.5, '→', 13, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.7, cl(c['kicker']), 18)
    note(sh, 'Pricing reflects the proposed business model, not signed contracts.', 6.38)
    footer(sh)


def p_competition(sh):
    c = D['competition']
    header(sh, c['label'], 4, cl(c['claim']))
    xs, ws = [0.6, 3.75, 6.55, 9.45], [3.0, 2.65, 2.75, 3.28]
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 0.98, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, c['heads'])):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    y0, rh = 2.04, 0.98
    for i, r in enumerate(c['rows']):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.85, rh, [para([R(r['alt'], 16, C['ink'], SERIF)])] + ([para([R(r['names'], 8.5, C['grey'], MONO)], before=3)] if r['names'] else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.15, rh, r['pro'], 11, C['body'], anchor='ctr', line=1.1, gap=2)
        sh.t(xs[2], y, ws[2] - 0.15, rh, r['con'], 11, C['body'], anchor='ctr', line=1.1, gap=2)
        sh.t(xs[3], y, ws[3] - 0.1, rh, r['ours'], 11.5, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=2)
    sh.rule(0.6, y0 + 3 * rh, W)
    g = c['gap']
    sh.rect(0.6, 5.22, W, 0.86, C['ink'])
    sh.t(0.85, 5.22, 1.8, 0.86, g['head'], 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.75, 5.22, W - 2.4, 0.86, [(R(g['text'], 12.5, C['onDark']),), (R(g['accent'], 13, C['onDarkHi'], SANS, True),)], 12.5, anchor='ctr', line=1.1, gap=3)
    note(sh, src_note('13'), 6.3)
    footer(sh)


def p_strategy(sh):
    c = D['strategy']
    header(sh, c['label'], 4, cl(c['claim']))
    sh.t(0.6, 1.6, W, 0.3, [(R('Use of funds  ', 10, C['accent'], MONO, True), R(c['split'], 12.5, C['ink'], SANS, True))], 12.5, anchor='ctr')
    cw, xs = cols(3, 0.3)
    top, ch = 2.02, 4.08
    for i, (p, x) in enumerate(zip(c['pillars'], xs)):
        dark = i == 2
        main, body, acc, grey = ((C['onDarkHi'], C['onDark'], C['accentLt'], C['onDark']) if dark else (C['ink'], C['body'], C['accent'], C['grey']))
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.28, cw - 0.56
        sh.t(ix, top + 0.16, iw, 0.44, p['name'], 22, acc if dark else C['ink'], SERIF)
        sh.t(ix, top + 0.64, iw, 0.24, 'Benchmark  ' + p['bench'], 8.5, grey, line=1.05)
        sh.t(ix, top + 1.02, iw, 0.2, 'Now', 9, grey, MONO)
        sh.t(ix, top + 1.22, iw, 0.42, p['now'], 10.5, body, line=1.08)
        sh.t(ix, top + 1.76, iw, 0.2, 'Investment', 9, acc, MONO, True)
        sh.t(ix, top + 1.96, iw, 0.8, p['invest'], 10, main, line=1.08, gap=2)
        sh.rect(ix, top + 2.76, iw, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(ix, top + 2.9, iw, 0.2, 'Next', 9, acc, MONO)
        sh.t(ix, top + 3.12, iw, 0.7, p['m6'], 12, main, SANS, True, line=1.1)
    kicker(sh, 6.18, cl(c['kicker']), 18)
    note(sh, src_note('14'), 6.66, 7.5)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_speed, 9, {336}, 4),
    (p_progress, 11, {401, 402, 403, 404, 407}, 5),
    (p_problem, 4, {160}, 6),
    (p_why_now, 3, {126}, 7),
    (p_solution, 6, {237}, 8),
    (p_products, 8, {300}, 9),
    (p_agent, 15, {562}, 10),
    (p_market, 13, {470}, 11),
    (p_business, 10, {376}, 12),
    (p_competition, 14, {511}, 13),
    (p_strategy, 16, {607}, 14),
]
SCHOOLS = list(range(429, 441))                   # university logos on v17's network slide, copied onto progress
SCREENSHOT = os.path.join(ASSETS, 'repos-2026-10-02.png')
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_progress:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            schools = copy_pics(slides[12], s, SCHOOLS)
            build(sh, logos, poly, schools)
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_products:
            s.shapes.add_picture(SCREENSHOT, Inches(0.6), Inches(1.72), width=Inches(5.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
