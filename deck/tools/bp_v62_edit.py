"""Make SimReal BP v62 from v61: new product pictures, and the competition column that ran into the next one.

Page 9: the three hand-drawn mock panels (6.5-7.5pt text) and their one-liners give way to one picture per card
from visuals_v62.py (what the AI does, how it is scored, what came out); tag, name and stars stay editable.
Page 14: the 弱在哪 text boxes were wider than their column and ran under the tinted 衍真 column; they are narrowed
to the column and the longest line is shortened so it stays on one line.

Usage: python3 visuals_v62.py; python3 bp_v62_edit.py SimReal-BP-v61.pptx out.pptx
"""
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets', 'v62')
CARDS = [(0.60, 'xitadel.png'), (4.70, 'mlbench.png'), (8.80, 'forecast.png')]


def inch(v):
    return v / 914400


def products(slide):
    for shp in list(slide.shapes):
        if 2.6 <= inch(shp.top) < 6.0:
            shp._element.getparent().remove(shp._element)
        elif shp.has_text_frame and shp.text_frame.text.startswith('界面为示意图'):
            run = shp.text_frame.paragraphs[0].runs[0]
            run.text = 'Xitadel分数来自其GitHub公开排行榜（人类最佳=80，满分100）；事件预测为示例。'
            for r in shp.text_frame.paragraphs[0].runs[1:]:
                r.text = ''
            shp.width = Inches(10)
    for x, img in CARDS:
        slide.shapes.add_picture(os.path.join(ASSETS, img), Inches(x + 0.2), Inches(2.62), width=Inches(3.52))


def competition(slide):
    for shp in slide.shapes:
        if shp.name in ('Text 16', 'Text 28', 'Text 33'):
            shp.width = Inches(2.2)
        if shp.name == 'Text 33':
            shp.text_frame.paragraphs[0].runs[0].text = '缺规则缺专家，占研究人力'


def main(src, dst):
    prs = Presentation(src)
    products(prs.slides[8])
    competition(prs.slides[13])
    prs.save(dst)
    print('wrote', dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
