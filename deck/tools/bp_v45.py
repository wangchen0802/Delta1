"""Build SimReal (衍真) BP v45 (13 slides, Chinese, no round terms) from v17's slides.

v44 after a line-by-line edit for tone and VC reading: fewer claims a VC would question (no "客户是",
no 82x headline, no "估值参照"), no mirrored phrasing, one name per thing.


v43 rewritten for how a VC reads: what it is, who, traction, then problem, timing, product, moat, market,
competition, model, and the team's edge last. Team on page 3. Plain short sentences and numbers; no
slogans. Facts as v41.

Usage: python3 bp_v45.py SimReal-BP-v17.pptx out.pptx
       then embed_cjk.py adds the Chinese font subsets to the shipping PPTX.
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['我们是谁', '问题与时机', '方案', '市场与竞争', '商业模式', '团队优势']
CONF = '机密  ·  仅供受邀投资机构内部评估使用  ·  2026年10月'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.36, 6.98, 0.5, 0.26, '衍真', 9, C['ink'], SANS, True, anchor='ctr')
    sh.t(1.85, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def contact(sh, y):
    sh.t(0.6, y, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)


def p_cover(sh):
    sh.t(1.62, 0.55, 1.0, 0.25, '衍真', 13, C['ink'], SANS, True, anchor='ctr')
    sh.t(0.6, 1.35, 8.6, 1.9, ['让AI在真实世界里', '自我进化'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, '强化学习训练环境与AI数据', 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, '面向AI实验室、企业和个人Agent团队', 16, C['accent'])
    sh.rule(0.6, 4.95, 7.6)
    for (k, v), x in zip([('商业计划书', '2026年10月'), ('联系', 'business@simreal.co')], [0.6, 4.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, '概述', 0, [('给AI实验室做强化学习训练环境，', F), ('从交易做起', T)])
    cells = [
        ('做什么', '训练环境与数据', ['交易、AI研究、事件预测已上线', '卖环境授权、Agent轨迹、专家和评测数据'], F),
        ('结果', '+12%', ['Qwen3.8-27B在Xitadel训练后', '未见过的交易日上最高提升12%，多次独立复现'], T),
        ('速度', '14天7款产品', ['5个公开仓库，512星', '零外部融资'], F),
        ('团队', '05后量化团队', ['剑桥、LSE、杜克数学', 'Jane Street、Citadel、Optiver、Millennium'], F),
        ('市场', '85亿美元', ['训练数据与RL环境供应商年收入', 'Mercor年化毛营收16个月涨27倍'], F),
        ('下一步', '个人Agent', ['练习环境开发中'], T),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.86 + (i // 3) * 2.2
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.6, v, 30, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.14, cw, 0.8, d, 11.5, C['body'], line=1.15, gap=2)
    footer(sh)


def p_team(sh):
    header(sh, '团队', 0, [('三位05后创始人，都在量化机构工作或实习过，', F), ('放弃转正offer全职做衍真', T)])
    b, sz = C['body'], 11
    people = [
        ('Charles', 'CEO', [(R('United Stables首位员工：U稳定币', sz, b), BR(sz), R('一年从0到14亿美元，一个月上线Binance', sz, b)),
                            '负责机构关系，参与SIG、DRW等合作', (R('汇丰港元稳定币项目唯一实习生，', sz, b), BR(sz), R('参与HKMA合规', sz, b)),
                            '伦敦Citadel对冲基金实习', 'X博主，内容数百万浏览'],
         ['2005年生 · LSE数学 · 深国交', '美国数学奥林匹克（USAMO）入围']),
        ('Henry', 'CTO', [(R('剑桥机器学习暑研，导师Po-Ling Loh', sz, b), BR(sz), R('（国际数理统计学会会士）', sz, b)),
                          '剑桥研究中心最年轻的本科AI研究员', 'Jane Street、Citadel、Optiver量化经历', '设计了5个基准和RL环境'],
         ['2005年生 · 剑桥数学一等荣誉（奖学金）· 深国交', '剑桥数学竞赛全球前30 · 英国物理竞赛超级金奖']),
        ('Amaris', 'COO', ['Millennium香港数据科学家', 'Millennium另类数据团队首位应届招聘',
                           (R('参与筹办Plug and Play香港首场活动', sz, b), BR(sz), R('（联合香港科技园，200+人）', sz, b)),
                           '强生MedTech科技峰会主持人（200+人）', '17岁出版作家，原创内容10万+互动'],
         ['2005年生 · 杜克数学与统计 · 上海包玉刚']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 4.2
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.18, cw - 0.6, 0.5, [para([R(name, 26, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.82, cw - 0.6, 2.4, lines, sz, b, line=1.12, gap=5)
        sh.rule(x + 0.3, y0 + 3.3, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.4, cw - 0.6, 0.7, edu, 9.5, C['grey'], line=1.15, gap=1.5)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_progress(sh, logos, polymarket, schools):
    header(sh, '进展', 0, [('14天上线7款产品，', F), ('2家前沿实验室在谈', T)])
    stats = [('2家', '前沿实验室在谈', '争取首个付费试点'), ('14天', '上线7款产品', '交易、AI研究、事件预测已上线'),
             ('512', 'GitHub星标', '5个公开仓库'), ('0', '外部融资', '')]
    cw, xs = cols(4, 0.3)
    for i, ((v, k, m), x) in enumerate(zip(stats, xs)):
        acc = i == 0
        sh.rect(x, 1.72, cw, 0.02, C['accent'] if acc else C['ink'])
        sh.t(x, 1.8, cw, 0.74, v, 40, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, 2.56, cw, 0.3, k, 12.5, C['ink'], SANS, True)
        if m:
            means(sh, x, 2.88, cw, m, 11)
    sh.t(0.6, 3.28, W, 0.22, '截至2026年9月29日；星标截至2026年10月2日', 8.5, C['grey'])
    label(sh, 0.6, 3.66, 8, '支持研发的从业者来自')
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 4.16, 1.25 if pic.shape_id == 403 else 1.0)
    if polymarket is not None:
        place(polymarket, slots[4], 4.16)
    sh.rule(0.6, 4.62, W, C['ink'])
    net = [('20万+', '可触达、可验证的专家', T), ('7,000+', '专家候补', F), ('21所', '高校的学生和校友', F)]
    for i, (v, k, acc) in enumerate(net):
        y = 4.74 + i * 0.52
        sh.t(0.6, y, 1.6, 0.52, v, 26, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(2.25, y, 2.4, 0.52, k, 11.5, C['body'], anchor='ctr')
    gx, gw = 4.95, W + 0.6 - 4.95
    cell = gw / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, gx + (i % 6 + 0.5) * cell, 5.14 + (i // 6) * 0.68, min(1.0 / w, 0.46 / h, 1.0))
    note(sh, '标识仅表示支持者任职机构或网络覆盖的部分高校，不代表机构或学校背书；网络覆盖不等于已注册或参与交付。', 6.45, 8)
    footer(sh)


def p_problem(sh):
    header(sh, '问题', 1, [('模型会推理，', F), ('做不好真实世界里的事', T)])
    rows = [('交易', '回测一路上涨', '没见过的行情里亏钱'), ('软件工程', '测试全部通过', '上线出错'),
            ('财务', '账看着做完了', '月结对不平'), ('事件预测', '分析头头是道', '结果押错')]
    y0, rh = 2.2, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, '看起来', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, '实际', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 18, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 18, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    sh.t(0.6, 5.0, 1.8, 0.6, '6/32', 26, C['accent'], SERIF, anchor='ctr')
    sh.t(2.45, 5.0, 9, 0.6, '前沿模型真钱交易32轮，6轮赚钱（Alpha Arena）', 12, C['body'], anchor='ctr')
    kicker(sh, 5.9, [('缺少', F), ('按真实结果打分的训练环境', T), ('。', F)], 20)
    footer(sh)


def p_why_now(sh):
    header(sh, '为什么是现在', 1, [('实验室开始花钱买训练环境，', F), ('专业领域供给不足', T)])
    items = [('01', '训练方式变了', ['推理模型靠可自动验证的结果做强化学习', '公开人类文本预计2026–2032年用尽'],
              '能自动打分的环境成了稀缺品'),
             ('02', '钱已经在花', ['每个RL任务200–2,000美元', '环境合同每季度六到七位数美元', 'Mercor 2026年7月收购环境公司Deeptune'],
              '环境成了实验室的采购项'),
             ('03', '个人Agent上线', ['Meta Muse、OpenAI Dots', '2026年9月上线，开始替人做事'],
              '下一批客户出现')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for (n, k, ev, m), x in zip(items, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 22, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.1, cw - 0.6, 1.2, ev, 12, C['body'], line=1.15, gap=5)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, m, 12.5)
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, '窗口期', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R('约20家早期RL环境公司，预计收敛到3–5家（Wing VC）。', 12.5, C['onDark']),
                                       R('我们从交易切入，对错可以自动验证', 12.5, C['onDarkHi'], SANS, True))],
         12.5, anchor='ctr', line=1.1)
    note(sh, '来源：Epoch AI（公开文本存量预测；《An FAQ on RL environments》，2026年1月）；SiliconANGLE、TechCrunch（2026年7月）；'
             'Meta、OpenAI发布（2026年9月）；Wing VC（2026年）。Muse、Dots为行业示例，不是衍真客户。', 6.3)
    footer(sh)


def p_solution(sh):
    header(sh, '方案', 2, [('按任务结果训练，', F), ('在没见过的任务上验收', T)])
    steps = [('接入', ['客户已有的模型', '和工具调用'], F), ('适配', ['监督微调', '人类反馈'], F),
             ('训练', ['按结果给奖励', '强化学习更新模型'], T), ('验收', ['没见过的任务', '达标才交付'], F)]
    cw, xs = cols(4, 0.4)
    top, bh = 1.72, 1.35
    for i, ((k, lines, acc), x) in enumerate(zip(steps, xs)):
        sh.rect(x, top, cw, bh, C['accent'] if acc else C['tint'])
        sh.t(x + 0.25, top + 0.14, 1, 0.2, f'0{i + 1}', 9, C['onDarkHi'] if acc else C['accent'], MONO)
        sh.t(x + 0.25, top + 0.36, cw - 0.5, 0.42, k, 21, C['onDarkHi'] if acc else C['ink'], SERIF)
        sh.t(x + 0.25, top + 0.86, cw - 0.5, 0.48, lines, 11, C['onDarkHi'] if acc else C['body'], line=1.1, gap=0)
        if i < 3:
            sh.t(x + cw, top, 0.4, bh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    cx = xs[2] + cw / 2
    sh.rect(cx, top + bh, 0.012, 0.3, C['mid'])
    sh.t(cx + 0.08, top + bh + 0.02, 2.4, 0.26, '训练暴露的薄弱任务', 9, C['grey'], MONO, anchor='ctr')
    lt = 3.42
    sh.rect(0.6, lt, W, 1.25, C['tint'])
    sh.t(0.85, lt + 0.16, 4, 0.22, '自动出题 · 研发中', 9.5, C['accent'], MONO, True)
    loop = [('找弱点', '分析训练失败'), ('出新题', '智能体生成或挑选'), ('校验', '规则、执行器、专家'), ('回到训练', '通过校验进入下一轮')]
    lw, lxs = cols(4, 0.4, 0.85, W - 0.5)
    for i, ((k, d), x) in enumerate(zip(loop, lxs)):
        sh.t(x, lt + 0.46, lw, 0.34, k, 16, C['accent'] if i == 3 else C['ink'], SERIF)
        sh.t(x, lt + 0.82, lw, 0.28, d, 11, C['body'])
        if i < 3:
            sh.t(x + lw, lt + 0.46, 0.4, 0.34, '→', 12, C['grey'], algn='ctr', anchor='ctr')
    sh.t(0.6, 4.78, W, 0.26, '验收任务不参与训练和出题', 10, C['grey'])
    sh.rect(0.6, 5.2, W, 0.66, C['ink'])
    sh.t(0.85, 5.2, 2.2, 0.66, '交付物', 14, C['accentLt'], SERIF, anchor='ctr')
    sh.t(3.0, 5.2, W - 2.6, 0.66, '训练环境  ·  专业数据  ·  评分工具  ·  评测报告', 15, C['onDarkHi'], SERIF, anchor='ctr')
    note(sh, '参考：InstructGPT（arXiv 2203.02155）；DeepSeek-R1（arXiv 2501.12948）；Absolute Zero（arXiv 2505.03335）。', 6.2)
    footer(sh)


def p_products(sh):
    header(sh, '产品', 2, [('5个基准已开源，', F), ('同等资源下已评分尝试+64%', T)])
    rows = [('SimReal-MLBench', '60个真实竞赛任务，测机器学习研究能力'),
            ('Xitadel-QuantBench', '回放订单簿，和同资产上的专业交易员比；保留轨迹'),
            ('MathmoBench', '高阶数学题，测推理和解题'),
            ('Puzzle Benchmark', '749道人工推理题，15类主题'),
            ('FuturePredict Bench', '截止前锁定概率和证据，按真实结局打分')]
    x0, y0, rh = 6.55, 1.72, 0.54
    sh.rule(x0, y0, W + 0.6 - x0, C['ink'])
    for i, (k, d) in enumerate(rows):
        y = y0 + i * rh
        sh.t(x0, y + 0.04, 6.2, 0.26, k, 13.5, C['ink'], SERIF)
        sh.t(x0, y + 0.28, 6.2, 0.24, d, 10.5, C['body'])
        sh.rule(x0, y + rh, W + 0.6 - x0)
    label(sh, 0.6, 4.78, 8, '训练基础设施', C['accent'])
    infra = [('1/4', '单轮耗时', 'microVM隔离'), ('+64%', '同等资源的已评分尝试', '任务资源分配'),
             ('−1/3', '存档耗时，存档次数减半', '智能存档')]
    cw, xs = cols(3)
    for (v, k, how), x in zip(infra, xs):
        sh.rect(x, 5.08, cw, 0.02, C['ink'])
        sh.t(x, 5.16, 1.55, 0.62, v, 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x + 1.6, 5.16, cw - 1.6, 0.62, [k, (R(how, 9.5, C['grey']),)], 11.5, C['ink'], SANS, True, anchor='ctr', line=1.1, gap=0)
    note(sh, '公开仓库截至2026年10月2日。三项指标为公司内部测试，比较口径不同，不能相乘或相加。', 6.3)
    footer(sh)


def p_moat(sh):
    header(sh, '技术壁垒', 2, [('计分准，', F), ('训练有效', T)])
    cards = [('计分', ['交易台规则：逐笔撮合，盈亏扣回撤再乘夏普', '最后一个交易日留出，结果哈希提前公开'],
              ['GPT 6得77分，其他前沿模型24–30分', '人类最佳80分，还没有模型超过']),
             ('训练', ['同一环境既评测又训练', 'microVM隔离、任务资源分配、智能存档'],
              ['Qwen3.8-27B最高提升12%', '未见过的交易日，多次独立复现'])]
    cw, xs = cols(2, 0.3)
    top, ch = 1.72, 3.35
    for (k, how, res), x in zip(cards, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.35, top + 0.24, cw - 0.7, 0.5, k, 26, C['ink'], SERIF)
        sh.t(x + 0.35, top + 0.92, cw - 0.7, 0.2, '做法', 9, C['grey'], MONO)
        sh.t(x + 0.35, top + 1.16, cw - 0.7, 0.8, how, 13, C['body'], line=1.15, gap=4)
        sh.rule(x + 0.35, top + 2.0, cw - 0.7, C['mid'])
        sh.t(x + 0.35, top + 2.14, cw - 0.7, 0.2, '结果', 9, C['accent'], MONO)
        sh.t(x + 0.35, top + 2.38, cw - 0.7, 0.85, res, 15.5, C['ink'], SANS, True, line=1.15, gap=4)
    sh.rect(0.6, 5.32, W, 0.72, C['ink'])
    sh.t(0.85, 5.32, 1.6, 0.72, '积累', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.5, 5.32, W - 2.15, 0.72, [(R('任务库、评分器、失败案例和验证过的专家随交付积累，', 12.5, C['onDark']),
                                      R('新客户直接复用', 12.5, C['onDarkHi'], SANS, True))], 12.5, anchor='ctr', line=1.1)
    note(sh, '来源：Xitadel-QuantBench公开报告与计分规范；训练基础设施指标为公司内部测试。', 6.3)
    footer(sh)


def p_market(sh):
    header(sh, '市场', 3, [('今天85亿美元，', F), ('2030年公司情景测算7,000亿美元', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, '今天：训练数据与RL环境')
    sh.t(0.9, 2.25, 5, 0.72, '85亿美元 / 年', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, '50余家供应商收入估算（Deedy Das，2026年7月）', 11, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030年：公司情景测算', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '7,000亿美元 / 年', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, '假设AI经济7万亿美元，训练投入占10%', 11, C['body'])
    growth = [('27倍', 'Mercor年化毛营收，16个月从7,500万到20亿美元', '专家数据需求在涨'),
              ('10倍', 'Mercor估值，17个月从20亿到200亿美元（洽谈中）', '资本在加注'),
              ('18倍', 'Snorkel AI数据服务近一年增长，年化收入3.75亿美元', '需求转向专家和环境')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 11.5, C['body'], line=1.1)
        means(sh, x, 5.34, cw, m, 12)
    sh.t(0.6, 5.86, W, 0.3, [(R('可比交易  ', 10, C['accent'], MONO, True), R('据彭博，阿里拟领投UniPat 3亿美元，估值约25亿美元', 12, C['ink']))], 12, anchor='ctr')
    note(sh, '来源：Mercor（TechCrunch、Sacra、Dealroom）；Snorkel AI公告（2026年9月）；Deedy Das行业图谱（2026年7月）；UniPat（彭博，2026年9月）。2030年为公司情景测算，非独立研究预测。', 6.36)
    footer(sh)


def p_competition(sh):
    header(sh, '竞争', 3, [('规模不如对手，', F), ('专业领域的计分和速度占优', T)])
    xs, ws = [0.6, 3.75, 6.55, 9.45], [3.0, 2.65, 2.75, 3.28]
    heads = ['客户现有选择', '强项', '短板', '衍真']
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 1.32, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    rows = [('数据与环境供应商', 'Scale AI、AfterQuery、Mercor', ['专家多，实验室关系深'],
             ['通用领域为主；交易撮合', '和风险计分需要交易台经验'], ['量化机构出身做撮合和计分', '同一策略，分数可复现']),
            ('公开基准', 'Alpha Arena等榜单', ['方便比较模型'], ['只评测，不产训练数据', '公开后很快饱和'],
             ['同一环境产训练数据', '储备任务持续换题']),
            ('客户自建', '', ['贴合自己的业务'], ['缺专业规则和专家', '占研究人力'],
             ['交易、事件预测环境现成', '专家网络按需调用'])]
    y0, rh = 2.04, 1.32
    for i, (k, names, pro, con, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.85, rh, [para([R(k, 18, C['ink'], SERIF)])] + ([para([R(names, 9, C['grey'], MONO)], before=3)] if names else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.15, rh, pro, 13, C['body'], anchor='ctr', line=1.15, gap=0)
        sh.t(xs[2], y, ws[2] - 0.15, rh, con, 13, C['body'], anchor='ctr', line=1.15, gap=0)
        sh.t(xs[3], y, ws[3] - 0.1, rh, us, 13.5, C['ink'], SANS, True, anchor='ctr', line=1.15, gap=0)
    sh.rule(0.6, y0 + 3 * rh, W)
    footer(sh)


def p_business(sh):
    header(sh, '商业模式', 4, [('先卖给实验室，', F), ('再卖给企业和个人Agent团队', T)])
    segs = [('现在', 'AI实验室', '2家前沿实验室在谈'), ('拓展', '企业Agent团队', '上线前训练、测试、验收'),
            ('下一步', '个人Agent团队', '长任务、跨应用、失败恢复练习')]
    cw, xs = cols(3, 0.3)
    label(sh, 0.6, 1.66, 4, '谁付钱')
    for i, ((stage, name, line), x) in enumerate(zip(segs, xs)):
        now = i == 0
        sh.rect(x, 1.92, cw, 1.05, C['ink'] if now else C['tint'])
        sh.t(x + 0.28, 2.02, cw - 0.5, 0.2, stage, 9, C['accentLt'] if now else C['accent'], MONO, True)
        sh.t(x + 0.28, 2.24, cw - 0.5, 0.36, name, 17, C['onDarkHi'] if now else C['ink'], SERIF)
        sh.t(x + 0.28, 2.62, cw - 0.5, 0.3, line, 11, C['onDark'] if now else C['body'])
    label(sh, 0.6, 3.12, 4, '买什么')
    lines = [('训练环境', '可执行任务和评分工具', '按期限、范围、使用权授权'),
             ('AI数据', 'Agent轨迹、专家数据、评测数据', '按量或批次'),
             ('联合训练', '训练实验、迁移验证、托管运行', '项目或持续服务合同')]
    for i, ((k, what, fee), x) in enumerate(zip(lines, xs)):
        sh.rect(x, 3.38, cw, 1.42, C['tint'])
        sh.t(x + 0.28, 3.5, cw - 0.5, 0.4, k, 19, C['ink'], SERIF)
        sh.t(x + 0.28, 3.94, cw - 0.5, 0.3, what, 11, C['body'])
        sh.t(x + 0.28, 4.32, cw - 0.5, 0.36, [(R('收费  ', 9, C['accent'], MONO), R(fee, 11.5, C['ink'], SANS, True))], 11.5)
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(['需求与验收', '付费试点', '授权交付', '持续更新'], pxs)):
        last = i == 3
        sh.rect(x, 5.0, pw, 0.5, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 5.0, pw - 0.5, 0.5, k, 15, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 5.0, 0.4, 0.5, '→', 13, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.7, [('模型一升级就要换任务和测试集，', F), ('这是复购逻辑', T), ('。', F)], 18)
    note(sh, '收费方式为拟议商业模式，不代表已签署合同。', 6.38)
    footer(sh)


def p_edge(sh):
    header(sh, '团队优势', 5, [('按周交付新环境，', F), ('个人Agent练习场已在开发', T)])
    cards = [
        ('交付', '14天，7个领域', '交易、AI研究、事件预测、数学、推理、财务结账、软件工程',
         [('自有环境构建工具和训练基础设施', '单轮耗时降到1/4'),
          ('20万+专家网络', '7,000+已报名候补'),
          ('验收后交付', '没见过的任务，达标才交')],
         '', F),
        ('下一站 · 开发中', '个人Agent练习场', '面向个人Agent开发团队',
         [('复刻真实应用', '虚拟机和浏览器里的邮件、日历、文档'),
          ('跨应用长任务', '按最终结果打分'),
          ('失败恢复', '弹窗、断网、登录失效时能否自己恢复')],
         '', T),
    ]
    cw, xs = cols(2, 0.3)
    top, ch = 1.72, 3.5
    for (tag, big, sub, items, m, dark), x in zip(cards, xs):
        main, body, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.38, cw - 0.76
        sh.t(ix, top + 0.24, iw, 0.22, tag, 9.5, acc, MONO, True)
        sh.t(ix, top + 0.5, iw, 0.56, big, 30, acc if dark else C['ink'], SERIF)
        sh.t(ix, top + 1.14, iw, 0.26, sub, 11, body)
        sh.rect(ix, top + 1.5, iw, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(ix, top + 1.64, iw, 1.7, [(R(h, 13, main, SANS, True), BR(13), R(d, 11.5, body)) for h, d in items], 13, line=1.12, gap=8)
        if m:
            means(sh, ix, top + 3.46, iw, m, 13, dark)
    contact(sh, 5.6)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_progress, 11, {401, 402, 403, 404, 407}, 4),
    (p_problem, 4, {160}, 5),
    (p_why_now, 3, {126}, 6),
    (p_solution, 6, {237}, 7),
    (p_products, 8, {300}, 8),
    (p_moat, 15, {562}, 9),
    (p_market, 13, {470}, 10),
    (p_competition, 14, {511}, 11),
    (p_business, 10, {376}, 12),
    (p_edge, 16, {607}, 13),
]
SCHOOLS = list(range(429, 441))                   # university logos on v17's network slide, copied onto progress
SCREENSHOT = os.path.join(ASSETS, 'repos-2026-10-02.png')
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out


def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_progress:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            schools = copy_pics(slides[12], s, SCHOOLS)
            build(sh, logos, poly, schools)
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_products:
            s.shapes.add_picture(SCREENSHOT, Inches(0.6), Inches(1.72), width=Inches(5.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
