"""Build SimReal (衍真) BP v74 (18 slides, Chinese) from v17's slides.

v73 with pages 7-9 cut down so each says one thing and does not repeat page 6's RSI loop:
  - 7 强化学习环境: the two real screenshots, larger and without labels, beside three facts (历史行情 / 交易规则 /
    专业打分); the four-step loop, the loop arrow and the closing lines go (page 6 already shows the loop).
  - 8 产品: the footnote is shorter; the strips are as in v73.
  - 9 交付物: the example trajectory has 5 steps instead of 7, the three cards say one thing each, no footnote.
Everything else is as in v73.

Usage: python3 visuals_v58.py; python3 visuals_v69.py; python3 bp_v74.py SimReal-BP-v17.pptx out.pptx
       then embed_cjk.py adds the Chinese font subsets to the shipping PPTX.
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, e, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['我们是谁', '问题与机会', '方法与产品', '进展与商业', '优势与竞争', '融资']
CONF = '机密  ·  仅供受邀投资机构内部评估使用  ·  2026年10月'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.36, 6.98, 0.5, 0.26, '衍真', 9, C['ink'], SANS, True, anchor='ctr')
    sh.t(1.85, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def contact(sh, y):
    sh.t(0.6, y, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)


def p_cover(sh):
    sh.t(1.62, 0.55, 1.0, 0.25, '衍真', 13, C['ink'], SANS, True, anchor='ctr')
    sh.t(0.6, 1.35, 8.6, 1.9, ['让AI在真实世界里', '自我进化'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.5, 9.2, 0.46, '用真实数据，驱动AI精准行动', 22, C['accent'], SERIF)
    sh.rule(0.6, 4.95, 9.0)
    for (k, v), x in zip([('成立', '2026年9月10日'), ('本轮融资', '人民币4,000万元'), ('商业计划书', '2026年10月')], [0.6, 3.6, 6.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, '公司概述', 0, [('衍真是一家AI新实验室（Neolab），', F), ('为下一代AI构建数据、强化学习环境与递归自我改进（RSI）', T)])
    sh.t(0.6, 1.5, W, 0.26, 'SimReal, a neolab building data, RL environments and recursive self-improvement (RSI) for future AI.', 10, C['grey'], SANS)
    cells = [
        ('收入', 'ARR 2,000万元', ['人民币，成立24天', '零外部融资'], T),
        ('开源', 'GitHub 602星', ['成立14天上线7款产品', '其中5个开源基准'], T),
        ('研究成果', '+12%', ['成立两周在交易环境跑通训练', '没见过的交易日上表现最高提升12%'], F),
        ('团队', '05后量化团队', ['剑桥、LSE、杜克数学', 'Jane Street、Citadel、Optiver、Millennium'], F),
        ('市场', '85亿美元', ['训练数据与RL环境，今天的年收入规模', 'Mercor 16个月收入涨27倍'], F),
        ('本轮融资', '4,000万元', ['人民币', '投向环境、RSI与交付'], T),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.98 + (i // 3) * 2.05
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, v, 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.8, d, 11.5, C['body'], line=1.15, gap=1)
    footer(sh)


def p_team(sh):
    header(sh, '团队', 0, [('三位21岁创始人，', F), ('放弃首年合计600万+薪酬的工作，搭建下一代训练引擎', T)])
    b, sz = C['body'], 11
    people = [
        ('Charles', 'CEO', [(R('United Stables首位员工：U稳定币', sz, b), BR(sz), R('一年从0做到14亿美元，一个月上线币安', sz, b)),
                            '负责机构合作，对接SIG、DRW等',
                            (R('汇丰港元稳定币项目唯一实习生，', sz, b), BR(sz), R('参与香港金管局合规', sz, b)),
                            '伦敦Citadel对冲基金实习'],
         ['LSE数学 · 深国交', '美国数学奥林匹克（USAMO）入围']),
        ('Henry', 'CTO', ['剑桥数学一等荣誉，奖学金获得者',
                          (R('剑桥AI研究中心最年轻的研究员，', sz, b), BR(sz), R('导师Po-Ling Loh（国际数理统计学会会士）', sz, b)),
                          'Jane Street、Citadel、Optiver量化经历', '主导设计全部5个开源基准'],
         ['深国交', '剑桥数学竞赛全球前30 · 英国物理竞赛超级金奖']),
        ('Amaris', 'COO', ['Millennium香港数据科学家', 'Millennium另类数据团队首位应届生',
                           (R('参与筹办Plug and Play香港首场活动', sz, b), BR(sz), R('（联合香港科技园，200+人）', sz, b))],
         ['杜克数学与统计 · 上海包玉刚']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.75
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO),
                                                          R('   21岁', 11, C['ink'], SANS, True)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.8, cw - 0.6, 2.2, lines, sz, b, line=1.12, gap=6)
        sh.rule(x + 0.3, y0 + 2.98, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.08, cw - 0.6, 0.6, edu, 9.5, C['grey'], line=1.12, gap=1)
    footer(sh)


def p_problem(sh):
    header(sh, '问题', 1, [('AI会推理，', F), ('但干不好真实的活', T)])
    rows = [('交易', '回测一路上涨', '实盘行情里亏钱'), ('软件工程', '测试全部通过', '一上线就出错'),
            ('财务', '账看着做完了', '月底对不上'), ('事件预测', '分析头头是道', '结果一出就错')]
    y0, rh = 2.15, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, '看上去', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, '实际上', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10.5, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 19, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 19, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    sh.t(0.6, 5.0, 1.8, 0.6, '6/32', 30, C['accent'], SERIF, anchor='ctr')
    sh.t(2.5, 5.0, 9.5, 0.6, '前沿大模型拿真钱做交易，32轮只有6轮赚钱（Alpha Arena）', 13, C['body'], anchor='ctr')
    kicker(sh, 5.85, [('缺的是一个', F), ('按真实结果打分的训练场', T), ('。', F)], 21)
    footer(sh)


def p_why_now(sh):
    header(sh, '为什么是现在', 1, [('训练的瓶颈从公开文本转到可验证的数据和环境，', F), ('钱已经在花，专业领域的供给还很少', T)])
    items = [('01', '从读书到做题', ['公开的人类文本预计2026–2032年用完', '推理模型改靠强化学习提升：在环境里做题、自动判对错，产出可验证的训练数据'],
              '可验证的数据和环境，成了稀缺资源'),
             ('02', '钱已经在花', ['一道强化学习训练题卖200–2,000美元', 'Mercor 2026年7月收购环境公司Deeptune', 'RSI新实验室Recursive估值46.5亿美元'],
              '实验室和资本为数据、环境、RSI买单'),
             ('03', '专业数据和环境还很少', ['Meta、OpenAI 9月推出替人办事的Agent', '但做交易、做财务，大模型仍频繁出错'],
              '先从对错最清楚的交易做起')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for (n, k, ev, m), x in zip(items, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 22, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.1, cw - 0.6, 1.2, ev, 12, C['body'], line=1.15, gap=6)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, m, 12.5)
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, '窗口期', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R('RL环境赛道约20家早期公司，预计最后剩3–5家。', 13, C['onDark']),
                                       R('先做出效果的，先拿下长期合同', 13, C['onDarkHi'], SANS, True))], 13, anchor='ctr', line=1.1)
    note(sh, '来源：Epoch AI（2026年1月）；TechCrunch（2026年7月）；SiliconANGLE（2026年5月）；Meta、OpenAI发布（2026年9月）；Wing VC（2026年）。', 6.3)
    footer(sh)


def p_market(sh):
    header(sh, '市场空间', 3, [('今天85亿美元，', F), ('头部公司16个月收入涨27倍', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, '今天：训练数据与RL环境')
    sh.t(0.9, 2.25, 5, 0.72, '85亿美元 / 年', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, '50余家供应商的收入合计', 11.5, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030年：公司测算', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '7,000亿美元 / 年', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, '按AI经济7万亿美元、10%投入训练估算', 11.5, C['body'])
    growth = [('27倍', ['Mercor年化收入16个月', '从7,500万涨到20亿美元'], '专家数据需求在爆发'),
              ('10倍', ['Mercor估值17个月', '从20亿涨到200亿美元（洽谈中）'], '资本在加注'),
              ('18倍', ['Snorkel AI新数据业务', '一年增长18倍'], '需求转向专家和环境')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 12, C['body'], line=1.1, gap=0)
        means(sh, x, 5.38, cw, m, 12.5)
    note(sh, '来源：Deedy Das行业图谱（2026年7月）；Mercor（TechCrunch、Dealroom）；Snorkel AI公告（2026年9月）。2030年为公司测算。', 6.3)
    footer(sh)


def p_customers(sh):
    header(sh, '客户', 3, [('我们卖环境、数据和评测：', F), ('前沿实验室现在付费，企业和个人Agent是下一步', T)])
    cards = [('现在', '前沿AI实验室', '环境、经过验证的训练题和评测，用来做后训练', '按环境、按题计费，签季度合同'),
             ('下一步', '企业Agent团队', '定制环境和验收评测，证明Agent上线前能把活干好', '先付费试点，再签年度授权'),
             ('再下一步 · C端', '个人用户', ['让AI快速摸清自己的习惯，', '每天用得顺手'], '按月订阅，或经Agent应用分发')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.5
    for (tag, k, buy, charge), x in zip(cards, xs):
        dark = tag == '现在'
        main, sub, acc, rule = ((C['onDarkHi'], C['onDark'], C['accentLt'], DARK_RULE) if dark else (C['ink'], C['body'], C['accent'], C['mid']))
        ix, iw = x + 0.32, cw - 0.64
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(ix, top + 0.24, iw, 0.22, tag, 10, acc, SANS, True)
        sh.t(ix, top + 0.52, iw, 0.5, k, 24, main, SERIF)
        sh.t(ix, top + 1.3, iw, 0.2, '他们买什么', 9, sub, SANS)
        sh.t(ix, top + 1.56, iw, 0.7, buy, 13, main, line=1.2, gap=0)
        sh.rect(ix, top + 2.4, iw, 0.01, rule)
        sh.t(ix, top + 2.54, iw, 0.2, '怎么收费', 9, acc, SANS, True)
        sh.t(ix, top + 2.8, iw, 0.6, charge, 13.5, main, SANS, True, line=1.2)
    kicker(sh, 5.6, [('同一套环境服务三类客户，', F), ('模型每升级一次都需要新题', T)], 20)
    footer(sh)


RSI_IMG = os.path.join(ASSETS, 'v58', 'rsi.png')


def p_rsi(sh):
    header(sh, '解决方案', 2, [('递归自我改进（RSI）：模型在环境里做题、被打分、补短板，', F), ('一圈比一圈强', T)])
    rx, rw = 8.35, 4.38
    label(sh, rx, 1.86, rw, '衍真强在哪', C['accent'])
    rows = [('打分专业', '量化标准，最强模型77分，其余24–30分'), ('出题又快又准', '20万+专家把关')]
    for i, (k, d) in enumerate(rows):
        y = 2.14 + i * 1.0
        sh.rule(rx, y, rw, C['mid'])
        sh.t(rx, y + 0.12, rw, 0.4, k, 18, C['ink'], SERIF)
        sh.t(rx, y + 0.56, rw, 0.3, d, 11.5, C['body'])
    sh.rule(rx, 4.14, rw, C['mid'])
    sh.rect(rx, 4.36, rw, 1.3, C['ink'])
    sh.t(rx + 0.3, 4.48, 3.9, 0.22, '效果已验证', 9.5, C['accentLt'], MONO, True)
    sh.t(rx + 0.3, 4.72, 1.6, 0.7, '+12%', 32, C['accentLt'], SERIF, anchor='ctr')
    sh.t(rx + 1.75, 4.72, 2.5, 0.7, ['没见过的交易日上', '表现最高提升12%'], 11.5, C['onDarkHi'], line=1.1, gap=0, anchor='ctr')
    note(sh, '分数来自Xitadel公开测评；+12%为多次独立复现中的最佳一次。', 6.36)
    footer(sh)


RL_IMG = os.path.join(ASSETS, 'v70')        # orderbook.png, replay.png: cropped from the team's v62 page "强化学习环境"


def p_rl_env(sh):
    header(sh, '强化学习环境', 2, [('让AI做量化研究员的工作，', F), ('用同一套标准考核', T)])
    rx, rw = 7.45, W + 0.6 - 7.45
    rows = [('历史行情', '订单簿逐时点数据，三档买卖报价'), ('交易规则', '按规则撮合成交，有仓位限制'),
            ('专业打分', '独立测试日，按收益、回撤、夏普计分')]
    for i, (k, d) in enumerate(rows):
        y = 1.82 + i * 1.63
        sh.rule(rx, y, rw, C['mid'])
        sh.t(rx, y + 0.3, rw, 0.46, k, 20, C['ink'], SERIF, anchor='ctr')
        sh.t(rx, y + 0.82, rw, 0.3, d, 12, C['body'])
    sh.rule(rx, 6.7, rw, C['mid'])
    footer(sh)


PRODUCT_IMG = os.path.join(ASSETS, 'v69')
DETAIL = [('交易环境', 'Xitadel', '102', '公开成绩', 'xitadel.png'),
          ('AI研究环境', 'SimReal-MLBench', '302', '基线成绩，并非AI Agent', 'mlbench.png'),
          ('事件预测环境', 'FuturePredict', '27', '数字为示例', 'forecast.png')]
DY0, DRH, DGAP = 1.84, 1.42, 0.06
STEP_COLS = [(3.24, '1', '给AI的题'), (6.14, '2', 'AI做什么'), (9.04, '3', '和谁比')]   # panel lefts 2 / 292 / 582 px in the strip


def p_products(sh):
    header(sh, '产品', 2, [('三个训练环境已上线，', F), ('都拿真实结果给AI打分', T)])
    for x, n, lab in STEP_COLS:
        sh.t(x, 1.58, 2.6, 0.22, [(R(n, 9.5, C['accent'], MONO, True), R('  ' + lab, 9.5, C['grey'], SANS))], 9.5, anchor='ctr')
    for i, (tag, name, stars, prov, _img) in enumerate(DETAIL):
        y = DY0 + i * (DRH + DGAP)
        sh.rect(0.6, y, 2.53, DRH, C['tint'])                 # left column only: the storyboard picture brings its own tint
        sh.rect(12.66, y, W + 0.6 - 12.66, DRH, C['tint'])
        sh.rect(0.6, y, 0.05, DRH, C['accent'])
        sh.t(0.85, y + 0.24, 2.2, 0.22, tag, 10, C['accent'], SANS, True)
        sh.t(0.85, y + 0.46, 2.2, 0.4, name, 18, C['ink'], SERIF)
        sh.t(0.85, y + 0.9, 1.4, 0.2, f'★  {stars}', 9.5, C['accent'], MONO, True)
        sh.t(0.85, y + 1.12, 2.2, 0.22, prov, 9.5, C['grey'])
    end = DY0 + 3 * DRH + 2 * DGAP
    sh.rect(0.6, end + 0.06, W, 0.36, C['tint'])
    sh.t(0.85, end + 0.06, W - 0.4, 0.36, [(R('5个开源基准，GitHub 602星', 12, C['accent'], SANS, True),
                                           R('：SimReal-MLBench 302  ·  Xitadel 102  ·  MathmoBench 101  ·  Puzzle 70  ·  FuturePredict 27'
                                             '（截至2026年10月4日）', 12, C['ink']))], 12, anchor='ctr')
    note(sh, '人类最佳：IMC Prosperity（全球交易竞赛）每项任务的最佳真实策略；基线：提前锁定的参考预测。', 6.68, 8)
    footer(sh)


def pill(sh, x, y, w, h, text, sz, color, fill=None, line=None):
    """A small rounded tag with centred bold text; fill and/or outline."""
    i = sh._id()
    fl = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else '<a:noFill/>'
    ln = f'<a:ln w="12700"><a:solidFill><a:srgbClr val="{line}"/></a:solidFill></a:ln>' if line else '<a:ln><a:noFill/></a:ln>'
    sh.xml.append(
        f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="SimReal {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 22000"/></a:avLst></a:prstGeom>{fl}{ln}</p:spPr>'
        f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
        f'{para([R(text, sz, color, SANS, True)], "ctr")}</p:txBody></p:sp>')


def p_deliverables(sh):
    header(sh, '交付物', 2, [('环境每跑一遍，交付一条带验证的轨迹：', F), ('不只是数据，还有证据', T)])
    px, py, pw, ph = 0.6, 1.72, 7.26, 3.84
    sh.rect(px, py, pw, ph, C['ink'])
    sh.t(px + 0.25, py + 0.18, 4, 0.22, '一条交易轨迹  ·  XITADEL', 8.5, C['onDark'], MONO)
    sh.t(px + pw - 1.25, py + 0.18, 1.0, 0.22, '示意', 8.5, C['onDark'], SANS, algn='r')
    cx = [px + 0.25, px + 0.85, px + 1.95]
    steps = [('01', '观察', '订单簿：买一 101.05，卖一 101.10'), ('02', '行动', '限价买入 100股 @ 101.05'),
             ('03', '世界反应', '成交 60股，余量改价后全部成交'), ('04', '行动', '收盘前卖出 100股 @ 101.30'),
             ('05', '结算', '按结果计分，写入奖励')]
    style = {'观察': ('4A4844', None, C['onDarkHi']), '行动': (C['accent'], None, 'FFFFFF'),
             '世界反应': (None, C['onDarkHi'], C['onDarkHi']), '结算': ('E8E6E1', None, C['ink'])}
    rh, y0 = 0.48, py + 0.52
    sh.rect(px + 0.25, y0, pw - 0.5, 0.01, DARK_RULE)
    for i, (n, k, d) in enumerate(steps):
        y = y0 + i * rh
        sh.t(cx[0], y, 0.5, rh, n, 9.5, C['onDark'], MONO, anchor='ctr')
        fill, line, color = style[k]
        pill(sh, cx[1], y + 0.11, 0.86, 0.26, k, 9.5, color, fill, line)
        sh.t(cx[2], y, pw - 2.2, rh, d, 12.5, C['onDarkHi'], anchor='ctr')
        sh.rect(px + 0.25, y + rh, pw - 0.5, 0.01, DARK_RULE)
    sh.t(px + 0.25, py + ph - 0.62, pw - 0.5, 0.36, [(R('验证器：通过 ✓', 12.5, C['accentLt'], SANS, True), R('    ·    ', 12.5, C['onDark']),
                                                      R('结果哈希已预先公开', 12.5, C['accentLt'], SANS, True))], 12.5, anchor='ctr')
    cards = [('01', '完整轨迹', '每一步都可回放，直接用于微调和强化学习'),
             ('02', '自动验证', '按真实结果判定，评分器上线前先经攻击测试'),
             ('03', '质检报告', '每批附污染检测、难度分布和验证通过率')]
    rx = px + pw + 0.18
    rw, chh, cg = W + 0.6 - rx, 1.2, 0.12
    for i, (n, k, d) in enumerate(cards):
        y = py + i * (chh + cg)
        sh.rect(rx, y, rw, chh, C['tint'])
        sh.t(rx + 0.26, y + 0.16, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(rx + 0.26, y + 0.36, rw - 0.5, 0.42, k, 20, C['ink'], SERIF)
        sh.t(rx + 0.26, y + 0.8, rw - 0.5, 0.36, d, 11, C['body'], line=1.15)
    kicker(sh, 5.86, [('数商卖人力，', F), ('衍真卖验证。', T)], 22)
    footer(sh)


def p_why_us(sh):
    header(sh, '为什么是我们', 4, [('2026年RSI的三个变化，', F), ('每一个我们都已经在交付', T)])
    xl, wl, xr = 0.6, 6.2, 7.45
    wr = W + 0.6 - xr
    sh.rect(xr - 0.25, 1.72, wr + 0.25, 0.34 + 3 * 1.08, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    sh.t(xl + 0.2, 1.74, 4, 0.32, '2026年的变化', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(xr, 1.74, 4, 0.32, '衍真', 9.5, C['accent'], MONO, True, anchor='ctr')
    rows = [('AI开始自己做研究', 'OpenAI研究团队每人每天配3.1个Agent工作日', 'OpenAI，2026年9月',
             ['MLBench：AI做研究的考场', 'GitHub 302星']),
            ('AI开始自己改进自己', 'Agent改写自己的代码，8天完成7轮自我改进', 'Weco AIDE²，2026年9月',
             ['交易环境里跑通RSI', '没见过的交易日上+12%']),
            ('卡点变成环境', '环境一乱，最强Agent通过率从83.9%掉到57.6%', 'Breaking the Environment Wall，2026年9月',
             ['20万+专家', '3周交付数百万需求'])]
    y0, rh = 2.06, 1.08
    for i, (k, ev, src, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.t(xl + 0.2, y + 0.14, wl, 0.42, k, 20, C['ink'], SERIF)
        sh.t(xl + 0.2, y + 0.56, wl, 0.26, ev, 12, C['body'])
        sh.t(xl + 0.2, y + 0.82, wl, 0.2, src, 8.5, C['grey'], MONO)
        sh.t(xr - 0.7, y, 0.4, rh, '→', 16, C['accent'], algn='ctr', anchor='ctr')
        sh.t(xr, y, wr - 0.1, rh, [(R(us[0], 15, C['ink'], SANS, True),), (R(us[1], 12, C['body']),)], 15, anchor='ctr', line=1.1, gap=3)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.rect(0.6, 5.62, W, 0.68, C['ink'])
    sh.t(0.85, 5.62, 1.3, 0.68, '所以', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.0, 5.62, W - 1.6, 0.68, [(R('RSI卡在环境，环境正是衍真的研究核心。', 13.5, C['onDark']),
                                      R('多数AI新实验室还没有收入，衍真成立24天ARR 2,000万元', 13.5, C['onDarkHi'], SANS, True))], 13.5, anchor='ctr')
    note(sh, '来源：OpenAI（2026年9月6日）；Weco AI，arXiv 2609.26457；Breaking the Environment Wall，arXiv 2609.29773；Deedy Das新实验室名单（2026年5月）。+12%为多次独立复现中的最佳一次。', 6.45)
    footer(sh)


def p_competition(sh):
    header(sh, '竞争', 4, [('对手强在规模，', F), ('我们强在专业领域的打分和速度', T)])
    xs, ws = [0.6, 3.88, 6.6, 9.45], [3.1, 2.6, 2.75, 3.28]
    heads = ['客户现在的选择', '强在哪', '弱在哪', '衍真']
    y0, rh = 2.04, 0.74
    sh.rect(xs[3] - 0.57, 1.72, W + 0.6 - xs[3] + 0.57, 0.32 + 3 * rh, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.3, h, 9.5, C['accent'] if i == 3 else C['grey'], SANS, i == 3, anchor='ctr')
    rows = [('数据外包公司', 'Scale AI、Mercor等', '人多，实验室关系深', '靠堆人，做不了专业打分', '量化出身，打分专业'),
            ('公开榜单', 'Alpha Arena等', '方便比较模型', '只能测不能训，很快被刷爆', '同一个环境既能测也能训'),
            ('实验室自建', '', '贴合自己的需求', '缺规则缺专家，占研究人力', '现成环境，20万+专家随时调用')]
    for i, (k, names, pro, con, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.9, rh, [para([R(k, 18, C['ink'], SERIF)])] + ([para([R(names, 8.5, C['grey'], MONO)], before=2)] if names else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.2, rh, pro, 12, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[2], y, ws[2] - 0.2, rh, con, 12, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[3], y, ws[3] - 0.15, rh, us, 12.5, C['ink'], SANS, True, anchor='ctr', line=1.1)
    sh.rule(0.6, y0 + 3 * rh, W)
    sh.t(0.6, 4.5, 3, 0.22, '为什么难被抄', 9.5, C['accent'], SANS, True)
    moat = [('失效地图', '持续跑前沿模型，知道它们在哪类题上失手', '前沿模型24–77分，人类最佳80分'),
            ('成本随规模下降', '验证器建好后，合成加自动核验批量生产', '单轮耗时降到1/4，同等资源多完成64%尝试'),
            ('能自证', '无污染可审计，训练增益可验证', '结果哈希预先公开；训练后+12%')]
    cw, cxs = cols(3, 0.35)
    for (k, d, ev), x in zip(moat, cxs):
        sh.rule(x, 4.8, cw, C['ink'])
        sh.t(x, 4.88, cw, 0.4, k, 17, C['ink'], SERIF)
        sh.t(x, 5.32, cw, 0.24, d, 11, C['body'])
        sh.t(x, 5.58, cw, 0.24, ev, 11, C['accent'], SANS, True)
    kicker(sh, 6.0, [('对手拼规模，我们拼加速度：', F), ('14天上线7款产品，24天做到ARR 2,000万元', T)], 20)
    footer(sh)


def p_business(sh):
    header(sh, '商业模式', 3, [('环境授权、数据交付、联合训练，', F), ('模型每升级一次就要复购', T)])
    lines = [('环境授权', '按期限和范围收费', '交易、AI研究、事件预测等训练环境'),
             ('数据交付', '按量收费', ['环境每跑一遍就产出数据：', '轨迹、专家示范、评测题']),
             ('联合训练', '按项目收费', '用我们的环境，帮客户把模型练强')]
    cw, xs = cols(3, 0.3)
    for i, ((k, fee, what), x) in enumerate(zip(lines, xs)):
        sh.rect(x, 1.72, cw, 2.05, C['tint'])
        sh.t(x + 0.32, 1.92, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.32, 2.2, cw - 0.6, 0.5, k, 24, C['ink'], SERIF)
        sh.t(x + 0.32, 2.82, cw - 0.64, 0.5, what, 11.5, C['body'], line=1.15, gap=0)
        sh.t(x + 0.32, 3.36, cw - 0.6, 0.3, fee, 13, C['ink'], SANS, True)
    label(sh, 0.6, 4.08, 6, '合作路径')
    path = ['需求与验收标准', '付费试点', '授权交付', '持续更新']
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(path, pxs)):
        last = i == 3
        sh.rect(x, 4.36, pw, 0.62, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 4.36, pw - 0.5, 0.62, k, 16, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 4.36, 0.4, 0.62, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.35, [('模型每升级一次都要换题，', F), ('环境和数据要持续更新，客户持续付费', T)], 18)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_traction(sh):
    header(sh, '当前进展', 3, [('2026年9月10日成立，', F), ('24天做到ARR 2,000万元', T)])
    ms = [('第14天', '7款', '产品上线', '交易环境同期跑通训练'),
          ('第3周', '数百万', '交付需求', ''),
          ('第24天', '2,000万元', 'ARR', '人民币，零外部融资'),
          ('第24天', '602', 'GitHub星标', '5个开源基准')]
    cw, xs = cols(4, 0.3)
    ax = 2.08
    sh.rect(0.6, ax, W, 0.02, C['ink'])
    for i, ((day, v, k, d), x) in enumerate(zip(ms, xs)):
        acc = i >= 2
        sh.rect(x, ax - 0.07, 0.16, 0.16, C['accent'])
        sh.t(x + 0.24, ax - 0.4, cw - 0.3, 0.26, day, 10.5, C['accent'], MONO, True, anchor='ctr')
        sh.t(x, ax + 0.22, cw, 0.7, v, 40 if len(v) <= 3 else 34, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(x, ax + 0.96, cw, 0.3, k, 13, C['ink'], SANS, True)
        if d:
            sh.t(x, ax + 1.28, cw, 0.3, d, 10.5, C['body'])
    sh.t(0.6, 3.72, W, 0.24, '“第N天”从成立日算起；截至2026年10月4日', 8.5, C['grey'])
    sh.rule(0.6, 4.12, W, C['ink'])
    label(sh, 0.6, 4.24, 6, '客户')
    sh.t(0.6, 4.5, 2.6, 0.7, '2家', 34, C['ink'], SERIF, anchor='ctr')
    sh.t(0.6, 5.22, 2.7, 0.3, '前沿实验室在谈', 13, C['ink'], SANS, True)
    label(sh, 3.75, 4.24, 8, '训练基础设施：同样算力，跑更多训练', C['accent'])
    infra = [('1/4', '单轮耗时降到1/4'), ('+64%', '同样资源，多完成64%的训练尝试'), ('−1/3', '存档耗时减少1/3，存档次数减半')]
    iw, ixs = cols(3, 0.3, 3.75, W + 0.6 - 3.75)
    for (v, k), x in zip(infra, ixs):
        sh.rect(x, 4.5, iw, 0.02, C['ink'])
        sh.t(x, 4.56, iw, 0.66, v, 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x, 5.22, iw, 0.5, k, 11.5, C['body'], line=1.1)
    note(sh, '基础设施指标为公司内部测试，口径不同，不能相乘或相加。', 6.3)
    footer(sh)


def p_network_pro(sh, logos, polymarket):
    header(sh, '专家网络', 4, [('参与研发的从业者来自顶级交易机构，', F), ('他们的产出由验证器逐条检验', T)])
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 2.25, 1.55 if pic.shape_id == 403 else 1.2)
    if polymarket is not None:
        place(polymarket, slots[4], 2.25, 1.1)
    sh.rule(0.6, 2.95, W, C['ink'])
    cells = [('20万+', '可触达的专家', T), ('7,000+', '已报名的候补专家', F), ('出题 · 打分 · 把关', '专家在环境里做的事', F)]
    cw, xs = cols(3)
    for (v, k, acc), x in zip(cells, xs):
        sh.t(x, 3.08, cw, 0.8, v, 44 if len(v) < 8 else 28, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(x, 3.94, cw, 0.3, k, 13, C['ink'], SANS, True)
    vs = [('数据外包公司', ['专家按小时计费，卖的是工时', '质量靠人工抽检，项目做完就散'], False),
          ('衍真', ['一线从业者出题、打分、把关，产出由验证器逐条检验', '题目和验证器沉淀成自有环境，客户越多越值钱'], True)]
    vw, vxs = cols(2, 0.3)
    for (k, lines, dark), x in zip(vs, vxs):
        main, sub = (C['onDarkHi'], C['accentLt']) if dark else (C['ink'], C['grey'])
        sh.rect(x, 4.55, vw, 1.45, C['ink'] if dark else C['tint'])
        sh.t(x + 0.3, 4.72, vw - 0.6, 0.22, k, 10, sub, SANS, True)
        sh.t(x + 0.3, 5.02, vw - 0.6, 0.9, lines, 13, main, SANS, dark, line=1.15, gap=6)
    note(sh, '标识仅表示从业者任职机构，不代表背书。', 6.4, 8)
    footer(sh)


def p_network_edu(sh, schools):
    header(sh, '高校网络', 4, [('21所顶尖高校的学生和校友，', F), ('覆盖顶级入门岗位和前沿科研', T)])
    cell = W / 6
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, 0.6 + (i % 6 + 0.5) * cell, 2.2 + (i // 6) * 0.98, min(1.45 / w, 0.68 / h, 1.8))
    sh.rule(0.6, 3.92, W, C['ink'])
    cover = [('顶级入门岗位', '交易、金融、软件等竞争最激烈的校招岗位', '为替人处理日常工作的Agent提供训练数据'),
             ('前沿科研', 'AI、数学与量化研究', '为自动做研究的Agent提供训练数据')]
    cw, xs = cols(2, 0.3)
    for (k, d, m), x in zip(cover, xs):
        sh.rect(x, 4.12, cw, 1.3, C['tint'])
        sh.rect(x, 4.12, 0.05, 1.3, C['accent'])
        sh.t(x + 0.32, 4.26, cw - 0.6, 0.42, k, 20, C['ink'], SERIF)
        sh.t(x + 0.32, 4.7, cw - 0.6, 0.26, d, 12, C['body'])
        means(sh, x + 0.32, 5.0, cw - 0.6, m, 12.5)
    kicker(sh, 5.6, [('这两类工作，', F), ('正是今天Agent主攻的两个方向', T)], 18)
    note(sh, '图为网络覆盖的部分高校；标识不代表学校背书。', 6.4, 8)
    footer(sh)


def p_raise(sh):
    header(sh, '本轮融资', 5, [('融资4,000万元，', F), ('投向环境、RSI和交付', T)])
    top, ch, lw = 1.72, 3.86, 2.9
    sh.rect(0.6, top, lw, ch, C['ink'])
    sh.t(0.9, top + 0.3, lw - 0.6, 0.22, '本轮融资（人民币）', 10, C['accentLt'], SANS, True)
    sh.t(0.9, top + 0.66, lw - 0.6, 0.8, '4,000万元', 36, C['accentLt'], SERIF)
    sh.t(0.9, top + ch - 0.62, lw - 0.6, 0.4, '投向三个方向，各有本轮目标  →', 11, C['onDark'], line=1.1)
    dirs = [('01', '环境与数据', '把交易的打法复制到更多专业领域', '环境工程师  ·  专家网络扩容', '10+个专业领域环境'),
            ('02', 'RSI', '把自我进化扩到更大的模型、更多的行业', '算力  ·  研究员', '开源大模型实现RSI'),
            ('03', '客户交付', '前沿实验室从试点做到长期合同', '交付团队  ·  商务', '累计20家客户')]
    x0 = 0.6 + lw + 0.25
    cw, xs = cols(3, 0.25, x0, W + 0.6 - x0)
    for (n, k, d, inv, goal), x in zip(dirs, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        ix, iw = x + 0.28, cw - 0.5
        sh.t(ix, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(ix, top + 0.48, iw, 0.5, k, 24, C['ink'], SERIF)
        sh.t(ix, top + 1.12, iw, 0.62, d, 13, C['ink'], SANS, True, line=1.15)
        sh.rect(ix, top + 1.96, iw, 0.01, C['mid'])
        sh.t(ix, top + 2.08, iw, 0.2, '投入', 9, C['body'], SANS)
        sh.t(ix, top + 2.3, iw, 0.42, inv, 12, C['ink'], line=1.1)
        sh.t(ix, top + 2.84, iw, 0.2, '本轮目标', 9, C['accent'], SANS, True)
        sh.t(ix, top + 3.06, iw, 0.62, goal, 16, C['accent'], SERIF, line=1.05)
    contact(sh, 5.86)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_problem, 4, {160}, 4),
    (p_why_now, 3, {126}, 5),
    (p_rsi, 6, {237}, 6),
    (p_rl_env, 9, {336}, 7),
    (p_products, 8, {300}, 8),
    (p_deliverables, 5, {169}, 9),
    (p_traction, 16, {607}, 10),
    (p_market, 13, {470}, 11),
    (p_customers, 20, {666}, 12),
    (p_business, 10, {376}, 13),
    (p_why_us, 15, {562}, 14),
    (p_competition, 14, {511}, 15),
    (p_network_pro, 11, {401, 402, 403, 404, 407}, 16),
    (p_network_edu, 12, set(range(429, 441)) | {442}, 17),
    (p_raise, 17, {607}, 18),
]
SCHOOLS = list(range(429, 441))
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out



def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_network_edu:
            build(sh, sorted((p for p in s.shapes if p.shape_id in SCHOOLS), key=lambda p: SCHOOLS.index(p.shape_id)))
        elif build is p_network_pro:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            build(sh, logos, poly)
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_rl_env:
            s.shapes.add_picture(os.path.join(RL_IMG, 'orderbook.png'), Inches(0.6), Inches(1.82), width=Inches(6.3))
            s.shapes.add_picture(os.path.join(RL_IMG, 'replay.png'), Inches(0.6), Inches(3.92), width=Inches(6.3))
            build(sh)
        elif build is p_products:
            for i, (*_rest, img) in enumerate(DETAIL):
                s.shapes.add_picture(os.path.join(PRODUCT_IMG, img), Inches(3.12), Inches(DY0 + i * (DRH + DGAP)), width=Inches(9.55))
            build(sh)
        elif build is p_rsi:
            s.shapes.add_picture(RSI_IMG, Inches(0.45), Inches(1.6), width=Inches(7.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
