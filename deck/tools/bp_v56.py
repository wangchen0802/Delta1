"""Build SimReal (衍真) BP v56 (16 slides, Chinese) from v17's slides.

v55 plus: the cover gets v39's intro lines and the founding date (2026-09-10), so every "成立第N天" later
lands; page 2 says 05后; an RSI roadmap page grounds the loop in published work (DeepSeek-R1, Absolute Zero,
AlphaEvolve, Darwin Gödel Machine ...) and ends on the big claim; progress splits again into traction (a
timeline from founding, plus the training-infrastructure numbers) and the expert network; the raise page says
where the money goes, not how much per line; the data page folds into the business model.

Usage: python3 visuals_v55.py; python3 bp_v56.py SimReal-BP-v17.pptx out.pptx
       then embed_cjk.py adds the Chinese font subsets to the shipping PPTX.
"""
import copy
import io
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches

from bp_v17 import C, MONO, SANS, SERIF, para
from bp_v34 import BR, CUR, DARK_RULE, NS, R, S, W, recolor, title_runs

F, T = False, True
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, '..', 'assets')
SECTIONS = ['我们是谁', '问题与机会', '方法与产品', '优势与竞争', '商业与进展', '融资']
CONF = '机密  ·  仅供受邀投资机构内部评估使用  ·  2026年10月'


# ------------------------------------------------------------------ frame ---
def header(sh, label, section, parts):
    sh.t(0.6, 0.42, 3, 0.24, f'{CUR["n"]:02d}', 10, C['accent'], MONO, True, anchor='ctr')
    nav = []
    for i, name in enumerate(SECTIONS):
        if i:
            nav.append(R('   ', 8.5, C['faint']))
        nav.append(R(name, 8.5, C['accent'] if i == section else C['faint'], SANS, i == section))
    sh.text(6.6, 0.42, 6.13, 0.24, [para(nav, 'r')], 'ctr')
    sh.t(0.6, 0.6, W, 0.56, label, 30, C['ink'], SERIF)
    sh.text(0.6, 1.12, W, 0.36, [para(list(title_runs(parts, 16)))])


def footer(sh):
    sh.t(1.36, 6.98, 0.5, 0.26, '衍真', 9, C['ink'], SANS, True, anchor='ctr')
    sh.t(1.85, 6.98, 8.0, 0.26, CONF, 8, C['grey'], anchor='ctr')
    if CUR['n'] is not None:
        sh.t(11.73, 6.98, 1.0, 0.26, f'{CUR["n"]:02d}', 9, C['grey'], MONO, algn='r', anchor='ctr')


def kicker(sh, y, parts, sz=20):
    sh.text(0.6, y, W, 0.46, [para(list(title_runs(parts, sz)))], 'ctr')


def note(sh, text, y=6.42, sz=8):
    sh.t(0.6, y, W, 0.4, text, sz, C['grey'], line=1.1, gap=0)


def label(sh, x, y, w, text, color=None):
    sh.t(x, y, w, 0.22, text, 9.5, color or C['grey'], MONO)


def cols(n, gap=0.35, x0=0.6, w=W):
    cw = (w - (n - 1) * gap) / n
    return cw, [x0 + i * (cw + gap) for i in range(n)]


def means(sh, x, y, w, text, sz=11.5, dark=False):
    """'→ what it means for us': the line every number on this deck ends with."""
    sh.t(x, y, w, 0.5, [(R('→ ', sz, C['accentLt'] if dark else C['accent']),
                         R(text, sz, C['onDarkHi'] if dark else C['ink'], SANS, True))], sz, line=1.1)


# ------------------------------------------------------------------ pages ---
def contact(sh, y):
    sh.t(0.6, y, W, 0.3, [(R('business@simreal.co', 11, C['ink'], MONO), R('   ·   ', 11, C['grey']), R('simreal.com.cn', 11, C['ink'], MONO))], 11)


def p_cover(sh):
    sh.t(1.62, 0.55, 1.0, 0.25, '衍真', 13, C['ink'], SANS, True, anchor='ctr')
    sh.t(0.6, 1.35, 8.6, 1.9, ['让AI在真实世界里', '自我进化'], 44, C['ink'], SERIF, line=1.05)
    sh.t(0.6, 3.42, 9.2, 0.46, '用真实数据重建真实世界', 22, C['ink'], SERIF)
    sh.t(0.6, 4.1, 9.6, 0.36, 'AI实验室、企业与个人Agent的训练场', 16, C['accent'])
    sh.rule(0.6, 4.95, 9.0)
    for (k, v), x in zip([('成立', '2026年9月10日'), ('本轮融资', '人民币4,000万元'), ('商业计划书', '2026年10月')], [0.6, 3.6, 6.6]):
        sh.t(x, 5.12, 3.2, 0.22, k, 9.5, C['grey'], MONO)
        sh.t(x, 5.4, 3.4, 0.36, v, 15, C['ink'])
    sh.t(0.6, 6.98, 8, 0.26, CONF, 8, C['grey'], anchor='ctr')


def p_overview(sh):
    header(sh, '公司概述', 0, [('衍真是一家Neolab，', F), ('为下一代AI搭建数据基础设施和强化学习环境，研发递归自我改进（RSI）', T)])
    sh.t(0.6, 1.5, W, 0.26, 'SimReal, a neolab building Data Infra, RL environments and recursive self-improvement (RSI) for future AI.', 10, C['grey'], SANS)
    cells = [
        ('收入', 'ARR 3,000万元', ['成立24天，零外部融资', '约450万美元'], T),
        ('开源', 'GitHub 602星', ['成立14天上线7款产品', '其中5个开源基准'], T),
        ('效果', '+12%', ['成立两周在交易环境跑通训练', '没见过的交易日上表现最高提升12%'], F),
        ('团队', '05后量化团队', ['剑桥、LSE、杜克数学', 'Jane Street、Citadel、Optiver、Millennium'], F),
        ('市场', '85亿美元', ['训练数据与RL环境，今天的年收入规模', 'Mercor 16个月收入涨27倍'], F),
        ('本轮融资', '4,000万元', ['人民币，投后估值5亿元', '投向环境、RSI研究与交付'], T),
    ]
    cw, xs = cols(3)
    for i, (k, v, d, acc) in enumerate(cells):
        x, y = xs[i % 3], 1.98 + (i // 3) * 2.05
        sh.rect(x, y, cw, 0.02, C['accent'] if acc else C['ink'])
        label(sh, x, y + 0.14, cw, k, C['accent'] if acc else None)
        sh.t(x, y + 0.42, cw, 0.56, v, 26, C['accent'] if acc else C['ink'], SERIF)
        sh.t(x, y + 1.08, cw, 0.8, d, 11.5, C['body'], line=1.15, gap=1)
    footer(sh)


def p_team(sh):
    header(sh, '团队', 0, [('三位21岁创始人，', F), ('放弃首年合计600万+薪酬的工作，搭建下一代训练引擎', T)])
    b, sz = C['body'], 11
    people = [
        ('Charles', 'CEO', [(R('United Stables首位员工：U稳定币', sz, b), BR(sz), R('一年从0做到14亿美元，一个月上线币安', sz, b)),
                            '负责机构合作，对接SIG、DRW等',
                            (R('汇丰港元稳定币项目唯一实习生，', sz, b), BR(sz), R('参与香港金管局合规', sz, b)),
                            '伦敦Citadel对冲基金实习'],
         ['LSE数学 · 深国交', '美国数学奥林匹克（USAMO）入围']),
        ('Henry', 'CTO', ['剑桥数学一等荣誉，奖学金获得者',
                          (R('剑桥机器学习暑研，导师Po-Ling Loh', sz, b), BR(sz), R('（国际数理统计学会会士）', sz, b)),
                          'Jane Street、Citadel、Optiver量化经历', '主导设计全部5个开源基准'],
         ['深国交', '剑桥数学竞赛全球前30 · 英国物理竞赛超级金奖']),
        ('Amaris', 'COO', ['Millennium香港数据科学家', 'Millennium另类数据团队首位应届生',
                           (R('参与筹办Plug and Play香港首场活动', sz, b), BR(sz), R('（联合香港科技园，200+人）', sz, b))],
         ['杜克数学与统计 · 上海包玉刚']),
    ]
    cw, xs = cols(3, 0.3)
    y0, ch = 2.42, 3.75
    for (name, role, lines, edu), x in zip(people, xs):
        sh.rect(x, y0, cw, ch, C['tint'])
        sh.text(x + 0.3, y0 + 0.16, cw - 0.6, 0.5, [para([R(name, 24, C['ink'], SERIF), R('   ' + role, 10, C['accent'], MONO),
                                                          R('   21岁', 11, C['ink'], SANS, True)])], 'ctr')
        sh.t(x + 0.3, y0 + 0.8, cw - 0.6, 2.2, lines, sz, b, line=1.12, gap=6)
        sh.rule(x + 0.3, y0 + 2.98, cw - 0.6, C['mid'])
        sh.t(x + 0.3, y0 + 3.08, cw - 0.6, 0.6, edu, 9.5, C['grey'], line=1.12, gap=1)
    footer(sh)


def p_problem(sh):
    header(sh, '问题', 1, [('AI会推理，', F), ('但干不好真实的活', T)])
    rows = [('交易', '回测一路上涨', '实盘行情里亏钱'), ('软件工程', '测试全部通过', '一上线就出错'),
            ('财务', '账看着做完了', '月底对不上'), ('事件预测', '分析头头是道', '结果一出就错')]
    y0, rh = 2.15, 0.66
    sh.rect(7.35, y0 - 0.36, 5.38, 0.36 + rh * len(rows), C['tint'])
    sh.t(2.4, y0 - 0.32, 4, 0.26, '看上去', 9.5, C['grey'], MONO, anchor='ctr')
    sh.t(7.6, y0 - 0.32, 4, 0.26, '实际上', 9.5, C['accent'], MONO, anchor='ctr')
    sh.rule(0.6, y0, W, C['ink'])
    for i, (k, a, b) in enumerate(rows):
        y = y0 + i * rh
        sh.t(0.6, y, 1.7, rh, k, 10.5, C['grey'], MONO, anchor='ctr')
        sh.t(2.4, y, 4.4, rh, a, 19, C['ink'], SERIF, anchor='ctr')
        sh.t(6.8, y, 0.5, rh, '→', 14, C['grey'], algn='ctr', anchor='ctr')
        sh.t(7.6, y, 5.0, rh, b, 19, C['ink'], SERIF, anchor='ctr')
        sh.rule(0.6, y + rh, W)
    sh.t(0.6, 5.0, 1.8, 0.6, '6/32', 30, C['accent'], SERIF, anchor='ctr')
    sh.t(2.5, 5.0, 9.5, 0.6, '前沿大模型拿真钱做交易，32轮只有6轮赚钱（Alpha Arena）', 13, C['body'], anchor='ctr')
    kicker(sh, 5.85, [('缺的是一个', F), ('按真实结果打分的训练场', T), ('。', F)], 21)
    footer(sh)


def p_why_now(sh):
    header(sh, '为什么是现在', 1, [('训练的瓶颈从数据转到环境，', F), ('钱已经在花，专业领域的好环境还很少', T)])
    items = [('01', '训练方式变了', ['推理模型靠能自动判对错的任务做强化学习', '公开的人类文本预计2026–2032年用完'], '能自动打分的环境成了稀缺资源'),
             ('02', '钱已经在花', ['一个强化学习任务卖200–2,000美元', 'Mercor 2026年7月收购环境公司Deeptune'], '训练环境成了实验室的固定采购'),
             ('03', '专业任务还没解决', ['大模型做交易、做财务仍频繁出错', 'Meta、OpenAI 9月推出替人办事的Agent'], '交易对错一目了然，最适合先做')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.3
    for (n, k, ev, m), x in zip(items, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.3, top + 0.22, 1, 0.22, n, 10, C['accent'], MONO, True)
        sh.t(x + 0.3, top + 0.48, cw - 0.6, 0.46, k, 22, C['ink'], SERIF)
        sh.t(x + 0.3, top + 1.1, cw - 0.6, 1.2, ev, 12, C['body'], line=1.15, gap=6)
        sh.rule(x + 0.3, top + 2.4, cw - 0.6, C['mid'])
        means(sh, x + 0.3, top + 2.52, cw - 0.6, m, 12.5)
    sh.rect(0.6, 5.28, W, 0.72, C['ink'])
    sh.t(0.85, 5.28, 1.3, 0.72, '窗口期', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.2, 5.28, W - 1.85, 0.72, [(R('RL环境赛道约20家早期公司，预计最后剩3–5家。', 13, C['onDark']),
                                       R('先做出效果的，先拿下长期合同', 13, C['onDarkHi'], SANS, True))], 13, anchor='ctr', line=1.1)
    note(sh, '来源：Epoch AI（2026年1月）；TechCrunch（2026年7月）；Meta、OpenAI发布（2026年9月）；Wing VC（2026年）。', 6.3)
    footer(sh)


def p_market(sh):
    header(sh, '市场空间', 1, [('今天85亿美元，', F), ('头部公司16个月收入涨27倍', T)])
    sh.rect(0.6, 1.72, W, 1.9, C['tint'])
    label(sh, 0.9, 1.97, 5, '今天：训练数据与RL环境')
    sh.t(0.9, 2.25, 5, 0.72, '85亿美元 / 年', 38, C['ink'], SERIF)
    sh.t(0.9, 3.05, 5.2, 0.3, '50余家供应商的收入合计', 11.5, C['body'])
    sh.t(6.0, 1.72, 0.7, 1.9, '→', 26, C['grey'], algn='ctr', anchor='ctr')
    label(sh, 6.9, 1.97, 5.5, '2030年：公司测算', C['accent'])
    sh.t(6.9, 2.25, 5.6, 0.72, '7,000亿美元 / 年', 38, C['accent'], SERIF)
    sh.t(6.9, 3.05, 5.6, 0.3, '按AI经济7万亿美元、10%投入训练估算', 11.5, C['body'])
    growth = [('27倍', ['Mercor年化收入16个月', '从7,500万涨到20亿美元'], '专家数据需求在爆发'),
              ('10倍', ['Mercor估值17个月', '从20亿涨到200亿美元（洽谈中）'], '资本在加注'),
              ('18倍', ['Snorkel AI新数据业务', '一年增长18倍'], '需求转向专家和环境')]
    cw, xs = cols(3)
    for (v, d, m), x in zip(growth, xs):
        sh.rect(x, 4.0, cw, 0.02, C['ink'])
        sh.t(x, 4.1, cw, 0.66, v, 36, C['ink'], SERIF)
        sh.t(x, 4.84, cw, 0.5, d, 12, C['body'], line=1.1, gap=0)
        means(sh, x, 5.38, cw, m, 12.5)
    note(sh, '来源：Deedy Das行业图谱（2026年7月）；Mercor（TechCrunch、Dealroom）；Snorkel AI公告（2026年9月）。2030年为公司测算。', 6.3)
    footer(sh)


def p_customers(sh):
    header(sh, '客户', 2, [('先卖给AI实验室，', F), ('再到企业和个人Agent团队', T)])
    cards = [('AI实验室', '现在', '买环境和数据，训练更强的模型', '环境授权、训练数据、评测'),
             ('企业Agent团队', '下一步', ['把业务流程做成练习题，', '上线前先练、先验收'], '定制环境与验收评测'),
             ('个人Agent团队', '再下一步', ['替人办事的Agent，', '要练长任务、跨应用操作'], '长任务练习环境')]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.6
    for (k, tag, need, deliver), x in zip(cards, xs):
        dark = tag == '现在'
        main, sub, acc = (C['onDarkHi'], C['onDark'], C['accentLt']) if dark else (C['ink'], C['body'], C['accent'])
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        sh.t(x + 0.32, top + 0.24, cw - 0.6, 0.22, tag, 10, acc, MONO, True)
        sh.t(x + 0.32, top + 0.52, cw - 0.6, 0.5, k, 24, main, SERIF)
        sh.t(x + 0.32, top + 1.3, cw - 0.6, 0.2, '他们要什么', 9, sub, MONO)
        sh.t(x + 0.32, top + 1.56, cw - 0.64, 1.0, need, 13, main, line=1.2, gap=0)
        sh.rect(x + 0.32, top + 2.6, cw - 0.64, 0.01, DARK_RULE if dark else C['mid'])
        sh.t(x + 0.32, top + 2.72, cw - 0.6, 0.2, '我们卖什么', 9, acc, MONO)
        sh.t(x + 0.32, top + 2.96, cw - 0.6, 0.5, deliver, 13.5, main, SANS, True)
    kicker(sh, 5.6, [('同一套环境，', F), ('卖给三类客户', T)], 20)
    footer(sh)


RSI_IMG = os.path.join(ASSETS, 'v55', 'rsi.png')


def p_rsi(sh):
    header(sh, 'RSI方法论', 2, [('模型在环境里做题、被打分、补短板，', F), ('一圈比一圈强', T)])
    rx = 8.35
    label(sh, rx, 1.72, 4.3, '以交易为例', C['accent'])
    rows = [('做题', '模型写交易策略'), ('打分', '回放行情，按盈亏和风险打分'), ('找弱点', '找出哪类行情亏钱'),
            ('出新题', '多造这类行情的题'), ('再训练', '用高分记录训练模型')]
    for i, (k, d) in enumerate(rows):
        y = 2.0 + i * 0.5
        sh.rule(rx, y, 4.38, C['mid'])
        sh.t(rx, y, 1.05, 0.5, k, 13, C['accent'] if i == 3 else C['ink'], SERIF, anchor='ctr')
        sh.t(rx + 1.1, y, 3.3, 0.5, d, 11.5, C['body'], anchor='ctr')
    sh.rule(rx, 4.5, 4.38, C['mid'])
    sh.rect(rx, 4.72, 4.38, 1.35, C['ink'])
    sh.t(rx + 0.3, 4.82, 3.9, 0.22, '已验证 · 成立两周跑通', 9.5, C['accentLt'], MONO, True)
    sh.t(rx + 0.3, 5.06, 1.6, 0.6, '+12%', 32, C['accentLt'], SERIF, anchor='ctr')
    sh.t(rx + 1.75, 5.06, 2.5, 0.9, ['在交易环境里训练后，', '没见过的交易日上', '表现最高提升12%'], 11, C['onDarkHi'], line=1.1, gap=0, anchor='ctr')
    note(sh, '第4步“AI自动出题”研发中；+12%为多次独立复现中的最佳一次。', 6.36)
    footer(sh)


def p_rsi_road(sh):
    header(sh, 'RSI路线图', 2, [('从人出题，到AI自己出题，', F), ('再到AI改进AI', T)])
    stages = [
        ('01 · 已跑通', '在环境里练', '人出题，按真实结果打分，模型越练越强',
         [('DeepSeek-R1（2025）', '只用自动判分做强化学习，模型自己学会推理'),
          ('Silver & Sutton（2025）', 'AlphaGo负责人与图灵奖得主：下一代AI要在环境里，从真实反馈中学习')],
         ['成立两周在交易上跑通：', '没见过的交易日上最高+12%'], F),
        ('02 · 研发中', 'AI自己出题', '模型找到自己的弱点，针对性出新题',
         [('Absolute Zero（清华等，2025）', '不用人工数据，自己出题自己做，数学和代码推理超过用人工题训练的同类模型'),
          ('Self-Challenging Agents（Meta，2025）', 'Agent自己出任务自己练，工具调用成绩翻倍')],
         ['论文多在数学和代码上验证；', '我们把它做到交易和预测里'], F),
        ('03 · 终局', 'AI改进AI', 'AI做AI研究，改进自己的训练方法',
         [('AlphaEvolve（DeepMind，2025）', 'AI优化训练代码，Gemini训练耗时减少1%'),
          ('Darwin Gödel Machine（2025）', 'Agent改写自己的代码，SWE-bench 20%→50%'),
          ('OpenAI（2025年10月）', '公开目标：2028年3月做出全自动AI研究员')],
         ['MLBench：60道真实机器学习竞赛题，', '就是AI做研究的考场'], T),
    ]
    cw, xs = cols(3, 0.3)
    top, ch = 1.66, 3.92
    for i, ((tag, k, d, papers, us, dark), x) in enumerate(zip(stages, xs)):
        main, sub, acc, rule = ((C['onDarkHi'], C['onDark'], C['accentLt'], DARK_RULE) if dark else (C['ink'], C['body'], C['accent'], C['mid']))
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.28, cw - 0.56
        sh.t(ix, top + 0.18, iw, 0.22, tag, 9.5, acc, MONO, True)
        sh.t(ix, top + 0.4, iw, 0.46, k, 22, main, SERIF)
        sh.t(ix, top + 0.86, iw, 0.3, d, 11.5, main)
        sh.rect(ix, top + 1.2, iw, 0.01, rule)
        sh.t(ix, top + 1.28, iw, 0.2, '论文已证明', 9, sub, MONO)
        ps = []
        for j, (name, res) in enumerate(papers):
            ps.append(para([R(name, 10.5, main, SANS, True)], before=0 if j == 0 else 7))
            ps.append(para([R(res, 10.5, sub)], before=1, line=1.1))
        sh.text(ix, top + 1.5, iw, 1.58, ps)
        sh.rect(ix, top + 3.12, iw, 0.01, rule)
        sh.t(ix, top + 3.2, iw, 0.2, '衍真', 9, acc, MONO, True)
        sh.t(ix, top + 3.42, iw, 0.46, us, 11.5, main, SANS, True, line=1.1, gap=0)
        if i < 2:
            sh.t(x + cw, top, 0.3, ch, '→', 13, C['grey'], algn='ctr', anchor='ctr')
    sh.rect(0.6, 5.72, W, 0.62, C['ink'])
    sh.t(0.85, 5.72, 1.3, 0.62, '终局', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.0, 5.72, W - 1.6, 0.62, [(R('方法论文里已经有了，专业领域缺的是考场和裁判。', 13, C['onDark']),
                                      R('每个行业最强的AI，都出自衍真的环境', 13, C['onDarkHi'], SANS, True))], 13, anchor='ctr', line=1.1)
    note(sh, '来源：DeepSeek-R1（arXiv 2501.12948）；Silver & Sutton《Welcome to the Era of Experience》（2025）；Absolute Zero（arXiv 2505.03335）；'
             'Self-Challenging Language Model Agents（arXiv 2506.01716）；AlphaEvolve（Google DeepMind，2025年5月）；Darwin Gödel Machine（arXiv 2505.22954）；OpenAI直播（2025年10月）。', 6.42, 7.5)
    footer(sh)


SCREENSHOT = os.path.join(ASSETS, 'repos-2026-10-04.jpg')


def p_products(sh):
    header(sh, '产品', 2, [('成立14天上线7款产品，', F), ('5个开源基准拿下GitHub 602星', T)])
    sh.t(0.6, 1.6, 2.6, 0.95, '602', 64, C['accent'], SERIF, anchor='ctr')
    sh.t(3.0, 1.6, 3.9, 0.95, ['GitHub星标', '5个开源基准，截至2026年10月4日'], 13, C['ink'], SANS, True, anchor='ctr', line=1.15, gap=2)
    rows = [('SimReal-MLBench', '302星', '60道真实机器学习竞赛题，测AI做研究的能力'),
            ('Xitadel-QuantBench', '102星', '测AI做交易，和同一资产上的专业交易员比'),
            ('MathmoBench', '101星', '高难度数学题'),
            ('Puzzle Benchmark', '70星', '749道人工编写的推理题'),
            ('FuturePredict', '27星', '预测真实事件，结果揭晓后打分')]
    x0, y0, rh = 7.25, 1.72, 0.66
    sh.rule(x0, y0, W + 0.6 - x0, C['ink'])
    for i, (k, star, d) in enumerate(rows):
        y = y0 + i * rh
        sh.text(x0, y + 0.06, 5.5, 0.3, [para([R(k, 14, C['ink'], SERIF), R('   ' + star, 10.5, C['accent'], MONO, True)])], 'ctr')
        sh.t(x0, y + 0.36, 5.5, 0.26, d, 11, C['body'])
        sh.rule(x0, y + rh, W + 0.6 - x0)
    sh.t(x0, 5.2, 5.48, 0.6, [(R('已上线的训练环境', 10, C['grey'], MONO),), (R('交易  ·  AI研究  ·  事件预测', 14, C['ink'], SERIF),)], 11, line=1.2, gap=4)
    footer(sh)


def p_why_us(sh):
    header(sh, '为什么是我们', 3, [('懂行、有效、够快，', F), ('三件事都有结果', T)])
    cards = [('懂行', ['量化出身，按交易台的标准打分'], ['最强模型77分，其余24–30分，', '人类最佳80分：分数拉得开']),
             ('有效', ['同一个环境既能测，也能训'], ['训练后，没见过的交易日上', '表现最高提升12%']),
             ('够快', ['自己的环境搭建工具', '20万+可触达的专家'], ['成立14天上线7款产品，', '第3周交付数百万需求'])]
    cw, xs = cols(3, 0.3)
    top, ch = 1.72, 3.2
    for (k, how, res), x in zip(cards, xs):
        sh.rect(x, top, cw, ch, C['tint'])
        sh.t(x + 0.32, top + 0.24, cw - 0.6, 0.5, k, 26, C['ink'], SERIF)
        sh.t(x + 0.32, top + 0.9, cw - 0.6, 0.2, '怎么做', 9, C['grey'], MONO)
        sh.t(x + 0.32, top + 1.14, cw - 0.64, 0.7, how, 12.5, C['body'], line=1.15, gap=2)
        sh.rule(x + 0.32, top + 1.92, cw - 0.64, C['mid'])
        sh.t(x + 0.32, top + 2.04, cw - 0.6, 0.2, '结果', 9, C['accent'], MONO)
        sh.t(x + 0.32, top + 2.28, cw - 0.64, 0.8, res, 13.5, C['ink'], SANS, True, line=1.15, gap=0)
    sh.rect(0.6, 5.2, W, 0.72, C['ink'])
    sh.t(0.85, 5.2, 2.0, 0.72, '越做越难追', 15, C['accentLt'], SERIF, anchor='ctr')
    sh.t(2.95, 5.2, W - 2.6, 0.72, [(R('每接一个客户，题库、打分器和验证过的专家都沉淀下来，', 13, C['onDark']),
                                      R('下一个客户直接复用', 13, C['onDarkHi'], SANS, True))], 13, anchor='ctr', line=1.1)
    note(sh, '分数来自Xitadel公开测评；+12%为多次独立复现中的最佳一次。', 6.3)
    footer(sh)


def p_competition(sh):
    header(sh, '竞争', 3, [('对手强在规模，', F), ('我们强在专业领域的打分和速度', T)])
    xs, ws = [0.6, 3.6, 6.5, 9.45], [2.9, 2.8, 2.85, 3.28]
    heads = ['客户现在的选择', '强在哪', '弱在哪', '衍真']
    sh.rect(xs[3] - 0.15, 1.72, W + 0.6 - xs[3] + 0.15, 0.32 + 3 * 1.0, C['tint'])
    sh.rule(0.6, 1.72, W, C['ink'])
    for i, (x, w, h) in enumerate(zip(xs, ws, heads)):
        sh.t(x + (0.2 if i == 0 else 0), 1.74, w, 0.32, h, 9.5, C['accent'] if i == 3 else C['grey'], MONO, i == 3, anchor='ctr')
    rows = [('数据外包公司', 'Scale AI、Mercor等', '人多，实验室关系深', '靠堆人，做不了专业打分', '量化出身，打分专业'),
            ('公开榜单', 'Alpha Arena等', '方便比较模型', '只能测不能训，很快被刷爆', '同一个环境既能测也能训'),
            ('实验室自建', '', '贴合自己的需求', '缺行业规则和专家，占研究人力', '现成环境，20万+专家随时调用')]
    y0, rh = 2.04, 1.0
    for i, (k, names, pro, con, us) in enumerate(rows):
        y = y0 + i * rh
        sh.rule(0.6, y, W)
        sh.text(0.8, y, 2.7, rh, [para([R(k, 17, C['ink'], SERIF)])] + ([para([R(names, 9, C['grey'], MONO)], before=3)] if names else []), 'ctr')
        sh.t(xs[1], y, ws[1] - 0.2, rh, pro, 12.5, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[2], y, ws[2] - 0.2, rh, con, 12.5, C['body'], anchor='ctr', line=1.1)
        sh.t(xs[3], y, ws[3] - 0.15, rh, us, 13, C['ink'], SANS, True, anchor='ctr', line=1.1)
    sh.rule(0.6, y0 + 3 * rh, W)
    kicker(sh, 5.45, [('可比交易：', F), ('阿里拟领投UniPat 3亿美元，估值约25亿美元', T)], 18)
    note(sh, '来源：彭博（2026年9月）。', 6.0)
    footer(sh)


def p_business(sh):
    header(sh, '商业模式', 4, [('卖环境、卖数据、做联合训练，', F), ('模型一升级就要复购', T)])
    lines = [('环境授权', '按期限和范围收费', '交易、AI研究、事件预测等训练环境'),
             ('数据交付', '按量收费', ['环境每跑一遍就产出数据：', '轨迹、专家示范、评测题']),
             ('联合训练', '按项目收费', '用我们的环境，帮客户把模型练强')]
    cw, xs = cols(3, 0.3)
    for i, ((k, fee, what), x) in enumerate(zip(lines, xs)):
        sh.rect(x, 1.72, cw, 2.05, C['tint'])
        sh.t(x + 0.32, 1.92, 1, 0.22, f'0{i + 1}', 10, C['accent'], MONO, True)
        sh.t(x + 0.32, 2.2, cw - 0.6, 0.5, k, 24, C['ink'], SERIF)
        sh.t(x + 0.32, 2.82, cw - 0.64, 0.5, what, 11.5, C['body'], line=1.15, gap=0)
        sh.t(x + 0.32, 3.36, cw - 0.6, 0.3, fee, 13, C['ink'], SANS, True)
    label(sh, 0.6, 4.08, 6, '合作路径')
    path = ['需求与验收标准', '付费试点', '授权交付', '持续更新']
    pw, pxs = cols(4, 0.4)
    for i, (k, x) in enumerate(zip(path, pxs)):
        last = i == 3
        sh.rect(x, 4.36, pw, 0.62, C['ink'] if last else C['tint'])
        sh.t(x + 0.25, 4.36, pw - 0.5, 0.62, k, 16, C['accentLt'] if last else C['ink'], SERIF, anchor='ctr')
        if i < 3:
            sh.t(x + pw, 4.36, 0.4, 0.62, '→', 14, C['grey'], algn='ctr', anchor='ctr')
    kicker(sh, 5.35, [('模型每升级一次都要换题，', F), ('环境和数据要持续更新，客户持续付费', T)], 18)
    footer(sh)


def place(pic, cx, cy, scale=1.0):
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left, pic.top = int(cx * 914400 - pic.width / 2), int(cy * 914400 - pic.height / 2)


def p_traction(sh):
    header(sh, '当前进展', 4, [('2026年9月10日成立，', F), ('24天做到ARR 3,000万元', T)])
    ms = [('第14天', '7款', '产品上线', '交易环境同期跑通训练'),
          ('第3周', '数百万', '交付需求', ''),
          ('第24天', '3,000万元', 'ARR', '约450万美元，零外部融资'),
          ('第24天', '602', 'GitHub星标', '5个开源基准')]
    cw, xs = cols(4, 0.3)
    ax = 2.08
    sh.rect(0.6, ax, W, 0.02, C['ink'])
    for i, ((day, v, k, d), x) in enumerate(zip(ms, xs)):
        acc = i >= 2
        sh.rect(x, ax - 0.07, 0.16, 0.16, C['accent'])
        sh.t(x + 0.24, ax - 0.4, cw - 0.3, 0.26, day, 10.5, C['accent'], MONO, True, anchor='ctr')
        sh.t(x, ax + 0.22, cw, 0.7, v, 40 if len(v) <= 3 else 34, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(x, ax + 0.96, cw, 0.3, k, 13, C['ink'], SANS, True)
        if d:
            sh.t(x, ax + 1.28, cw, 0.3, d, 10.5, C['body'])
    sh.t(0.6, 3.72, W, 0.24, '“第N天”从成立日算起；截至2026年10月4日', 8.5, C['grey'])
    sh.rule(0.6, 4.12, W, C['ink'])
    label(sh, 0.6, 4.24, 6, '客户')
    sh.t(0.6, 4.5, 2.6, 0.7, '2家', 34, C['ink'], SERIF, anchor='ctr')
    sh.t(0.6, 5.22, 2.7, 0.3, '前沿实验室在谈', 13, C['ink'], SANS, True)
    label(sh, 3.75, 4.24, 8, '训练基础设施：同样算力，跑更多训练', C['accent'])
    infra = [('1/4', '单轮耗时降到1/4'), ('+64%', '同样资源，多完成64%的训练尝试'), ('−1/3', '存档耗时减少1/3，存档次数减半')]
    iw, ixs = cols(3, 0.3, 3.75, W + 0.6 - 3.75)
    for (v, k), x in zip(infra, ixs):
        sh.rect(x, 4.5, iw, 0.02, C['ink'])
        sh.t(x, 4.56, iw, 0.66, v, 32, C['accent'], SERIF, anchor='ctr')
        sh.t(x, 5.22, iw, 0.5, k, 11.5, C['body'], line=1.1)
    note(sh, '基础设施指标为公司内部测试，口径不同，不能相乘或相加。', 6.3)
    footer(sh)


def p_network(sh, logos, polymarket, schools):
    header(sh, '专家网络', 4, [('20万+可触达的专家，', F), ('覆盖21所顶尖高校', T)])
    label(sh, 0.6, 1.72, 8, '参与研发的从业者来自')
    slots = [0.6 + (i + 0.5) * W / 5 for i in range(5)]
    for pic, cx in zip(logos, slots):
        place(pic, cx, 2.3, 1.5 if pic.shape_id == 403 else 1.2)
    if polymarket is not None:
        place(polymarket, slots[4], 2.3, 1.2)
    sh.rule(0.6, 2.9, W, C['ink'])
    net = [('20万+', '可触达的专家', T), ('7,000+', '已报名的候补专家', F), ('21所', '高校的学生和校友', F)]
    for i, (v, k, acc) in enumerate(net):
        y = 3.12 + i * 0.86
        sh.t(0.6, y, 2.0, 0.7, v, 34, C['accent'] if acc else C['ink'], SERIF, anchor='ctr')
        sh.t(2.65, y, 2.2, 0.7, k, 12.5, C['body'], anchor='ctr')
    gx, gw = 4.95, W + 0.6 - 4.95
    cell = gw / 4
    for i, pic in enumerate(schools):
        w, h = pic.width / 914400, pic.height / 914400
        place(pic, gx + (i % 4 + 0.5) * cell, 3.45 + (i // 4) * 0.9, min(1.45 / w, 0.62 / h, 1.5))
    note(sh, '标识仅表示从业者任职机构或网络覆盖的部分高校，不代表背书。', 6.4, 8)
    footer(sh)


def p_raise(sh):
    header(sh, '本轮融资', 5, [('融资4,000万元，', F), ('投向环境、RSI研究和交付', T)])
    tiles = [('本轮融资（人民币）', '4,000万元', T), ('投后估值', '5亿元', F), ('本轮出让', '8%', F)]
    cw, xs = cols(3, 0.25)
    for (k, v, dark), x in zip(tiles, xs):
        sh.rect(x, 1.66, cw, 0.86, C['ink'] if dark else C['tint'])
        sh.t(x + 0.28, 1.76, cw - 0.5, 0.2, k, 9, C['accentLt'] if dark else C['grey'], MONO)
        sh.t(x + 0.28, 1.96, cw - 0.5, 0.5, v, 28, C['accentLt'] if dark else C['ink'], SERIF)
    dirs = [('01', '环境与数据', '把交易的打法复制到更多专业领域', '环境工程师  ·  专家网络扩容', '10+个专业领域环境', F),
            ('02', 'RSI研究', '让模型自己找弱点、自己出题、自己变强', '算力  ·  研究员', '开源大模型跑通RSI', T),
            ('03', '客户交付', '前沿实验室从试点做到长期合同', '交付团队  ·  商务', '累计20家客户', F)]
    cw, xs = cols(3, 0.25)
    top, ch = 2.72, 3.3
    for (n, k, d, inv, goal, dark), x in zip(dirs, xs):
        main, sub, acc, rule = ((C['onDarkHi'], C['onDark'], C['accentLt'], DARK_RULE) if dark else (C['ink'], C['body'], C['accent'], C['mid']))
        sh.rect(x, top, cw, ch, C['ink'] if dark else C['tint'])
        ix, iw = x + 0.3, cw - 0.6
        sh.t(ix, top + 0.2, 1, 0.22, n, 10, acc, MONO, True)
        sh.t(ix, top + 0.46, iw, 0.5, k, 24, main, SERIF)
        sh.t(ix, top + 1.04, iw, 0.6, d, 13, main, SANS, True, line=1.15)
        sh.rect(ix, top + 1.78, iw, 0.01, rule)
        sh.t(ix, top + 1.9, iw, 0.2, '投入', 9, sub, MONO)
        sh.t(ix, top + 2.12, iw, 0.3, inv, 12, main)
        sh.t(ix, top + 2.56, iw, 0.2, '本轮目标', 9, acc, MONO)
        sh.t(ix, top + 2.78, iw, 0.36, goal, 16, acc, SERIF)
    contact(sh, 6.3)
    footer(sh)


# ---------------------------------------------------------------- assembly ---
PAGES = [
    (p_cover, 0, {16, 17}, None),
    (p_overview, 1, {64}, 2),
    (p_team, 2, {74, 78, 79, 81, 82, 83}, 3),
    (p_problem, 4, {160}, 4),
    (p_why_now, 3, {126}, 5),
    (p_market, 13, {470}, 6),
    (p_customers, 20, {666}, 7),
    (p_rsi, 6, {237}, 8),
    (p_rsi_road, 5, {169}, 9),
    (p_products, 8, {300}, 10),
    (p_why_us, 15, {562}, 11),
    (p_competition, 14, {511}, 12),
    (p_business, 10, {376}, 13),
    (p_traction, 16, {607}, 14),
    (p_network, 11, {401, 402, 403, 404, 407}, 15),
    (p_raise, 17, {607}, 16),
]
SCHOOLS = list(range(429, 441))
POLYMARKET = os.path.join(ASSETS, 'logos', 'polymarket.png')


def copy_pics(src, dst, ids):
    """Copy pictures (crop and effects intact) from one slide to another, relinking their images."""
    out = []
    for shp in sorted((p for p in src.shapes if p.shape_id in ids), key=lambda p: ids.index(p.shape_id)):
        el = copy.deepcopy(shp._element)
        blip = el.find('.//' + qn('a:blip'))
        blip.set(qn('r:embed'), dst.part.relate_to(src.part.related_part(blip.get(qn('r:embed'))), RT.IMAGE))
        el.find('.//' + qn('p:cNvPr')).set('id', str(9000 + shp.shape_id))
        dst.shapes._spTree.append(el)
        out.append(next(p for p in dst.shapes if p.shape_id == 9000 + shp.shape_id))
    return out



def main(src, dst):
    prs = Presentation(src)
    slides = list(prs.slides)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    order = []
    for build, idx, keep, n in PAGES:
        CUR['n'] = n
        s, sid = slides[idx], ids[idx]
        for shp in list(s.shapes):
            if shp.shape_id not in keep:
                s.shapes._spTree.remove(shp._element)
        tree = s.shapes._spTree
        sh = S(5000)
        if build is p_network:
            for clr in s._element.cSld.bg.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr'):
                clr.set('val', C['paper'])                   # this page was dark in v17
            logos = [shp for shp in s.shapes if shp.shape_id in (401, 402, 403, 404)]
            for pic in logos + [shp for shp in s.shapes if shp.shape_id == 407]:
                recolor(pic, C['ink'])
            poly = s.shapes.add_picture(POLYMARKET, 0, 0, height=Inches(0.3)) if os.path.exists(POLYMARKET) else None
            build(sh, logos, poly, copy_pics(slides[12], s, SCHOOLS))
        elif build is p_team:
            pics = sorted([shp for shp in s.shapes if shp.shape_id in (78, 79, 81, 82, 83)], key=lambda p: p.left)
            for i, p in enumerate(pics):
                cx = 0.6 + (i + 0.5) * W / len(pics)
                p.left, p.top = int((cx - p.width / 914400 / 2) * 914400), int((1.98 - p.height / 914400 / 2) * 914400)
            build(sh)
        elif build is p_products:
            from PIL import Image
            im = Image.open(SCREENSHOT)
            buf = io.BytesIO()
            im.crop((0, 280, im.width, im.height)).save(buf, 'PNG')       # the pinned repos, without the page header
            buf.seek(0)
            s.shapes.add_picture(buf, Inches(0.6), Inches(2.62), width=Inches(6.3))
            build(sh)
        elif build is p_rsi:
            s.shapes.add_picture(RSI_IMG, Inches(0.45), Inches(1.6), width=Inches(7.6))
            build(sh)
        else:
            build(sh)
        frag = etree.fromstring(f'<p:spTree {NS}>' + ''.join(sh.xml) + '</p:spTree>')
        for child in list(frag):
            tree.append(child)
        order.append(sid)
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)
    for el in ids:
        if el not in order:
            prs.part.drop_rel(el.rId)
    prs.save(dst)
    print('wrote', dst, len(order), 'slides')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
